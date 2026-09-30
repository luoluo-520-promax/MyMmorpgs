# -*- coding: utf-8 -*-
"""在原 Word 文档中增补 P16 权威纵深加固（不新建文档）。"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC_PATH = Path(r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx")
BACKUP = DOC_PATH.with_suffix(".docx.bak-before-P16")


def set_paragraph_text(p, text: str) -> None:
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for run in p.runs[1:]:
        run.text = ""


def insert_paragraph_after(paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    if text:
        new_para.add_run(text)
    return new_para


def find_paragraph(doc, predicate):
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if predicate(t, p):
            return p
    return None


def already_has(doc, needle: str) -> bool:
    return any(needle in (p.text or "") for p in doc.paragraphs)


def main() -> None:
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    shutil.copy2(DOC_PATH, BACKUP)
    doc = Document(str(DOC_PATH))

    # ── 封面：生成日期 / 修订说明 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：") and "P16" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月25日（P16 权威纵深加固；P15 生成式 AI 沉浸补齐；"
                "P14 体验纵深补齐；P13 体验优化；P12 性能加固）",
            )
            break

    rev = find_paragraph(doc, lambda t, _: t.startswith("修订说明："))
    if rev and "P16" not in (rev.text or ""):
        set_paragraph_text(
            rev,
            "修订说明：在原有 P6–P15（含性能加固、体验优化、体验纵深补齐、七大短板加固与生成式 AI 沉浸）基础上，"
            "增补 P16 权威纵深加固：PhysicsAuthorityService 物理哈希软拉回、TerrainTopologyGraph 碎岩因果链、"
            "EcoGraphService 捕食矩阵涟漪、多端 HitBox/AimAssist、WorldExplorationBalancer 热力掉落、"
            "CoopNarrativeProxy 联机隔膜、CompanionGhost 伙伴路径、ShadowRecovery 断线快进、"
            "AuctionHouse PriceStability/跨服货架、PredictionVerdict 划痕补偿、"
            "OpenWorldConfigPatch 双缓冲 Schema Versioning；技术栈截至 2026-08-25。"
            "请在 Word 中全选后按 F9 更新目录。",
        )

    # ── 1.4 能力关键词 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if "P15 动态叙事" in t and "P16" not in t and "对齐二游" in t:
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；以及 P16 物理哈希权威校验、地形破坏因果链、生态图负反馈、热力探索调度、"
                "联机叙事代理、伙伴幽灵实体、断线快进和解、拍卖通胀熔断与配置热更双缓冲。",
            )
            break

    # ── 3.2 列表 bullet（插在 P15 bullet 后）──
    if not already_has(doc, "【P16 权威纵深加固】"):
        p15b = find_paragraph(doc, lambda t, _: t.startswith("【P15 生成式 AI 沉浸补齐】"))
        if p15b:
            insert_paragraph_after(
                p15b,
                "【P16 权威纵深加固】SceneMoveCmd.physicsStateHash + PhysicsAuthorityService（500ms 软拉回）+ "
                "TerrainTopologyGraph（DESTROY_CLIFF 射线+弹药因果）；EcoGraphService（PredatorPreyMatrix）→ "
                "POST /ecosystem/imbalance/ripple；CombatAssistService.hitboxScale / predictiveAimAssist；"
                "WorldExplorationBalancer + DynamicLootTier.heatDensityCoefficient（冷门×1.5，MsgId 2260）；"
                "CoopNarrativeProxy（WORLD_SHIFT_SHIELD MsgId 2451 + 助人之证）；"
                "CompanionGhostService（ENTITY=4 / PathNodes MsgId 2501）；"
                "NetworkEmulatorProxy + ServerShadowService.recoverWithFastForward；"
                "AuctionHouseService.PriceStabilityIndex + 跨服 10% 手续费；"
                "HitFeedbackService.PredictionVerdict + CompensateEffectId；"
                "OpenWorldConfigPatchService staging 双缓冲 + config_version/Git SHA。",
                style="List Bullet",
            )

    # ── 3.2.12 P16 产品说明（插在 P15 测试段之后、3.4 之前）──
    if not already_has(doc, "3.2.12 权威纵深加固"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("测试：AiP15CapabilitiesTest") and "ai-enhancement.md" in t,
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("3.2.11 生成式 AI 沉浸补齐"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "3.2.12 权威纵深加固（P16）",
                style="Heading 3",
            )
            sections = [
                (
                    "Normal",
                    "物理权威与碎岩因果：SceneMoveCmd 增加 physicsStateHash（速度矢量/碰撞法线 MD5）；"
                    "PhysicsAuthorityService 每 500ms 对比预期推演，偏差超阈值软拉回。"
                    "DESTROY_CLIFF 经 TerrainTopologyGraph 校验射线相交并扣除重型体力/炸弹弹药，杜绝隔空碎岩。",
                ),
                (
                    "Normal",
                    "生态图负反馈：EcoGraphService 内置 PredatorPreyMatrix；刷新看上下级存量比。"
                    "物种降至阈值触发 POST /ecosystem/imbalance/ripple → RegionImpactService.EventChain"
                    " 与 RegionTugOfWarService 初始贡献偏移（乱杀影响势力格局）。",
                ),
                (
                    "Normal",
                    "多端瞄准：移动端 BOSS 弱点 HitBox×1.15（PC 保持精准）；"
                    "predictiveAimAssist 下发 SuggestedTargetAngle，客户端阻尼吸附，服务端不强制改朝向。",
                ),
                (
                    "Normal",
                    "探索热力调度：RegionHeat 过去约 10 分钟格子停留 <5 人时 HeatDensityCoefficient=1.5，"
                    "透传「稀有度提升」（MsgId 2260）；WorldExplorationBalancer 按热度方差下发 compensateSpawnCmd 补高级精英。",
                ),
                (
                    "Normal",
                    "联机叙事：CoopRoom 非房主遇 STORY_INSTANCE_START 时播放 WORLD_SHIFT_SHIELD（MsgId 2451），"
                    "隐藏 CHAOS 层；访客助力关键剧情 BOSS 写入编年史「助人之证」微章。",
                ),
                (
                    "Normal",
                    "伙伴幽灵：ActorType COMPANION_GHOST（ENTITY=4）不挡战斗碰撞但进 AOI；"
                    "CompanionPathNodes（MsgId 2501）客户端插值；攀爬/游泳体力为玩家 50%，"
                    "玩家耗尽触发鼓励 Buff 短暂回体并挂钩好感。",
                ),
                (
                    "Normal",
                    "弱网与重连：NetworkEmulatorProxy 测随机丢 ACK/延迟 2s；"
                    "断线超 30s 保护期启用时间回溯仲裁 + ClientPredictedAction 队列一次性快进和解（不直接踢）。",
                ),
                (
                    "Normal",
                    "经济与配置：拍卖 PriceStabilityIndex（Z-Score>3σ 需定价理由+24h 冷却）；"
                    "跨服货架绑定 region_shard_id、手续费 10%。"
                    "配置写入 Redis:config:staging，Admin 发布后仅新 SceneInstance 生效；"
                    "进行中实例读 snapshot_cache，绑定 config_version/Git Commit SHA。",
                ),
                (
                    "Normal",
                    "预表现视觉补偿：HitFeedback MsgId 2090 增加 PredictionVerdict"
                    "（CONFIRMED_HIT / CONFIRMED_MISS / SOFT_ROLLBACK）；"
                    "SOFT_ROLLBACK 下发 CompensateEffectId（划痕火花）并允许 1s 弹性缓冲避免瞬移。",
                ),
                (
                    "List Bullet",
                    "API：POST .../physics/hash/validate、.../terrain/destroy-cliff、"
                    ".../ecosystem/imbalance/ripple、.../combat/assist/aim、.../hit/verdict、"
                    "GET .../explore/loot/heat-coeff、POST .../explore/balancer/tick、"
                    ".../coop/narrative/room|assist、.../companion/spawn|follow-path、"
                    ".../shadow/recover、.../auction/list-stable、"
                    ".../config/stage|publish|snapshot-instance；story/instance/start 支持 coopRoomId。",
                ),
                (
                    "List Bullet",
                    "协议：MsgId 2260 稀有度提升、2451 WORLD_SHIFT_SHIELD、2501 CompanionPathNodes；2090 载荷扩展 PredictionVerdict。",
                ),
                (
                    "List Bullet",
                    "测试：OpenWorldP16AuthorityFlowTest、OpenWorldP16EdgeCaseTest（common）；"
                    "OpenWorldP16BusinessFlowTest（scene Internal API 全链路）；"
                    "权威 docs/open-world-top-tier.md P16。",
                ),
            ]
            cur = h
            for style, text in sections:
                cur = insert_paragraph_after(cur, text, style=style)

    # ── 5.2.14 P16 实现对照 ──
    if not already_has(doc, "5.2.14 P16"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("测试：AiP15CapabilitiesTest / AiP15BusinessFlowTest"),
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("5.2.13 P15"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "5.2.14 P16 权威纵深加固（物理哈希 / 生态因果 / 多端瞄准 / 热力调度 / 联机叙事 / "
                "伙伴幽灵 / 断线快进 / 拍卖稳定 / 预测补偿 / 配置双缓冲）",
                style="Heading 3",
            )
            cur = insert_paragraph_after(
                h,
                "代码详见 docs/open-world-top-tier.md P16 节；门面 OpenWorldGameplayFacade；"
                "GET /gameplay/status 返回 p16 快照（physicsAuthority、ecoGraph、aimAssist、"
                "explorationBalancer、coopNarrative、companionGhost、shadowFastForward、"
                "auctionStability、predictionVerdict、configSchemaVersioning 等）。",
                style="Normal",
            )
            bullets = [
                "物理：SceneMoveCmd.physicsStateHash、PhysicsAuthorityService、TerrainTopologyGraph。",
                "生态：EcoGraphService + RegionImpactService.beginEventChain + RegionTugOfWarService。",
                "战斗辅助：CombatAssistService.hitboxScale / predictiveAimAssist；"
                "HitFeedbackService.PredictionVerdict；PrePlaybackService.compensateEffectId；"
                "MoveTrajectoryValidator.elasticBuffer。",
                "探索：WorldExplorationBalancer、CollectibleService.DynamicLootTier.heatDensityCoefficient。",
                "叙事：CoopNarrativeProxy、StoryStateMachine.startInstanceWithCoop、PlayerChronicle 助人之证。",
                "伙伴：CompanionGhostService、SceneActorService.ENTITY_COMPANION_GHOST=4。",
                "同步：NetworkEmulatorProxy、ServerShadowService.recoverWithFastForward。",
                "经济：AuctionHouseService.PriceStabilityIndex / Cross-Shard Shelf。",
                "配置：OpenWorldConfigPatchService.stagePatch|publishStaging|snapshotForInstance。",
                "测试：OpenWorldP16AuthorityFlowTest / OpenWorldP16EdgeCaseTest（common）；"
                "OpenWorldP16BusinessFlowTest（scene）。",
            ]
            for b in bullets:
                cur = insert_paragraph_after(cur, b, style="List Bullet")

    # ── 第十章成熟度 ──
    if not already_has(doc, "P16 权威纵深加固：物理哈希"):
        p15m = find_paragraph(doc, lambda t, _: t.startswith("P15 生成式 AI 沉浸补齐："))
        if p15m:
            insert_paragraph_after(
                p15m,
                "P16 权威纵深加固：物理哈希软拉回、碎岩因果链、生态图涟漪、热力掉落杠杆、"
                "联机 WORLD_SHIFT_SHIELD、伙伴幽灵路径、断线快进和解、拍卖 Z-Score 熔断、"
                "预测划痕补偿、配置双缓冲；已接入 OpenWorldGameplayFacade、InternalOpenWorldController。",
                style="List Bullet",
            )

    # ── 11.18 业务故事 ──
    if not already_has(doc, "11.18 物理哈希到配置双缓冲"):
        anchor = find_paragraph(doc, lambda t, _: t.startswith("设置页选择「感性」风格"))
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("11.17 世界共鸣叙事"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "11.18 物理哈希到配置双缓冲（P16）",
                style="Heading 2",
            )
            stories = [
                "客户端上报 SceneMoveCmd.physicsStateHash → POST /physics/hash/validate；"
                "篡改重力系数偏差超阈值 → softPullback 回正速度/重力。",
                "瞄准悬崖 POST /terrain/destroy-cliff（BOMB）→ 射线命中扣弹；"
                "隔空/无弹药返回 ray_miss 或 bomb_exhausted。",
                "过度猎杀野猪 → 存量低于阈值 POST /ecosystem/imbalance/ripple → "
                "狼群袭城事件链启动，拉锯深渊侧初始贡献上升。",
                "冷门格子 GET /explore/loot/heat-coeff → heatDensityCoefficient=1.5 并提示「稀有度提升」；"
                "POST /explore/balancer/tick 下发 compensateSpawnCmd。",
                "手机端 POST /combat/assist/device MOBILE → hitBoxScale=1.15；"
                "POST /combat/assist/aim 得 SuggestedTargetAngle；预测未命中 POST /hit/verdict SOFT_ROLLBACK 播划痕火花。",
                "联机房主带 coopRoomId 开传说任务 → 访客收 WORLD_SHIFT_SHIELD；"
                "助力 BOSS 后 POST /coop/narrative/assist 获「助人之证」。",
                "POST /companion/spawn + /companion/follow-path 下发路径节点插值跟随；"
                "玩家体力耗尽时伙伴鼓励 Buff 短暂回体。",
                "弱网丢 ACK（NetworkEmulatorProxy）后断线超 30s → POST /shadow/recover 快进和解，不踢下线。",
                "跨服挂单价超 3σ → 须填定价理由；POST /auction/list-stable crossShard=true 手续费 10%。",
                "Admin POST /config/stage（绑定 Git SHA）→ /config/publish；"
                "进行中深渊实例仍读 snapshot_cache 旧 floorHp，新开实例才吃新配置。",
            ]
            cur = h
            for s in stories:
                cur = insert_paragraph_after(cur, s, style="Normal")

    # ── 第十四章 API 速览追加 P16 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("内部 HTTP 前缀 /internal/**") and "P16：" not in t:
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；P16：/physics/hash/validate、/terrain/destroy-cliff、/ecosystem/imbalance/ripple、"
                "/combat/assist/aim、/hit/verdict、/explore/loot/heat-coeff、/explore/balancer/tick、"
                "/coop/narrative/*、/companion/spawn|follow-path、/shadow/recover、/auction/list-stable、"
                "/config/stage|publish|snapshot-instance；MsgId 2260/2451/2501。",
            )
            break

    # ── FAQ Q17 ──
    if not already_has(doc, "Q17：P16"):
        q16 = find_paragraph(doc, lambda t, _: t.startswith("Q16：P15"))
        anchor = q16
        for p in doc.paragraphs:
            if (p.text or "").startswith("答：ai-service InternalAiP15Controller"):
                anchor = p
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "Q17：P16 物理权威/生态涟漪/热力掉落/联机隔膜/伙伴幽灵/断线快进/拍卖熔断/配置双缓冲如何联调？",
                style="Heading 2",
            )
            insert_paragraph_after(
                h,
                "答：场景侧 InternalOpenWorldController：/physics/hash/validate、/terrain/destroy-cliff、"
                "/ecosystem/imbalance/ripple、/combat/assist/aim、/hit/verdict、/explore/loot/heat-coeff、"
                "/explore/balancer/tick、/coop/narrative/room|assist、/companion/spawn|follow-path、"
                "/shadow/recover、/auction/list-stable、/config/stage|publish|snapshot-instance；"
                "story/instance/start 带 coopRoomId。"
                "回归：mvn -pl mmorpg-common,scene-service test \"-Dtest=OpenWorldP16*\"。"
                "权威说明 docs/open-world-top-tier.md（P16）。",
                style="Normal",
            )

    # ── 总结：P5～P16 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if "它覆盖玩家入口、大世界（含 P5～P15" in t and "P16" not in t:
            set_paragraph_text(
                p,
                t.replace(
                    "大世界（含 P5～P15：探索纵深、客户端预测软回滚、动量继承、确定性破坏、探索活力、野外营地、周词缀与命运残响等）与生成式 AI 沉浸（动态叙事/常驻伙伴/战斗分析/截图路点/动态人格）",
                    "大世界（含 P5～P16：探索纵深、客户端预测软回滚、物理哈希权威、生态图涟漪、热力探索调度、联机叙事隔膜、伙伴幽灵、断线快进、拍卖熔断、配置双缓冲等）与生成式 AI 沉浸（动态叙事/常驻伙伴/战斗分析/截图路点/动态人格）",
                ),
            )
            break

    # ── 附录 C ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：") and "附录" not in t and "P16" not in t and "七大体验短板" in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月25日（P16 权威纵深加固；P15 生成式 AI 沉浸补齐；"
                "P14 体验纵深补齐 + 七大体验短板加固；P13 体验优化；P12 性能加固；P10/P11 原地增补）",
            )
            break
    # 附录列表里的生成日期 bullet
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：2026年08月25日（P15") and "P16" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月25日（P16 权威纵深加固；P15 生成式 AI 沉浸补齐；"
                "P14 体验纵深补齐 + 七大体验短板加固；P13 体验优化；P12 性能加固；P10/P11 原地增补）",
            )

    # ── 附录表格 open-world-top-tier ──
    for table in doc.tables:
        for row in table.rows:
            row_text = " ".join(c.text for c in row.cells)
            if "docs/open-world-top-tier.md" not in row_text:
                continue
            for cell in row.cells:
                if "P0～P14" in cell.text:
                    cell.text = cell.text.replace("P0～P14", "P0～P16")
                elif "P0～P15" in cell.text:
                    cell.text = cell.text.replace("P0～P15", "P0～P16")
                elif "P0～P11" in cell.text:
                    cell.text = cell.text.replace("P0～P11", "P0～P16")

    doc.save(str(DOC_PATH))
    print("OK: updated", DOC_PATH)
    print("BACKUP:", BACKUP)


if __name__ == "__main__":
    main()
