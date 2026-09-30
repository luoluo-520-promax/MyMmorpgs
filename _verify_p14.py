from docx import Document
path = r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx"
doc = Document(path)
out = r"c:\Users\ASUS\IdeaProjects\test\MyMmorpg\_tmp_docx_after_p14.txt"
with open(out, "w", encoding="utf-8") as f:
    for i,p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if not t: continue
        if any(k in t for k in ["P14","3.2.9","5.2.12","11.15","Q14","2026年08月24","P5～P14","vitality","predict/action"]):
            f.write(f"{i}|{p.style.name}|{t}\n")
    f.write("\n=== tables ===\n")
    for ti,t in enumerate(doc.tables):
        for ri,row in enumerate(t.rows):
            for ci,cell in enumerate(row.cells):
                if "open-world-top-tier" in cell.text or "P0" in cell.text:
                    f.write(f"t{ti} r{ri} c{ci}: {cell.text}\n")
print("done")