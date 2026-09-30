# -*- coding: utf-8 -*-
from docx import Document

path = r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx"
out = r"c:\Users\ASUS\IdeaProjects\test\MyMmorpg\_tmp_docx_outline.txt"
doc = Document(path)
lines = []
lines.append(f"=== paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        style = p.style.name if p.style else ""
        lines.append(f"{i}|{style}|{t}")
lines.append(f"=== tables: {len(doc.tables)}")
for ti, table in enumerate(doc.tables):
    lines.append(f"--- table {ti} rows={len(table.rows)} cols={len(table.columns)}")
    for ri, row in enumerate(table.rows):
        cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
        lines.append(f"  r{ri}: " + " | ".join(cells))
with open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("wrote", out)
