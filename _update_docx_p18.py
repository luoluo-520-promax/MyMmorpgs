# -*- coding: utf-8 -*-
"""在原 Word 文档中增补 P18 八大手感短板补齐（不新建文档）。"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC_PATH = Path(r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx")
BACKUP = DOC_PATH.with_suffix(".docx.bak-before-P18")


def set_paragraph_text(p, text: str) -> None:
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for run in p.runs[1:]:
        run.text = ""


def insert_paragraph_after(paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    if text:
        new_para.add_run(text)
    return new_para


def find_paragraph(doc, predicate):
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if predicate(t, p):
            return p
    return None


def already_has(doc, needle: str) -> bool:
    return any(needle in (p.text or "") for p in doc.paragraphs)


def main() -> None:
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    shutil.copy2(DOC_PATH, BACKUP)
    doc = Document(str(DOC_PATH))

    # ── 封面：生成日期 / 修订说明 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：") and "P18" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月26日（P18 八大手感短板补齐；P17 联机一致性加固；"
                "P16 权威纵深加固；P15 生成式 AI 沉浸补齐；P14 体验纵深补齐）",
            )
            break

    rev = find_paragraph(doc, lambda t, _: t.startswith("修订说明："))
    if rev and "P18" not in (rev.text or ""):
        set_paragraph_text(
            rev,
            "修订说明：在原有 P6–P17（含性能加固、体验优化、体验纵深补齐、七大短板加固、"
            "生成式 AI 沉浸、权威纵深加固、联机一致性加固）基础上，"
            "增补 P18 八大手感短板补齐：ActionState 动作状态机分层、CancelAction 取消优先级、"
            "钩锁预测 ACK 与异步审计、攀爬挂边/翻越、摄像机相对移动与视野锥软锁、"
            "冰面滑行摩擦参数、clientDeltaMs 动态输入窗口与 50ms 逻辑帧心跳、"
            "分帧采集打断续传；技术栈截至 2026-08-26。请在 Word 中全选后按 F9 更新目录。",
        )

    # ── 1.4 能力关键词（P17 bullet 后）──
    if not already_has(doc, "P18 八大手感"):
        p17b = find_paragraph(
            doc,
            lambda t, _: t.startswith("【P17 联机一致性加固】")
            or ("P17 联机一致性加固" in t and "WorldOwnershipContext" in t and t.startswith("【")),
        )
        if p17b is None:
            p17b = find_paragraph(doc, lambda t, _: "CoopNarrativeProxy HostMutexLock" in t)
        if p17b:
            insert_paragraph_after(
                p17b,
                "【P18 八大手感短板补齐】ActionState ASM（空中/地面/攀爬攻击分层）+ CancelAction 取消表"
                "（DODGE>HEAVY>NORMAL>JUMP）+ HitFeedback allowCancelDuringHitStop；"
                "GrappleNodeService.predictGrapple 即时 ACK + PhysicsAuthorityService 200ms 钩锁审计；"
                "MovementType.CLIMB_HANG/CLIMB_VAULT + HANG_RECOVER_RATE；"
                "SceneMoveCmd.cameraYaw/moveIntent/clientDeltaMs + 摄像机±60°软锁；"
                "TerrainMutationService 冰面 surfaceFriction + TraverseMomentumService.applyFriction；"
                "ServerShadowService.heartbeat 50ms + CollectibleService 分帧采集续传。",
                style="List Bullet",
            )

    # ── 3.2.14 P18 产品说明（插在 P17 测试段之后）──
    if not already_has(doc, "3.2.14 八大手感短板补齐"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("测试：OpenWorldP17CoopConsistencyFlowTest")
            and "open-world-top-tier.md P17" in t,
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("3.2.13 联机一致性加固"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "3.2.14 八大手感短板补齐（P18）",
                style="Heading 3",
            )
            sections = [
                (
                    "Normal",
                    "动作状态机（ASM）：SceneMoveCmd 增加 actionState（GROUND_IDLE/RUN、AIR_NORMAL/HEAVY、"
                    "CLIMB_ATTACK、SWIM_ATTACK）；MovementAdmissionService.validateActionState 绑定空中/地面/攀爬上下文；"
                    "AIR_NORMAL 下 1.5s 窗口最多 3 次空中轻击（airComboFloaty）；空中战斗消耗 AIR_COMBAT_COST。",
                ),
                (
                    "Normal",
                    "取消优先级：CancelAction 取消表 + ReactionValidator.tryCancel + PoiseService.resetStiffness；"
                    "InputBufferService.enqueue 携带 expectedCancelAction 与 clientDeltaMs 动态窗口；"
                    "HitFeedback MsgId 2090 增加 allowCancelDuringHitStop（小怪轻击可闪避取消，Boss 重击不可）。",
                ),
                (
                    "Normal",
                    "钩锁预测：GrapplePredictCsReq 即时 ACK（GRAPPLE_PREDICT_ACK=177）；"
                    "PhysicsAuthorityService 200ms 后审计落点偏差，>3m 软拉回；"
                    "GrapplePhysicsService.pullSpeedCurve 下发贝塞尔拉拽曲线，客户端本地模拟。",
                ),
                (
                    "Normal",
                    "攀爬挂边/翻越：MovementType.CLIMB_HANG（体力<5 且低速挂墙 3s，每秒恢复 10 体力）与 CLIMB_VAULT"
                    "（顶边 0.5m 内翻越瞬移至站立点，RetCode VAULT_SUCCESS=174，取消硬直）。",
                ),
                (
                    "Normal",
                    "摄像机相对移动：moveIntent（FORWARD/BACKWARD/LEFT/RIGHT）+ cameraYaw 重映射输入；"
                    "BACKWARD+DASH → DASH_BACKWARD；CombatAssistService 基于摄像机 ±60° 视野锥筛选软锁目标。",
                ),
                (
                    "Normal",
                    "冰面滑行：TerrainMutationService 冻结水面广播 surfaceFriction=0.1、brakeDeceleration=0.5；"
                    "TraverseMomentumService.applyFriction 客户端滑行衰减；PhysicsAuthorityService 冰面拉回阈值 2.0m。",
                ),
                (
                    "Normal",
                    "动态输入采样：clientDeltaMs + effectiveWindowMs=100+delta×2；"
                    "ServerShadowService.heartbeat 每 50ms 下发 logicFrameSeq；move/admit 回写 effectiveWindowMs。",
                ),
                (
                    "Normal",
                    "采集续传：CollectibleService.startCollect/tickCollect/pauseCollect；"
                    "受击 pauseCollect 后 5s 内可续传；进度>80% 自动完成；响应 collectedPercent（0~100）。",
                ),
                (
                    "List Bullet",
                    "API：POST .../move/admit（actionState/moveIntent/cameraYaw/clientDeltaMs）、"
                    ".../combat/cancel/try、.../input-buffer/enqueue、.../grapple/predict|audit、"
                    ".../shadow/heartbeat、.../collectible/start|tick|pause、"
                    ".../combat/assist/soft-lock（cameraYawDeg）；"
                    "physics/hash/validate（onIceSurface）。",
                ),
                (
                    "List Bullet",
                    "协议：scene_protocol ActionState/MoveIntent/GrapplePredictCsReq/VaultCsReq；"
                    "MovementType MOVE_CLIMB_HANG/MOVE_CLIMB_VAULT；RetCode 174~177。",
                ),
                (
                    "List Bullet",
                    "测试：OpenWorldP18FeelFlowTest、OpenWorldP18EdgeCaseTest（common）；"
                    "OpenWorldP18BusinessFlowTest（scene Internal API 全链路）；"
                    "权威 docs/open-world-top-tier.md P18。",
                ),
            ]
            cur = h
            for style, text in sections:
                cur = insert_paragraph_after(cur, text, style=style)

    # ── 5.2.16 P18 实现对照 ──
    if not already_has(doc, "5.2.16 P18"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("测试：OpenWorldP17CoopConsistencyFlowTest")
            and "OpenWorldP17BusinessFlowTest" in t,
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("5.2.15 P17"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "5.2.16 P18 八大手感短板补齐（ASM / 取消优先级 / 钩锁预测 / 攀爬挂边翻越 / "
                "摄像机移动 / 冰面滑行 / 动态输入 / 采集续传）",
                style="Heading 3",
            )
            cur = insert_paragraph_after(
                h,
                "代码详见 docs/open-world-top-tier.md P18 节；门面 OpenWorldGameplayFacade；"
                "GET /gameplay/status 返回 p18 快照（actionStateMachine、cancelPriority、grapplePredictAck、"
                "climbHangVault、cameraRelativeMove、iceSurfaceFriction、dynamicInputWindow、progressiveCollect 等）。",
                style="Normal",
            )
            bullets = [
                "ASM：ActionState、MovementAdmissionService.validateActionState、StaminaConsumeService.AIR_COMBAT_COST。",
                "取消：CancelAction、ReactionValidator.tryCancel、PoiseService.resetStiffness、"
                "InputBufferService、HitFeedbackService.allowCancelDuringHitStop。",
                "钩锁：GrappleNodeService.predictGrapple、PhysicsAuthorityService.scheduleGrappleAudit、"
                "GrapplePhysicsService.pullSpeedCurve。",
                "攀爬：MovementType.CLIMB_HANG/CLIMB_VAULT、StaminaConsumeService.HANG_RECOVER_RATE。",
                "摄像机：SceneMoveCmd.cameraYaw/moveIntent、CombatAssistService 视野锥软锁。",
                "冰面：TerrainMutationService.surfaceFriction、TraverseMomentumService.applyFriction、"
                "PhysicsAuthorityService.ICE_PULLBACK_METERS。",
                "输入：InputConfidenceAnalyzer.effectiveWindowMs、ServerShadowService.HEARTBEAT_INTERVAL_MS。",
                "采集：CollectibleService.startCollect/tickCollect/pauseCollect、collectedPercent。",
                "测试：OpenWorldP18FeelFlowTest / OpenWorldP18EdgeCaseTest（common）；"
                "OpenWorldP18BusinessFlowTest（scene）。",
            ]
            for b in bullets:
                cur = insert_paragraph_after(cur, b, style="List Bullet")

    # ── 第十章成熟度 ──
    if not already_has(doc, "P18 八大手感短板：ASM"):
        p17m = find_paragraph(doc, lambda t, _: t.startswith("P17 联机一致性加固："))
        if p17m:
            insert_paragraph_after(
                p17m,
                "P18 八大手感短板：ASM 动作分层、CancelAction 取消表、钩锁预测 ACK、攀爬挂边/翻越、"
                "摄像机相对移动、冰面滑行摩擦、clientDeltaMs 动态窗口、50ms 逻辑帧心跳、分帧采集续传；"
                "已接入 OpenWorldGameplayFacade、InternalOpenWorldController。",
                style="List Bullet",
            )

    # ── 11.20 业务故事 ──
    if not already_has(doc, "11.20 空中连击到采集续传"):
        anchor = find_paragraph(doc, lambda t, _: t.startswith("11.19 联机一致性"))
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: "P17" in t and "主机迁移" in t)
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "11.20 空中连击到采集续传（P18）",
                style="Heading 2",
            )
            stories = [
                "客户端 GLIDE + actionState=AIR_NORMAL 连续 3 次轻击 → POST /move/admit 返回 airComboHit 1~3，"
                "第 4 击 airComboCapped，服务端扣 AIR_COMBAT_COST。",
                "普攻后摇中 expectedCancelAction=DODGE → POST /combat/cancel/try 清除 Poise 硬直；"
                "Boss 重击 HitStop 返回 allowCancelDuringHitStop=false。",
                "POST /grapple/predict 得 ACK → 客户端立即出钩；200ms 后 POST /grapple/audit，"
                "偏差>3m 软拉回预测点。",
                "攀爬体力耗尽 movementType=CLIMB_HANG → 挂墙 3s 每秒回 10 体；"
                "顶边 CLIMB_VAULT → retcode=VAULT_SUCCESS 瞬移站立。",
                "moveIntent=BACKWARD + cameraYaw=90 + DASH → dashVariant=DASH_BACKWARD；"
                "软锁带 cameraYawDeg 仅在 ±60° 锥内修正。",
                "冰封河面 surfaceFriction=0.1 → 客户端 applyFriction 长距离滑行；"
                "physics/hash/validate onIceSurface=true 时 pullbackThresholdM=2.0。",
                "clientDeltaMs=33 → effectiveWindowMs=166；POST /shadow/heartbeat 递增 logicFrameSeq。",
                "POST /collectible/start → tick → pauseCollect（受击）→ 5s 内 start 续传 → collectedPercent 达 100 完成。",
            ]
            cur = h
            for s in stories:
                cur = insert_paragraph_after(cur, s, style="Normal")

    # ── 第十四章 API 速览追加 P18 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("内部 HTTP 前缀 /internal/**") and "P18：" not in t:
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；P18：/move/admit（actionState/cameraYaw/clientDeltaMs）、/combat/cancel/try、"
                "/input-buffer/enqueue、/grapple/predict|audit、/shadow/heartbeat、"
                "/collectible/start|tick|pause；RetCode 174~177。",
            )
            break

    # ── FAQ Q19 ──
    if not already_has(doc, "Q19：P18"):
        anchor = find_paragraph(doc, lambda t, _: t.startswith("答：场景侧 InternalOpenWorldController") and "/coop/ownership/bind" in t)
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("Q18：P17"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "Q19：P18 ASM/取消优先级/钩锁预测/攀爬挂边/摄像机移动/冰面滑行/动态输入/采集续传如何联调？",
                style="Heading 2",
            )
            insert_paragraph_after(
                h,
                "答：场景侧 InternalOpenWorldController：/move/admit（actionState、moveIntent、cameraYaw、"
                "clientDeltaMs）、/combat/cancel/try、/input-buffer/enqueue（expectedCancelAction）、"
                "/grapple/predict、/grapple/audit、/shadow/heartbeat、/collectible/start|tick|pause、"
                "/combat/assist/soft-lock（cameraYawDeg）、/physics/hash/validate（onIceSurface）。"
                "回归：mvn -pl mmorpg-common,scene-service test \"-Dtest=OpenWorldP18*\"。"
                "权威说明 docs/open-world-top-tier.md（P18）。",
                style="Normal",
            )

    # ── 总结：P5～P18 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if "P5～P17" in t and "P18" not in t:
            set_paragraph_text(
                p,
                t.replace("P5～P17", "P5～P18").replace(
                    "主机迁移、地形 TSV",
                    "主机迁移、地形 TSV、ASM 动作分层、取消优先级、钩锁预测、攀爬挂边翻越、"
                    "摄像机相对移动、冰面滑行、动态输入窗口、采集续传",
                ),
            )
            break

    # ── 附录 C 生成日期 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：2026年08月26日（P17") and "P18" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月26日（P18 八大手感短板补齐；P17 联机一致性加固；"
                "P16 权威纵深加固；P15 生成式 AI 沉浸补齐；P14 体验纵深补齐 + 七大体验短板加固）",
            )

    # ── 附录表格 open-world-top-tier ──
    for table in doc.tables:
        for row in table.rows:
            row_text = " ".join(c.text for c in row.cells)
            if "docs/open-world-top-tier.md" not in row_text:
                continue
            for cell in row.cells:
                if "P0～P17" in cell.text:
                    cell.text = cell.text.replace("P0～P17", "P0～P18")
                elif "P0～P16" in cell.text:
                    cell.text = cell.text.replace("P0～P16", "P0～P18")

    doc.save(str(DOC_PATH))
    print("OK: updated", DOC_PATH)
    print("BACKUP:", BACKUP)


if __name__ == "__main__":
    main()
