# -*- coding: utf-8 -*-
from docx import Document
path = r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx"
doc = Document(path)

rev = ("修订说明：在原有 P6–P12（含性能加固 Actor 信箱/对象池/AOI 增量等）基础上，新增 P13 体验优化"
       "（探索罗盘/地图标记/世界影响、区域特色移动/攀爬歇脚/通用探索套件、生态叙事/环境叙事/探索正向反馈、"
       "战斗辅助模式、资源自动化/灵活日常/智能养成），并增补 P14 体验纵深补齐（客户端预测与软回滚、立体移动动量继承、"
       "确定性破坏、探索活力/神瞳共鸣、野外营地/幻影借用/助战印记、多端辅助与连招宏、周词缀轮换/命运残响）；"
       "技术栈截至 2026-08-24。请在 Word 中全选后按 F9 更新目录。")

for p in doc.paragraphs:
    t = p.text or ""
    if t.startswith("修订说明："):
        p.text = rev
    if t.startswith("生成日期：") and "附录" not in t and p.style and p.style.name == "List Bullet":
        p.text = "生成日期：2026年08月24日（P14 体验纵深补齐；P13 体验优化；P12 性能加固；P10/P11 原地增补）"

for table in doc.tables:
    for row in table.rows:
        row_has_doc = any("docs/open-world-top-tier.md" in c.text for c in row.cells)
        if not row_has_doc:
            continue
        for cell in row.cells:
            if "P0～P11" in cell.text or "P0~P11" in cell.text:
                cell.text = cell.text.replace("P0～P11", "P0～P14").replace("P0~P11", "P0~P14")

doc.save(path)
print("fixed")