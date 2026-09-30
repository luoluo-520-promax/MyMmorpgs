# -*- coding: utf-8 -*-
"""修正 P15 插入后与既有「短板加固」章节的编号冲突。"""
from docx import Document

DOC_PATH = r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx"


def set_text(p, text: str) -> None:
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ""


def main() -> None:
    doc = Document(DOC_PATH)

    for p in doc.paragraphs:
        t = (p.text or "").strip()
        # 产品章：P15 改为 3.2.11（避开既有 3.2.10 短板加固）
        if t.startswith("3.2.10 生成式 AI 沉浸补齐"):
            set_text(p, t.replace("3.2.10", "3.2.11", 1))
        # 业务故事：P15 改为 11.17（避开既有 11.16 短板加固）
        if t.startswith("11.16 世界共鸣叙事到人格化建议"):
            set_text(p, t.replace("11.16", "11.17", 1))
        # FAQ：P15 改为 Q16（避开既有 Q15 短板加固）
        if t.startswith("Q15：P15 动态叙事"):
            set_text(p, t.replace("Q15：", "Q16：", 1))

    # 确认成熟度 P15 bullet 存在
    has_maturity = any(
        (p.text or "").startswith("P15 生成式 AI 沉浸补齐：") for p in doc.paragraphs
    )
    if not has_maturity:
        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph

        for p in doc.paragraphs:
            if (p.text or "").startswith("P14 体验纵深补齐："):
                new_p = OxmlElement("w:p")
                p._p.addnext(new_p)
                np = Paragraph(new_p, p._parent)
                try:
                    np.style = "List Bullet"
                except Exception:
                    pass
                np.add_run(
                    "P15 生成式 AI 沉浸补齐：动态叙事、常驻 AI 伙伴（MsgId 2500）、战斗连招诊断、"
                    "截图路点、动态人格；已接入 AiPlatformFacade、InternalAiP15Controller、PlayerAiController。"
                )
                break

    doc.save(DOC_PATH)
    print("OK: renumbered P15 headings to 3.2.11 / 11.17 / Q16")


if __name__ == "__main__":
    main()
