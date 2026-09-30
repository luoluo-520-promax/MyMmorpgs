# -*- coding: utf-8 -*-
"""生成 MyMmorpg 项目中文介绍 Word 文档（含目录字段与页码）。"""

from __future__ import annotations

import datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DESKTOP = Path.home() / "Desktop"
OUT_PATH = DESKTOP / "MyMmorpg项目详细介绍文档.docx"


def set_run_font(run, name_cn="微软雅黑", name_en="Calibri", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name_en
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)

    for style_name, size, color in [
        ("Heading 1", 18, RGBColor(0x1F, 0x4E, 0x79)),
        ("Heading 2", 14, RGBColor(0x2E, 0x75, 0xB6)),
        ("Heading 3", 12, RGBColor(0x2E, 0x75, 0xB6)),
    ]:
        st = styles[style_name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(14)
        st.paragraph_format.space_after = Pt(8)

    if "Caption Soft" not in [s.name for s in styles]:
        cap = styles.add_style("Caption Soft", WD_STYLE_TYPE.PARAGRAPH)
        cap.font.size = Pt(10)
        cap.font.italic = True
        cap.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        cap.font.name = "Calibri"
        cap._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_page(doc: Document):
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.5)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)


def add_page_number(paragraph):
    """在段落中插入「第 X 页 / 共 Y 页」。"""
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=RGBColor(0x66, 0x66, 0x66))

    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r1 = paragraph.add_run()
    r1._r.append(fld1)
    r1._r.append(instr)
    r1._r.append(fld2)
    set_run_font(r1, size=9, color=RGBColor(0x66, 0x66, 0x66))

    run2 = paragraph.add_run(" 页 / 共 ")
    set_run_font(run2, size=9, color=RGBColor(0x66, 0x66, 0x66))

    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText")
    instr2.set(qn("xml:space"), "preserve")
    instr2.text = " NUMPAGES "
    fld4 = OxmlElement("w:fldChar")
    fld4.set(qn("w:fldCharType"), "end")
    r2 = paragraph.add_run()
    r2._r.append(fld3)
    r2._r.append(instr2)
    r2._r.append(fld4)
    set_run_font(r2, size=9, color=RGBColor(0x66, 0x66, 0x66))

    run3 = paragraph.add_run(" 页")
    set_run_font(run3, size=9, color=RGBColor(0x66, 0x66, 0x66))


def setup_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("MyMmorpg 项目介绍文档")
    set_run_font(hr, size=9, color=RGBColor(0x88, 0x88, 0x88))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)


def enable_update_fields_on_open(doc: Document):
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def add_toc(doc: Document):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    hint = OxmlElement("w:t")
    hint.text = "（请用 Microsoft Word 打开本文件后，右键目录 →「更新域」→「更新整个目录」，即可生成带页码的完整目录。）"
    run._r.append(hint)
    run._r.append(fld_end)


def add_para(doc, text, *, first_line_indent=True, bold=False, size=12):
    p = doc.add_paragraph()
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=12)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run, size=12)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "2E75B6")
        shading.set(qn("w:val"), "clear")
        hdr[i]._tc.get_or_add_tcPr().append(shading)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    set_page(doc)
    setup_header_footer(doc)
    enable_update_fields_on_open(doc)

    # ========== 封面 ==========
    for _ in range(3):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MyMmorpg")
    set_run_font(r, size=36, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("大型多人在线角色扮演游戏（MMORPG）\n服务端项目详细介绍文档")
    set_run_font(r, size=18, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    today = datetime.date.today().strftime("%Y年%m月%d日")
    r = info.add_run(
        f"面向普通读者的通俗说明版\n"
        f"技术栈：Java 17 · Spring Boot 3 · 微服务可选\n"
        f"文档生成日期：{today}\n"
        f"版本说明：基于仓库当前实现能力整理"
    )
    set_run_font(r, size=12, color=RGBColor(0x55, 0x55, 0x55))

    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tip.add_run(
        "\n使用提示：用 Word 打开后若目录显示为提示文字，"
        "请按 Ctrl+A 全选，再按 F9 更新域，即可看到带页码的目录。"
    )
    set_run_font(r, size=10, color=RGBColor(0x88, 0x88, 0x88))

    page_break(doc)

    # ========== 目录 ==========
    h = doc.add_heading("目录", level=1)
    add_toc(doc)
    note = doc.add_paragraph()
    run = note.add_run(
        "说明：目录由 Word 自动根据各级标题生成，并带有对应页码。"
        "页脚显示「第 X 页 / 共 Y 页」。"
    )
    set_run_font(run, size=10, color=RGBColor(0x66, 0x66, 0x66))
    page_break(doc)

    # ========== 第1章 ==========
    doc.add_heading("第一章 写给普通人的开场白", level=1)

    doc.add_heading("1.1 这份文档是写给谁的？", level=2)
    add_para(
        doc,
        "这份文档不是只给程序员看的技术手册，而是希望让没有写过代码、或对游戏服务器只有模糊印象的人，"
        "也能搞清楚：MyMmorpg 是什么、它能做什么、它由哪些部分组成、为什么要拆成这么多模块、"
        "以及日常开发和运维大致是怎样运转的。"
    )
    add_para(
        doc,
        "阅读时可以把整份文档当成「一本介绍游戏后台工厂如何运转」的通俗读物。"
        "遇到英文术语时，文中会尽量用生活化比喻解释；文末还有名词对照表可供查阅。"
    )

    doc.add_heading("1.2 一句话讲清这个项目", level=2)
    add_para(
        doc,
        "MyMmorpg 是一套用 Java 编写的「网络游戏服务器」示例工程。"
        "它负责在玩家手机或电脑上的游戏客户端背后，默默处理登录、走路、打架、聊天、背包、任务、活动、充值等一切需要服务器裁决的事情。"
    )
    add_para(
        doc,
        "打个比方：客户端像是游戏的「前台门面」和「遥控器」；服务器才是真正的「裁判 + 账本 + 仓库」。"
        "谁赢了、道具有没有发出去、钱有没有到账，都要由服务器说了算，这样才不容易被作弊。"
    )

    doc.add_heading("1.3 为什么叫 MMORPG？", level=2)
    add_para(
        doc,
        "MMORPG 是英文 Massively Multiplayer Online Role-Playing Game 的缩写，中文常叫「大型多人在线角色扮演游戏」。"
        "典型特征是：很多人同时在线、各自扮演一个角色、在同一个（或分线的）世界里探索、社交、战斗和成长。"
        "本项目对标的就是这类游戏的服务端能力底座，而不是只做单机或小房间对战。"
    )

    doc.add_heading("1.4 项目当前定位", level=2)
    add_bullets(
        doc,
        [
            "它是一个可运行、可测试、可部署的服务端示例 / 教学与工程实践项目。",
            "支持「单体模式」（一个程序打包大部分能力，适合本地开发）和「微服务模式」（能力拆成多个独立服务，适合规模扩展）。",
            "部分玩法数值与结算仍是演示级，但架构骨架、安全校验、幂等发奖、热更新、观测指标等已经按生产思路落地。",
            "文档中的「已实现 / 演示 / 计划中」会如实区分，避免把愿景当成已完成功能。",
        ],
    )

    # ========== 第2章 ==========
    doc.add_heading("第二章 用生活比喻理解整套系统", level=1)

    doc.add_heading("2.1 把游戏服务器想成一座城", level=2)
    add_para(
        doc,
        "想象 MyMmorpg 是一座正在运转的游戏之城："
    )
    add_bullets(
        doc,
        [
            "城门口（网关 / 玩家服务入口）：检查你是不是合法居民，把你领到该去的地方。",
            "户籍处（玩家与账号）：记录你是谁、等级多少、有没有被封号。",
            "街道与广场（场景服务）：你在哪里走路、附近能看见谁、换地图怎么交接。",
            "竞技场（战斗服务）：真正算伤害、判定胜负。",
            "仓库（背包服务）：道具进出都要记账，防止「复制装备」。",
            "邮局与会客厅（大厅服务）：邮件、好友、排行榜。",
            "任务公告栏（任务服务）：接任务、交任务、记进度。",
            "活动中心（活动服务）：节日活动、世界 BOSS、充值进度等。",
            "商城收银台（商店服务）：下单、支付验签、对账、发货。",
            "市政厅（管理后台）：策划导入配置、处理投诉、一键热更新。",
            "城市广播站（聊天服务）：世界频道、附近频道等消息传递。",
            "技能学院（技能服务）：技能配置与释放相关能力。",
            "匹配大厅（匹配服务）：找水平接近的队友或对手。",
            "版本办公室（更新服务）：告诉客户端该下哪些资源包、版本是否过期。",
        ],
    )

    doc.add_heading("2.2 为什么不能全塞进一个文件？", level=2)
    add_para(
        doc,
        "早期可以「一个人干所有事」（单体），开发快、调试方便。"
        "但当同时在线人数变多、某个玩法（比如战斗）特别吃 CPU 时，"
        "就把「竞技场」单独扩建成更大的场馆（微服务拆分），其它部分不受影响。"
        "本项目刻意保留两种模式，方便从学习到扩容平滑过渡。"
    )

    doc.add_heading("2.3 客户端与服务器如何说话？", level=2)
    add_para(
        doc,
        "游戏里走路、放技能这类操作非常频繁，不适合每次都像打开网页那样慢悠悠请求。"
        "因此本项目主要用「长连接」：客户端连上服务器后保持一条（或几条）通道，"
        "用紧凑的二进制协议（Protobuf）来回传消息，延迟更低、流量更省。"
    )
    add_para(
        doc,
        "同时，后台管理、健康检查、内部服务互调等，仍会使用常见的 HTTP 接口。"
        "可以理解为：玩家游戏走「专线电话」，运维和内部协作走「公文系统」。"
    )

    # ========== 第3章 ==========
    doc.add_heading("第三章 玩家能体验到的能力全景", level=1)
    add_para(
        doc,
        "下面按「玩家视角」介绍系统已经具备或演示可用的能力。"
        "普通人可以把它当成功能清单；技术人员可据此对照模块。"
    )

    doc.add_heading("3.1 账号、角色与在线", level=2)
    add_bullets(
        doc,
        [
            "登录鉴权与会话管理：确认你是本人，并维持在线状态。",
            "重复登录踢下线：同一账号在别处登录时，会通知并关闭旧连接，避免「一人分身」造成数据混乱。",
            "在线状态可被好友系统读取（例如通过 Redis 记录谁在线）。",
        ],
    )

    doc.add_heading("3.2 大世界与场景", level=2)
    add_bullets(
        doc,
        [
            "场景分线：一张地图人太多时自动开新线，像地铁加开车厢。",
            "AOI（兴趣区域）：只同步你附近需要看见的人/物，而不是全图广播，省流量也更公平。",
            "动态域网格：按人口密度拆分/合并负责区域，为大规模同服做底座。",
            "无缝交接票据：跨服务器节点传送时，带着速度、朝向、区域等信息「交接护照」，减少瞬移感。",
            "断线重连宽限：短时间掉线可尝试恢复场景状态。",
        ],
    )

    doc.add_heading("3.3 战斗与成长玩法", level=2)
    add_bullets(
        doc,
        [
            "战斗开始 / 行动 / 结算流程由战斗服务裁决。",
            "挑战关卡：可联动真实开战流程。",
            "抽卡（卡池）：含保底、历史记录、背包幂等发奖，减少重复领取。",
            "肉鸽玩法：战斗结算回写，支持局内天赋加点。",
            "战斗日/周统计：便于排行、活动进度或运营分析。",
            "技能体系：技能配置与战斗释放相关能力独立成域。",
        ],
    )

    doc.add_heading("3.4 社交与轻社交", level=2)
    add_bullets(
        doc,
        [
            "好友与在线态。",
            "邮件及附件发奖（可落到背包，带幂等）。",
            "聊天频道能力。",
            "排行榜（等级、战力等），缓存未命中时可回退数据库 Top 列表并回填。",
            "世界事件临时小队：为某次活动临时组队，不强制绑定公会。",
        ],
    )

    doc.add_heading("3.5 任务、活动与商城", level=2)
    add_bullets(
        doc,
        [
            "任务配置可由 JSON 驱动，支持导入与热更新。",
            "活动系统：配置、进度、领奖；可对接战斗胜利、充值等事件推进进度。",
            "世界 BOSS / 动态事件：含伤害贡献与结算钩子（演示级结算可继续接完善）。",
            "商城订单、支付验签、发货、累充进度；强调幂等与对账，降低「付了钱没货 / 发两次货」风险。",
        ],
    )

    doc.add_heading("3.6 外观、匹配与更新", level=2)
    add_bullets(
        doc,
        [
            "皮肤衣柜与穿戴。",
            "匹配：按等级/战力/等待时间等做互补评分；支持队列分片与跨分片配对。",
            "客户端版本与资源更新清单，由更新服务与后台导入协同。",
        ],
    )

    doc.add_heading("3.7 AI 辅助（只读建议面）", level=2)
    add_para(
        doc,
        "项目提供 AI 助手兼容接口：可以回答玩家问题或给出规则型建议。"
        "默认可以是「规则教练」模式；若外部大模型不可用，会回退到进程内规则，不阻断游戏主流程。"
        "这意味着 AI 是「参谋」，不是「裁判」——真正改数值、发道具仍走正式业务链路。"
    )

    # ========== 第4章 ==========
    doc.add_heading("第四章 技术全景（尽量讲人话）", level=1)

    doc.add_heading("4.1 技术选型一览", level=2)
    add_table(
        doc,
        ["技术", "通俗理解", "在本项目中的作用"],
        [
            ["Java 17", "编写服务器的主流语言版本", "全部服务端业务代码的基础"],
            ["Spring Boot 3", "快速搭应用的「脚手架」", "每个服务的启动与配置框架"],
            ["Spring Cloud Alibaba", "微服务全家桶（可选）", "服务发现、远程调用等扩展"],
            ["Netty", "高性能网络引擎", "游戏长连接（TCP，以及可选 UDP/KCP）"],
            ["Protobuf", "紧凑的数据打包格式", "客户端与服务器消息协议"],
            ["MySQL", "可靠的关系型数据库", "账号、邮件、任务进度、订单等持久数据"],
            ["Redis", "超快的内存数据库", "在线态、匹配队列、排行榜、战斗状态、票据等"],
            ["RocketMQ", "消息队列（可开关）", "战斗结束、订单支付等异步事件"],
            ["Docker Compose", "一键拉起依赖环境", "本机 MySQL/Redis/Nacos 等"],
            ["TestNG / CI", "自动考试系统", "单元测试、集成测试、质量门禁"],
        ],
        col_widths=[3.2, 5.0, 7.0],
    )

    doc.add_heading("4.2 仓库里有哪些模块？", level=2)
    add_para(doc, "父工程名为 mmorpg-parent，下面挂着多个子模块。可以按「积木」理解：")
    add_table(
        doc,
        ["模块", "角色说明（通俗）"],
        [
            ["mmorpg-protocol", "「词典」：所有协议消息编号与结构定义"],
            ["mmorpg-domain-api", "「合同范本」：跨服务调用的接口约定"],
            ["mmorpg-common", "「公共工具箱」：网络、RPC、通用配置等"],
            ["mmorpg-gateway", "「总门岗」（可选）：统一入口、区域头、粘滞路由"],
            ["player-service", "「主大厅」：玩家入口、会话、默认单体时嵌入多域"],
            ["scene-service", "场景 / AOI / 怪物与大世界底座"],
            ["battle-service", "战斗计算与结算"],
            ["activity-service", "活动、世界事件、充值进度等"],
            ["bag-service", "背包与发奖"],
            ["shop-service", "商城、订单、支付与对账"],
            ["hall-service", "好友、邮件、排行、临时小队"],
            ["quest-service", "任务"],
            ["matchmaking-service", "匹配"],
            ["chat-service", "聊天"],
            ["skill-service", "技能"],
            ["admin-service", "运营后台、导入、热更、投诉、AI 草稿等"],
            ["update-service", "客户端版本与资源清单"],
            ["mmorpg-cli", "命令行小工具：导入配置等运维操作"],
        ],
        col_widths=[4.5, 10.7],
    )

    doc.add_heading("4.3 两种部署模式", level=2)
    doc.add_heading("4.3.1 单体模式（默认，适合学习与本地演示）", level=3)
    add_para(
        doc,
        "默认只启动 player-service（HTTP 端口通常为 8989），其它很多域能力以 Maven 依赖形式「嵌」在同一个进程里。"
        "就像把城中各部门暂时合署办公：沟通快、部署简单，一台电脑也能跑起来。"
    )
    add_numbered(
        doc,
        [
            "用 Docker Compose 启动 MySQL 与 Redis。",
            "复制 .env.example 为 .env（开发默认值通常可用）。",
            "启动 player-service（dev 配置）。",
            "客户端连接玩家服的 HTTP / 游戏端口即可联调。",
        ],
    )

    doc.add_heading("4.3.2 微服务模式（适合扩展与分工）", level=3)
    add_para(
        doc,
        "每个领域服务独立进程、独立端口。"
        "player-service 更像网关：把客户端命令转发到对应服务；服务之间用 Feign / HTTP 内部接口协作，"
        "关键链路还可用消息队列做最终一致。"
    )
    add_para(doc, "各服务常用 HTTP 端口如下（便于对照）：")
    add_table(
        doc,
        ["服务", "端口", "一句话职责"],
        [
            ["player-service", "8989", "客户端连接与玩家主入口"],
            ["scene-service", "8981", "场景与大世界"],
            ["chat-service", "8982", "聊天"],
            ["bag-service", "8983", "背包"],
            ["skill-service", "8984", "技能"],
            ["admin-service", "8985", "后台管理"],
            ["hall-service", "8986", "大厅社交"],
            ["quest-service", "8987", "任务"],
            ["matchmaking-service", "8988", "匹配"],
            ["shop-service", "8990", "商城与支付"],
            ["battle-service", "8991", "战斗"],
            ["activity-service", "8992", "活动"],
            ["update-service", "8993", "更新清单"],
            ["mmorpg-gateway", "8443（可配）", "可选统一网关"],
        ],
        col_widths=[4.5, 3.5, 7.2],
    )

    # ========== 第5章 ==========
    doc.add_heading("第五章 核心子系统详解", level=1)

    doc.add_heading("5.1 玩家入口与中心路由（Center）", level=2)
    add_para(
        doc,
        "当场景服务器不止一台时，需要有一个「总调度」知道玩家该去哪台机器。"
        "本项目支持本地目录模式与远端中心模式。"
        "跨节点传送时会签发一次性 session_ticket（可以想成临时通行证），目标节点验票后才能进入，防止伪造传送。"
    )

    doc.add_heading("5.2 场景、AOI 与大世界底座", level=2)
    add_para(
        doc,
        "AOI 用空间哈希把地图切成格子，只处理相邻格子里的对象，再按距离精筛。"
        "WorldZoneManager 则按密度管理 Zone。"
        "项目还提供 AOI 压力测试接口，用于单进程千人级查询基准（注意：这是基准测试，不等于真实集群千人同屏已完全验证）。"
    )
    add_para(
        doc,
        "反作弊服务会检查移动是否超速/瞬移、伤害是否异常等，并累计违规，必要时临时封禁。"
        "这相当于给城门口装测速摄像头与裁判复核。"
    )

    doc.add_heading("5.3 战斗域", level=2)
    add_para(
        doc,
        "战斗服务负责开战、行动推进与结算，并向外发出「战斗结束」等事件。"
        "活动、任务等系统可以订阅这些事件来更新进度，而不必每次同步死等，提高整体吞吐。"
        "战斗运行状态常放在 Redis；活跃战斗索引用集合结构维护，避免危险的全库扫描。"
    )

    doc.add_heading("5.4 背包、邮件与发奖幂等", level=2)
    add_para(
        doc,
        "网络游戏最怕两件事：道具复制、以及「网络抖动导致同一奖励发两次」。"
        "因此发奖链路强调幂等：同样的发奖凭证只成功一次。"
        "邮件附件也可以走统一发奖端口进入背包。"
    )

    doc.add_heading("5.5 商城、支付与对账", level=2)
    add_para(
        doc,
        "商城涉及真金白银，设计上更保守："
        "订单落库、支付事件走发件箱（outbox）、收件箱（inbox）防重复消费、钱包/道具流水账本、"
        "以及对账接口对比渠道账单差异。"
    )
    add_para(
        doc,
        "开发环境可用 MOCK 支付方便联调；生产环境会禁用 MOCK，并要求配置渠道密钥与验签。"
        "可以理解为：沙箱里可以用「假钞练手」，正式开业必须验真钞。"
    )

    doc.add_heading("5.6 活动与世界事件", level=2)
    add_para(
        doc,
        "活动服务管理活动配置与玩家进度，支持条件判定、阶段、商店货品、展示文案、导入校验等较完整的配置模型。"
        "它还能消费「商店已支付」「战斗结束」等消息，推进充值进度或胜利次数类任务。"
        "世界 BOSS / 动态事件提供排期、伤害统计与结算钩子，便于继续接邮件或背包发奖。"
    )

    doc.add_heading("5.7 匹配系统", level=2)
    add_para(
        doc,
        "匹配队列放在 Redis，用分布式锁语义（SET NX + Lua）降低并发下的错配。"
        "队列可分片，并支持跨分片配对。"
        "评分会考虑等级、战力、等待时间等因素，让「等太久的人」更容易被照顾到。"
    )

    doc.add_heading("5.8 管理后台与热更新", level=2)
    add_para(
        doc,
        "admin-service 面向运营与策划："
        "配置导入（活动、任务、商城、更新清单等）、投诉处理、操作日志、AI 辅助草稿、"
        "以及分阶段热更新（activity → update → player 配置缓存）。"
    )
    add_para(
        doc,
        "热更新就像「不关城门的情况下更换公告牌」：先更新活动，再更新客户端清单，再清玩家侧缓存，"
        "并留下发布审计记录，便于追溯「谁在什么时候发布了什么」。"
    )
    add_para(
        doc,
        "后台有 IP 白名单、HMAC 签名或 API Key、权限点（如导入活动、处理投诉）等安全措施。"
        "生产环境会强制开启鉴权并拒绝弱默认密钥。"
    )

    doc.add_heading("5.9 网关与就近接入", level=2)
    add_para(
        doc,
        "可选的 mmorpg-gateway 可做统一入口。"
        "区域感知过滤器会注入区域相关请求头，配合基础设施层的智能 DNS / GSLB，实现「玩家尽量连更近的接入点」。"
        "粘滞路由会尽量把同一账号粘到同一玩家服实例，减少会话漂移。"
    )

    # ========== 第6章 ==========
    doc.add_heading("第六章 数据、消息与一致性（通俗版）", level=1)

    doc.add_heading("6.1 为什么既有 MySQL 又有 Redis？", level=2)
    add_para(
        doc,
        "MySQL 像「保险箱账本」：适合需要长期保存、可查询、可对账的数据（账号、邮件、订单、任务进度等）。"
        "Redis 像「前台便签和计分板」：极快，适合在线状态、排行榜、匹配队列、短期票据、热点缓存。"
        "两者配合：重要结果落保险箱，高频读写用便签加速。"
    )

    doc.add_heading("6.2 多库拆分", level=2)
    add_para(
        doc,
        "Docker 初始化会准备多个库（如 player / battle / activity / update），"
        "对应微服务拆分后的数据边界，避免所有表挤在一个库里难以扩展。"
    )

    doc.add_heading("6.3 消息队列在做什么？", level=2)
    add_para(
        doc,
        "有些事情不必「当场同步办完」。例如战斗刚结束，活动进度可以稍后更新。"
        "消息队列像邮局：战斗服务寄出「战斗结束」信件，活动服务与任务服务各自领取处理。"
        "开发环境可关闭 MQ（用空实现），生产若依赖这些最终一致链路则需显式开启。"
    )
    add_para(doc, "当前强调必须可靠投递的事件包括但不限于：")
    add_bullets(
        doc,
        [
            "商店订单已支付 → 活动充值进度。",
            "战斗结束 → 活动胜利次数 / 日常任务进度（另有 HTTP 投影兜底路径）。",
        ],
    )

    doc.add_heading("6.4 「最终一致」是什么意思？", level=2)
    add_para(
        doc,
        "不是说可以错账，而是说：系统允许极短时间内各处数据暂时不完全一样，但通过重试、幂等和补偿，最终会达成正确状态。"
        "就像快递：下单瞬间仓库可能还没拣货，但流程走完后你一定能收到货（或进入可追溯的失败处理）。"
    )

    # ========== 第7章 ==========
    doc.add_heading("第七章 安全、风控与生产意识", level=1)

    doc.add_heading("7.1 内部接口为什么要签名？", level=2)
    add_para(
        doc,
        "服务与服务之间的 /internal/** 接口如果裸奔，内网被渗透后就可能直接发奖、改进度。"
        "因此使用 HMAC 签名与时间戳，证明请求来自可信调用方且未被篡改。"
        "生产环境启动时会校验密钥强度，弱默认值会直接拒绝启动（fail-fast）。"
    )

    doc.add_heading("7.2 后台安全", level=2)
    add_bullets(
        doc,
        [
            "IP 白名单：只允许办公网或堡垒机访问。",
            "API Key / HMAC：证明操作者身份。",
            "权限点（RBAC 思路）：导入配置与处理投诉等权限可分开。",
            "操作日志：出了问题能追溯。",
        ],
    )

    doc.add_heading("7.3 会话与传输", level=2)
    add_para(
        doc,
        "会话可用 RSA 密钥；开发未配置时可能临时生成，但生产必须固定密钥。"
        "网关与服务可配置 TLS，避免明文传输敏感信息。"
    )

    doc.add_heading("7.4 反作弊与支付风控", level=2)
    add_para(
        doc,
        "移动与伤害校验降低外挂收益；支付侧强调验签、对账、幂等履约。"
        "这些都是「防小人」的基础工程能力：不能指望玩家客户端诚实上报结果。"
    )

    # ========== 第8章 ==========
    doc.add_heading("第八章 本地运行、测试与质量保障", level=1)

    doc.add_heading("8.1 最小本地启动路径", level=2)
    add_numbered(
        doc,
        [
            "安装 JDK 17、Maven、Docker Desktop（或等价环境）。",
            "在项目根目录执行：docker compose up -d mysql redis。",
            "复制 .env.example 为 .env。",
            "启动：mvn -pl player-service -am spring-boot:run -Dspring-boot.run.profiles=dev。",
            "用客户端或调试工具连接玩家服进行验证。",
        ],
    )

    doc.add_heading("8.2 测试怎么分层？", level=2)
    add_table(
        doc,
        ["类型", "目的", "常见命令"],
        [
            ["单元测试", "快速验证单个类/方法逻辑", "mvn -B test"],
            ["集成测试", "带着真实依赖（如 MySQL/Redis）验证联动", "mvn -B test -Pintegration-tests"],
            ["质量门禁", "覆盖率与静态缺陷扫描", "mvn -B verify -Pquality-gate"],
        ],
        col_widths=[3.5, 6.5, 5.2],
    )
    add_para(
        doc,
        "测试框架以 TestNG 套件组织：每个模块通常有 testng-unit.xml 与 testng-integration.xml。"
        "CI（持续集成）会在代码推送后自动跑测试，减少「我本地能过、别人一拉就挂」的问题。"
    )

    doc.add_heading("8.3 可观测性", level=2)
    add_para(
        doc,
        "关键路径会打警告日志，并暴露 Prometheus 指标（如调度、MQ、商城相关计数）。"
        "运维可以据此看：请求是否变慢、消息是否堆积、支付是否异常增多。"
    )

    # ========== 第9章 ==========
    doc.add_heading("第九章 能力成熟度对照表", level=1)
    add_para(
        doc,
        "为避免误解，这里把能力分成三档。普通人看这一章，就能知道「哪些能当真用，哪些只是演示，哪些还在路上」。"
    )

    doc.add_heading("9.1 已实现（工程可用）", level=2)
    add_bullets(
        doc,
        [
            "单体 / 微服务双模式切换。",
            "内部 API HMAC 与生产密钥启动校验。",
            "Admin 导入、投诉、操作日志、分阶段热更与发布审计。",
            "AI 健康检查与问答兼容面（规则回退）。",
            "Center 路由、迁移票据、场景分线与重连宽限。",
            "AOI、动态域、无缝交接、世界事件与临时小队底座。",
            "反作弊底座、支付对账 SPI、Feign 熔断降级。",
            "好友/邮件/任务进度落库；匹配与排行榜等 Redis 能力。",
            "抽卡保底与历史、挑战开战、肉鸽局内成长等玩法链路。",
            "Docker Compose、CI、质量门禁等工程化设施。",
        ],
    )

    doc.add_heading("9.2 演示级（能跑但未完全产品化）", level=2)
    add_bullets(
        doc,
        [
            "匹配主要是双人等级/战力带宽配对。",
            "部分活动/战斗奖励数值为演示常量。",
            "世界事件奖励结算多为钩子，需继续接完整发奖闭环。",
            "AOI 压测是单进程基准，不等于集群千人同屏验收。",
            "RocketMQ 开发默认关闭。",
            "会话 RSA 未配置时开发态临时生成（生产会强制固定密钥）。",
        ],
    )

    doc.add_heading("9.3 计划中（路线图）", level=2)
    add_bullets(
        doc,
        [
            "继续拆分 battle/skill/activity 等 CommandPort，推进真正的 gateway-only。",
            "更完整的分布式锁方案、官方支付 SDK 直连、KCP 能力更深对齐。",
            "真实 GeoIP/GSLB Anycast、跨区域多活、反作弊规则引擎生产化、集群级压测。",
        ],
    )

    # ========== 第10章 ==========
    doc.add_heading("第十章 典型业务故事（帮助建立直觉）", level=1)

    doc.add_heading("10.1 故事一：小明登录并走进大世界", level=2)
    add_numbered(
        doc,
        [
            "小明打开客户端，完成登录，拿到会话凭证。",
            "玩家服确认身份，记录在线，并规划他应进入的场景节点/分线。",
            "小明在场景中移动；AOI 只把附近玩家同步给他。",
            "若该区域太拥挤，系统可能把他导向新分线或其它 Zone。",
            "若需要跨节点传送，系统签发一次性票据，目标节点验票后接收。",
        ],
    )

    doc.add_heading("10.2 故事二：打完一场挑战，活动进度+1", level=2)
    add_numbered(
        doc,
        [
            "客户端发起挑战，玩家服把开战命令转到战斗域。",
            "战斗服务计算过程并结算胜负。",
            "战斗结束事件发出（MQ 或兜底 HTTP）。",
            "活动服务幂等消费事件，更新「今日胜利次数」之类进度。",
            "小明在活动界面看到进度变化，达标后领奖进入背包。",
        ],
    )

    doc.add_heading("10.3 故事三：充值礼包到账", level=2)
    add_numbered(
        doc,
        [
            "小明在商城下单，生成订单记录。",
            "支付渠道回调，服务器验签确认「真的付了钱」。",
            "支付成功事件进入可靠投递链路。",
            "发货与累充进度更新都带幂等键，避免重复。",
            "若渠道账单与本地不一致，对账流程产出差异单，供人工或自动重试履约。",
        ],
    )

    doc.add_heading("10.4 故事四：策划上线新活动", level=2)
    add_numbered(
        doc,
        [
            "策划准备活动 JSON/CSV，经校验后通过 Admin 或 CLI 导入。",
            "管理员触发分阶段 reload。",
            "活动服、更新服、玩家服依次刷新配置缓存。",
            "发布审计记录本次操作。",
            "玩家无需停服即可看到新活动（视具体配置与客户端资源而定）。",
        ],
    )

    # ========== 第11章 ==========
    doc.add_heading("第十一章 团队角色如何协作", level=1)
    add_table(
        doc,
        ["角色", "主要关心什么", "会接触哪些模块"],
        [
            ["客户端开发", "协议、登录、场景同步、战斗表现", "protocol、player 入口、各玩法消息号"],
            ["服务端开发", "业务正确性、性能、拆分边界", "各 *-service、common、domain-api"],
            ["策划", "数值、任务、活动配置", "config、Admin 导入、热更"],
            ["运营 / 客服", "投诉、封禁、发奖补单", "admin、对账、操作日志"],
            ["运维 / SRE", "部署、监控、密钥、扩容", "Docker、网关、Redis、Prometheus"],
            ["测试", "回归、压测、支付与幂等用例", "TestNG、CI、质量门禁"],
        ],
        col_widths=[3.2, 5.5, 6.5],
    )

    # ========== 第12章 ==========
    doc.add_heading("第十二章 常见问题（FAQ）", level=1)

    faqs = [
        (
            "Q1：我没有编程基础，能看懂这个项目吗？",
            "可以先读第一到第三章和第十章故事。把每个服务当成城市部门，不必先学 Java。"
            "等建立直觉后，再按兴趣深入某一章技术细节。",
        ),
        (
            "Q2：这是完整上线的商业游戏吗？",
            "它是完整度较高的服务端工程示例与底座，含大量可运行能力；"
            "但部分数值与结算仍是演示级，商业上线还需要内容、合规、运维与更严格的压测打磨。",
        ),
        (
            "Q3：为什么文档里有那么多英文单词？",
            "工业界通用术语（如 Redis、MQ、HMAC）。本文尽量配了中文解释，文末名词表也可查阅。",
        ),
        (
            "Q4：单体和微服务我该选哪个？",
            "学习、毕业设计、本机演示：优先单体。"
            "需要多人协作开发、独立扩容战斗/场景：再开微服务开关。",
        ),
        (
            "Q5：没有客户端能不能玩？",
            "可以先起服务看健康检查与日志；完整玩法体验通常需要配套客户端或协议调试工具。"
            "仓库侧重服务端。",
        ),
        (
            "Q6：数据会不会丢？",
            "持久数据在 MySQL；热点在 Redis。"
            "关键发奖与支付链路强调幂等和账本。"
            "但仍需正规备份与生产级运维，工程示例不能替代完整容灾方案。",
        ),
        (
            "Q7：目录和页码在我电脑上不显示怎么办？",
            "请用 Microsoft Word 打开桌面上的 docx，全选后按 F9，或右键目录选择「更新域」。"
            "WPS 一般也支持，但个别版本对域更新支持略有差异。",
        ),
    ]
    for q, a in faqs:
        doc.add_heading(q, level=2)
        add_para(doc, a)

    # ========== 第13章 ==========
    doc.add_heading("第十三章 名词小词典", level=1)
    add_table(
        doc,
        ["名词", "一句话解释"],
        [
            ["MMORPG", "大型多人在线角色扮演游戏"],
            ["客户端", "玩家安装的游戏程序"],
            ["服务端", "架在机房/云上的裁判与账本程序"],
            ["单体", "多数功能装在一个进程里"],
            ["微服务", "按领域拆成多个可独立部署的服务"],
            ["网关", "统一入口与路由转发层"],
            ["协议 / Protobuf", "双方约定好的消息格式"],
            ["长连接", "连上后保持通道，适合高频交互"],
            ["AOI", "只同步兴趣区域内的可见对象"],
            ["分线", "同一地图开多个平行实例分流"],
            ["幂等", "同一请求执行多次结果仍正确（不重复发奖）"],
            ["HMAC", "用密钥生成的防篡改签名"],
            ["Outbox / Inbox", "发件箱/收件箱模式，保证消息可靠且不重复"],
            ["Feign", "声明式的服务间 HTTP 调用方式"],
            ["熔断降级", "下游挂了时快速失败并走备用逻辑，避免雪崩"],
            ["热更新", "尽量不停服刷新配置"],
            ["对账", "把渠道账单与自家订单逐笔核对"],
            ["CI", "代码变更后自动构建与测试"],
            ["Prometheus", "指标采集与监控体系中的常见组件"],
        ],
        col_widths=[4.0, 11.2],
    )

    # ========== 第14章 ==========
    doc.add_heading("第十四章 延伸阅读与仓库文档地图", level=1)
    add_para(doc, "若需要更偏技术实现的细节，可在仓库中继续阅读：")
    add_table(
        doc,
        ["文件", "内容"],
        [
            ["README.md", "快速开始与模块一览"],
            ["DEPLOYMENT.md", "部署模式、环境变量、已实现/演示/计划"],
            [".env.example", "全量环境变量样例"],
            ["TESTING.md", "TestNG 测试套件说明"],
            ["docs/service-boundaries.md", "服务边界与同步/异步契约"],
            ["docs/architecture-gateway.md", "网关化与 common 拆分路线"],
            ["docs/db-migration.md", "数据库迁移相关说明"],
            ["docs/resilience-observability.md", "韧性与可观测性"],
        ],
        col_widths=[5.5, 9.7],
    )

    # ========== 第15章 ==========
    doc.add_heading("第十五章 总结", level=1)
    add_para(
        doc,
        "MyMmorpg 试图回答一个实际问题：如何用现代 Java 技术栈，搭出一套既适合学习演示、又具备走向微服务扩展潜力的 MMORPG 服务端。"
        "它覆盖了玩家入口、大世界场景、战斗、社交、任务、活动、商城支付、后台运营、安全校验与工程化测试等关键板块。"
    )
    add_para(
        doc,
        "对普通人而言，记住三句话就够："
        "第一，客户端负责表现，服务器负责裁决与记账；"
        "第二，系统可以先「合署办公」（单体），再「分局扩容」（微服务）；"
        "第三，发奖与支付必须以幂等、验签、对账为底线，因为这是信任的基础。"
    )
    add_para(
        doc,
        "希望这份文档能帮助你建立整图直觉。若你是学生，可用它作为课程设计说明的骨架；"
        "若你是转行者，可用它作为阅读源码的导航；若你是管理者，可用它评估项目完整度与后续投入重点。"
    )

    # 附录
    doc.add_heading("附录 A 服务端口速查（再贴一次）", level=1)
    add_para(
        doc,
        "便于打印单页查阅：player 8989；scene 8981；chat 8982；bag 8983；skill 8984；"
        "admin 8985；hall 8986；quest 8987；matchmaking 8988；shop 8990；"
        "battle 8991；activity 8992；update 8993；gateway 8443（可配置）。"
    )

    doc.add_heading("附录 B 文档信息", level=1)
    add_bullets(
        doc,
        [
            f"生成日期：{today}",
            "文档语言：简体中文",
            "文档目标：面向普通人的项目全景介绍，兼顾必要技术准确性",
            "排版：自动目录域 + 页脚页码（第 X 页 / 共 Y 页）",
            f"输出路径：{OUT_PATH}",
        ],
    )

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH))
    print(f"OK: {OUT_PATH}")


if __name__ == "__main__":
    build()
