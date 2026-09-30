# -*- coding: utf-8 -*-
"""更新桌面原文档：写入 AI 平台集成说明（不新建文档）。"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

DOCX = Path(r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx")
BACKUP = DOCX.with_suffix(".docx.bak")

NEW_37 = (
    "3.7 AI 辅助面（AI 平台已集成）\n"
    "AI 是「参谋」不是「裁判」，分两条物理隔离管道：管道 A（战斗 Tick）仅用规则/行为树/DDA，禁止 LLM；"
    "管道 B（交互/运营）走 ai-service 或 player-service 内嵌 AiPlatformFacade，超时/熔断后回退 rule-coach。\n"
    "玩家侧：POST /api/player/ai/advise（个性化顾问，含推荐）· /recommend · /support（RAG 智能客服）· /tactical（实时战术）· /profile。\n"
    "NPC：短期会话记忆 scene-service NpcDialogueMemory + 长期情感向量 NpcLongTermMemory（信任/好感/愤怒/恐惧）。\n"
    "战斗：DynamicDifficultyService + BossStrategyTier 自适应档位；GET /internal/battle/dda/evaluate 含 strategyTier。\n"
    "运营：Admin 活动草案自动 contentValidate；流失预测联动 retention 干预（礼包/推送）；MLOps 模型注册/漂移告警。\n"
    "微服务：ai-service:8995；单体：game.ai.remote=false 时逻辑嵌入 player-service/admin-service。"
)

AI_MODULE_LINE = "ai-service\n统一 AI 推理平台：推荐/客服/NPC 长期记忆/内容校验/流失干预/MLOps（8995）"

NEW_SECTION_53 = (
    "5.3.1 AI 推理平台（ai-service + AiPlatformFacade）\n"
    "模块：ai-service（8995）+ mmorpg-common/ai/*。单体嵌入 player-service（game.ai.remote=false）；"
    "微服务 Feign AiPlatformClient + Gateway Path=/internal/ai/**。\n"
    "核心 API：POST /internal/ai/infer · /recommend/items · /support/ask · /tactical/advise · "
    "/retention/evaluate · /content/generate|validate · /npc/bond · GET /mlops/drift。\n"
    "玩家 API：/api/player/ai/advise|recommend|support|tactical|profile（需 Token）。\n"
    "测试：AiPlatformBusinessFlowTest、AiInferenceGatewayTest、DynamicDifficultyFlowTest（strategyTier）、NpcMemoryMoodFlowTest。"
)

OLD_REV_MARKERS = ("P14 体验纵深补齐", "P14 + AI 平台")
NEW_REV = "生成日期：2026年08月24日（P14 + AI 平台 ai-service 集成；P13 体验优化；P12 性能加固）"

OLD_SUMMARY = "智能 NPC / AI 顾问 / Boss BT / Admin AI 内容草稿。"
NEW_SUMMARY = (
    "智能 NPC（短期+长期记忆）/ AI 顾问+推荐+客服+战术 / Boss BT+DDA / "
    "Admin AI 草稿+自动校验+流失干预 / ai-service 8995。"
)


def set_paragraph_text(para, text: str) -> None:
    """整段替换文本（保留首 run 样式）。"""
    if not para.runs:
        para.add_run(text)
        return
    para.runs[0].text = text
    for run in para.runs[1:]:
        run.text = ""


def replace_paragraph_containing(doc: Document, needle: str, new_text: str) -> bool:
    for para in doc.paragraphs:
        if needle in para.text:
            set_paragraph_text(para, new_text)
            return True
    return False


def replace_in_paragraph(doc: Document, needle: str, old: str, new: str) -> bool:
    for para in doc.paragraphs:
        if needle in para.text and old in para.text:
            set_paragraph_text(para, para.text.replace(old, new))
            return True
    return False


def insert_before_heading(doc: Document, heading_needle: str, new_text: str) -> bool:
    if any(heading_needle in p.text and "5.3.1" in p.text for p in doc.paragraphs):
        return False
    for i, para in enumerate(doc.paragraphs):
        if heading_needle in para.text and "5.3.1" not in para.text:
            new_para = para.insert_paragraph_before(new_text)
            set_paragraph_text(new_para, new_text)
            return True
    return False


def insert_paragraph_after(paragraph, text: str = "") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def insert_module_after_guild(doc: Document) -> bool:
    for i, para in enumerate(doc.paragraphs):
        t = para.text.strip()
        if "guild-service" in t and "8998" in t and "远征" in t:
            if i + 1 < len(doc.paragraphs) and "ai-service" in doc.paragraphs[i + 1].text:
                return False
            title = insert_paragraph_after(para, "ai-service")
            set_paragraph_text(title, "ai-service")
            desc = insert_paragraph_after(title, AI_MODULE_LINE.split("\n", 1)[1])
            set_paragraph_text(desc, AI_MODULE_LINE.split("\n", 1)[1])
            return True
    return False


def main() -> None:
    if not DOCX.exists():
        raise SystemExit(f"文档不存在: {DOCX}")
    shutil.copy2(DOCX, BACKUP)
    doc = Document(str(DOCX))

    # 3.7 正文（跳过目录行）
    for para in doc.paragraphs:
        t = para.text.strip()
        if t.startswith("3.7 AI") and "辅助面" in t and "\t" not in t:
            set_paragraph_text(para, NEW_37)
            break
    else:
        replace_paragraph_containing(doc, "AI 是「参谋」不是「裁判」", NEW_37)

    insert_module_after_guild(doc)

    replace_in_paragraph(
        doc,
        "gacha",
        "gacha 8994；guild-service 8998",
        "gacha 8994；ai-service 8995（AI 推理网关）；guild-service 8998",
    )

    replace_in_paragraph(
        doc,
        "GACHA_REMOTE",
        "GAME_SCENE/CHAT/BAG/SKILL/HALL/QUEST/MATCH/BATTLE/ACTIVITY/UPDATE/SHOP/GACHA_REMOTE_ENABLED=true",
        "GAME_SCENE/CHAT/BAG/SKILL/HALL/QUEST/MATCH/BATTLE/ACTIVITY/UPDATE/SHOP/GACHA_REMOTE_ENABLED=true；"
        "GAME_AI_REMOTE_ENABLED=true（ai-service:8995）",
    )

    for para in doc.paragraphs:
        if para.text.startswith("生成日期：") and "2026年08月24日" in para.text:
            if "ai-service" not in para.text:
                set_paragraph_text(para, NEW_REV)
            break

    replace_paragraph_containing(doc, OLD_SUMMARY, NEW_SUMMARY)
    insert_before_heading(doc, "5.3 战斗域与 AI", NEW_SECTION_53)

    doc.save(str(DOCX))
    print(f"已更新: {DOCX}")
    print(f"备份: {BACKUP}")


if __name__ == "__main__":
    main()
