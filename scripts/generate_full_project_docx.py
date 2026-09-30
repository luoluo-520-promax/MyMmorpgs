# -*- coding: utf-8 -*-
"""生成 MyMmorpg 全方位详细介绍 Word（目录域 + 页脚页码），输出到桌面。"""

from __future__ import annotations

import datetime
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DESKTOP = Path.home() / "Desktop"
OUT_PATH = DESKTOP / "MyMmorpg项目全方位详细介绍文档.docx"


def set_run_font(run, name_cn="微软雅黑", name_en="Calibri", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(8)

    for style_name, size, color in [
        ("Heading 1", 18, RGBColor(0x1F, 0x4E, 0x79)),
        ("Heading 2", 14, RGBColor(0x2E, 0x75, 0xB6)),
        ("Heading 3", 12, RGBColor(0x2E, 0x75, 0xB6)),
    ]:
        st = doc.styles[style_name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        st.paragraph_format.space_before = Pt(14)
        st.paragraph_format.space_after = Pt(8)

    if "Caption Soft" not in [s.name for s in doc.styles]:
        cap = doc.styles.add_style("Caption Soft", WD_STYLE_TYPE.PARAGRAPH)
        cap.font.size = Pt(10)
        cap.font.italic = True
        cap.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        cap.font.name = "Calibri"
        cap._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_page(doc: Document):
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.8)
        section.right_margin = Cm(2.5)
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)


def add_page_number(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=RGBColor(0x66, 0x66, 0x66))

    for instr_text in (" PAGE ", None, " NUMPAGES "):
        if instr_text is None:
            run2 = paragraph.add_run(" 页 / 共 ")
            set_run_font(run2, size=9, color=RGBColor(0x66, 0x66, 0x66))
            continue
        fld1 = OxmlElement("w:fldChar")
        fld1.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = instr_text
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        r = paragraph.add_run()
        r._r.append(fld1)
        r._r.append(instr)
        r._r.append(fld2)
        set_run_font(r, size=9, color=RGBColor(0x66, 0x66, 0x66))

    run3 = paragraph.add_run(" 页")
    set_run_font(run3, size=9, color=RGBColor(0x66, 0x66, 0x66))


def setup_header_footer(doc: Document):
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("MyMmorpg 项目全方位详细介绍文档")
    set_run_font(hr, size=9, color=RGBColor(0x88, 0x88, 0x88))

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)


def enable_update_fields_on_open(doc: Document):
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)


def add_toc(doc: Document):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    hint = OxmlElement("w:t")
    hint.text = "（请用 Word 打开后按 Ctrl+A，再按 F9 更新域，即可生成带页码的完整目录。）"
    run._r.append(hint)
    run._r.append(fld_end)


def add_para(doc, text, *, first_line_indent=True, bold=False, size=12):
    p = doc.add_paragraph()
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=12)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run, size=12)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "2E75B6")
        shading.set(qn("w:val"), "clear")
        hdr[i]._tc.get_or_add_tcPr().append(shading)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def page_break(doc):
    doc.add_page_break()


def build():
    doc = Document()
    configure_styles(doc)
    set_page(doc)
    setup_header_footer(doc)
    enable_update_fields_on_open(doc)
    today = datetime.date.today().strftime("%Y年%m月%d日")

    # ========== 封面 ==========
    for _ in range(2):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MyMmorpg")
    set_run_font(r, size=40, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("大型多人在线角色扮演游戏（MMORPG）\n服务端项目全方位详细介绍文档")
    set_run_font(r, size=18, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = info.add_run(
        f"通俗可读 · 覆盖架构 / 玩法 / 运维 / 安全 / 测试\n"
        f"技术栈：Java 17 · Spring Boot 3.2 · Spring Cloud Alibaba · Netty · MySQL · Redis\n"
        f"文档生成日期：{today}\n"
        f"说明：基于仓库当前实现能力整理，已实现 / 演示 / 计划中如实区分"
    )
    set_run_font(r, size=12, color=RGBColor(0x55, 0x55, 0x55))

    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tip.add_run(
        "\n【重要】用 Microsoft Word 或兼容软件打开后：Ctrl+A 全选 → 按 F9 更新域，"
        "即可看到带页码的自动目录；页脚为「第 X 页 / 共 Y 页」。"
    )
    set_run_font(r, size=10, color=RGBColor(0x88, 0x88, 0x88))

    page_break(doc)

    # ========== 目录 ==========
    doc.add_heading("目录", level=1)
    add_toc(doc)
    note = doc.add_paragraph()
    run = note.add_run("说明：目录由 Word 根据标题自动生成并带页码；若仍显示提示文字，请执行「更新域」。")
    set_run_font(run, size=10, color=RGBColor(0x66, 0x66, 0x66))
    page_break(doc)

    # ========== 第1章 ==========
    doc.add_heading("第一章 写给所有人的开场白", level=1)

    doc.add_heading("1.1 这份文档是写给谁的？", level=2)
    add_para(
        doc,
        "这份文档面向所有需要理解 MyMmorpg 的人：没有写过代码的产品/策划/运营、刚接触游戏服务端的学生、"
        "要评估工程完整度的管理者，以及需要快速建立整图的开发与测试同事。"
        "我们尽量用生活化比喻解释专业概念；遇到英文术语会同步给中文含义；文末有名词小词典。"
    )
    add_para(
        doc,
        "你不必一口气读完。建议：普通人读第 1～3 章与「业务故事」章；开发读架构与子系统；运维读部署、安全与排障；"
        "策划读玩法与配置导入；测试读质量保障与成熟度对照。"
    )

    doc.add_heading("1.2 一句话讲清这个项目", level=2)
    add_para(
        doc,
        "MyMmorpg 是一套用 Java 编写的网络游戏服务器工程。"
        "它在玩家客户端背后处理登录、走路、打架、聊天、背包、任务、活动、抽卡、充值等一切需要「裁判」和「记账」的事情。"
        "客户端像前台门面与遥控器；服务器才是真正的裁判 + 账本 + 仓库——谁赢了、道具有没有发出去、钱有没有到账，都由服务器说了算。"
    )

    doc.add_heading("1.3 为什么叫 MMORPG？", level=2)
    add_para(
        doc,
        "MMORPG 是 Massively Multiplayer Online Role-Playing Game 的缩写，中文常称「大型多人在线角色扮演游戏」。"
        "特征是：很多人同时在线、各自扮演角色、在同一（或分线的）世界中探索、社交、战斗与成长。"
        "本项目对标的是这类游戏的服务端能力底座，而不是只做单机或小房间对战。"
    )

    doc.add_heading("1.4 项目当前定位", level=2)
    add_bullets(
        doc,
        [
            "可运行、可测试、可部署的服务端示例 / 教学与工程实践项目。",
            "支持「单体模式」（一个进程打包大部分能力，适合本地）与「微服务模式」（按域拆分，适合扩展）。",
            "大量能力已按生产思路落地（安全校验、幂等发奖、热更新、观测指标、对账等）；部分数值与结算仍为演示级。",
            "文档如实区分「已实现 / 演示 / 计划中」，避免把愿景当成已完成功能。",
            "对齐二游常见体验：元素反应、体力、装备随机词条、战令/月卡、签到、联机房间、智能 NPC 与 Boss 行为树等。",
        ],
    )

    # ========== 第2章 ==========
    doc.add_heading("第二章 用生活比喻理解整套系统", level=1)

    doc.add_heading("2.1 把游戏服务器想成一座城", level=2)
    add_bullets(
        doc,
        [
            "城门口（网关 / 玩家服务入口）：检查身份，把你领到该去的地方。",
            "户籍处（玩家与账号）：等级、经验、封禁、会话。",
            "街道与广场（场景服务）：走路、附近能看见谁、换地图怎么交接。",
            "竞技场（战斗服务）：算伤害、判胜负；Boss 用行为树而不是临场编故事。",
            "仓库（背包服务）：道具进出记账，防复制；还有体力（Resin）与装备词条。",
            "邮局与会客厅（大厅）：邮件、好友、队伍、助战、家园。",
            "任务公告栏（任务服务）：接交任务、记进度。",
            "活动中心（活动服务）：节日活动、世界 BOSS、签到、充值进度。",
            "商城收银台（商店）：下单、验签、对账、战令/月卡。",
            "扭蛋机房（抽卡服务）：保底、十连、概率公示与审计。",
            "市政厅（管理后台）：配置导入、投诉、热更、AI 草稿。",
            "广播站（聊天）：频道消息与举报。",
            "技能学院（技能）：技能配置与释放相关能力。",
            "匹配大厅（匹配）：找队友/对手，含跨服副本匹配。",
            "版本办公室（更新）：资源补丁、差分、CDN 预热。",
        ],
    )

    doc.add_heading("2.2 为什么要能拆成多个服务？", level=2)
    add_para(
        doc,
        "早期可以「一个人干所有事」（单体），开发快。当同时在线变多、战斗特别吃 CPU 时，"
        "就把竞技场单独扩建成更大场馆（微服务），其它部门不受影响。本项目刻意保留两种模式。"
    )

    doc.add_heading("2.3 客户端与服务器如何说话？", level=2)
    add_para(
        doc,
        "走路、放技能非常频繁，不适合每次都像打开网页那样慢请求。"
        "因此主要用「长连接」：连上后保持通道，用紧凑的二进制协议（Protobuf）收发消息。"
        "后台管理、健康检查、服务互调仍用 HTTP。可以理解为：玩家走专线电话，运维走公文系统。"
    )
    add_para(
        doc,
        "网络通道支持 Netty TCP，以及可选的 UDP/KCP；KCP 写缓冲有水位控制，不可写时丢弃推送，避免背压把进程拖垮。"
    )

    # ========== 第3章 ==========
    doc.add_heading("第三章 玩家能体验到的能力全景", level=1)
    add_para(doc, "按「玩家视角」罗列。普通人当功能清单；技术人员可对照模块与接口。")

    doc.add_heading("3.1 账号、角色与在线", level=2)
    add_bullets(
        doc,
        [
            "登录鉴权与会话；重复登录踢下线（KickPlayer），避免一人分身乱数据。",
            "在线态可读（如 Redis player:online:{id}），好友系统可展示谁在线。",
            "跨服好友可用 Redis SET（cross:friend:{id}）。",
            "全局唯一 ID（Snowflake / GlobalUID），便于合服与跨服。",
        ],
    )

    doc.add_heading("3.2 大世界与场景", level=2)
    add_bullets(
        doc,
        [
            "场景分线：人太多自动开线，像地铁加开车厢。",
            "AOI（兴趣区域）：只同步附近对象，省流量。",
            "动态域网格（WorldZoneManager + ZoneLoadBalancer）：按密度拆合 Zone。",
            "场景实例池、无缝交接票据（带速度/朝向/zone）。",
            "WorldState：时间天气、机关 Bitmap、房主世界隔离。",
            "世界等级动态数值；Boss 全局刷新锁；采集物权（全服共享 / 每人独立 / 队伍共享等）。",
            "Portal 预加载与资源预加载清单；断线重连约 30 秒保护（无敌/隐身）。",
            "移动延迟补偿（SnapshotBuffer）与客户端预测校验（MOVE/DASH/技能位移差异化阈值）。",
            "【P5 产品纵深】ECA 解谜规则（RuleTriggerService）、物理状态层（PhysicsLayerService）、"
            "10+ 解谜模板（PuzzleTemplateService）、收集物分级与神瞳（CollectibleService）、"
            "区域探索度/声望（RegionProgressService）、奇观入口条件与 LOD 引导、"
            "角色绑定世界技（WorldSkillService）、野外烹饪（WorldCookingService）、"
            "区域潮汐/混沌/事件链（RegionImpactService）、公共地图标记与幻影残影、"
            "区域频道精英击杀广播、程序化洒点与 grid_cell 配置热更。",
            "【P6 二游手感纵深】移动状态机 SceneMoveCmd/MovementType + MovementAdmissionService/"
            "StaminaConsumeService；PersistentEffectZone 持续伤害；GuidanceChainService 引导链；"
            "DynamicLootTier；CHAOS CurseLayer/CleanseAnchor；元素 GaugeUnit+精通；"
            "EquipEnhanceService；Boss 贡献百分位分档与 MVP 广播。",
        ],
    )

    doc.add_heading("3.2.1 大世界探索纵深循环（P5）", level=3)
    add_bullets(
        doc,
        [
            "看见 → 想办法 → 解锁：远处 LOD 高亮奇观，前置谜题未解则拒绝进入（WonderEntryCondition）。",
            "解谜不再只是 Bitmap 开关：事件-条件-动作（ECA）可组合元素/时间/目标，触发风场、隐藏路径等。",
            "物理层同步：火焰蔓延、水面冻结等按 AOI 网格广播，客户端可渲染环境状态。",
            "收集驱动：普通/精致/珍贵宝箱 + 神瞳类收集品提升探索技能等级与体力上限。",
            "区域探索度面板：传送点、收集物、谜题、世界任务四维进度，达 40%/80% 解锁声望奖励。",
            "角色生态位：移动手段绑定特定角色/道具；厨师双倍食材、矿工额外矿物掉落；篝火即时料理 Buff。",
            "区域动态生态：清营变 SAFE 可召唤 Boss；无锚点维护则潮汐退回 CONTESTED；可进入 CHAOS 负面天气。",
            "异步社交：公共标记（求救/炫耀/BOSS 提醒）、幻影回放、区域频道 JoinBattle 独立奖励蹭怪。",
            "内容量产：ProceduralPlacementService 程序化洒点 + OpenWorldConfigPatchService 增量热更。",
        ],
    )

    doc.add_heading("3.2.2 二游手感纵深（P6）", level=3)
    add_bullets(
        doc,
        [
            "所见即所达：攀爬须校验 ClimbableMeshId 与体力；滑翔校验风场/体力阈值（IsGlidingAvailable）。",
            "体力按动作类型 + 持续时间阶梯扣除（攀爬/游泳/滑翔/摆荡），不再只是上限加成。",
            "元素改世界：火焰 apply 注册 PersistentEffectZone，每 0.5s Tick 对进入实体造成持续伤害。",
            "收集引导链：起点→仙灵→风圈→终点宝箱；AOI 广播 GuideParticle。",
            "动态掉落：宝箱 GrantPlan 按 WorldLevel 与区域探索度抬升品质权重。",
            "CHAOS 硬约束：停留超 30 秒叠加 CurseLayer（攻降 + DoT + 禁传送）；交互 CleanseAnchor 临时净化。",
        ],
    )

    doc.add_heading("3.3 战斗与成长", level=2)
    add_bullets(
        doc,
        [
            "战斗开始 / 行动 / 结算；挑战关卡联动开战。",
            "元素反应引擎：附着(GaugeUnit 弱1/强2/超强4)→触发→反应→倍率×(1+精通转化)；"
            "蒸发/融化支持残留减半；客户端预测与 rollback 标记。",
            "AI 队友指挥 + Boss 行为树（狂暴/转阶段/锁仇恨）；实时战斗禁止依赖 LLM。",
            "抽卡：保底/十连/天井/历史；概率公示与审计摘要。",
            "肉鸽：结算回写；局内天赋加点。",
            "战斗日/周统计（Redis Hash）；排行榜（等级/战力 ZSET）。",
            "皮肤衣柜与穿戴（协议号段约 12xx）。",
        ],
    )

    doc.add_heading("3.4 社交与轻社交", level=2)
    add_bullets(
        doc,
        [
            "好友、邮件（附件幂等进背包）、聊天与举报。",
            "队伍：邀请码、踢人、队长转移。",
            "联机房间 CoopRoom：最多 4 人共斗（弱化全服 AOI 伤害广播）。",
            "事件临时小队；好友助战（冷却/报酬/统计）；家园访问与装饰同步。",
            "跨服 4 人副本匹配（match type=3）。",
        ],
    )

    doc.add_heading("3.5 任务、活动、签到与商城", level=2)
    add_bullets(
        doc,
        [
            "任务 JSON 驱动，支持 Admin/Internal 导入与 reload。",
            "活动：条件、阶段、商店货品、展示文案、导入校验；可订阅战斗结束与订单支付。",
            "世界 BOSS / 动态事件：排期、伤害与治疗/护盾辅助折算、按贡献百分位（前10%/30%/50%）分档结算 grantPlans，并广播 MVP。",
            "七日签到 / 月卡日领。",
            "战令 / 月卡 / 首充双倍（含自动续费、补偿领取、跨平台同步 stub）。",
            "商城订单双写、支付验签、账本、对账；渠道适配 Apple/Google/支付宝/微信 stub + MOCK。",
            "体力 Resin：自然恢复与每日购买次数。",
            "装备随机词条写入 affix_blob；EquipEnhanceService 支持强化升级（每 4 级副词条成长，服务端种子 + 分布式锁）。",
        ],
    )

    doc.add_heading("3.6 匹配与客户端更新", level=2)
    add_bullets(
        doc,
        [
            "匹配互补评分（等级/战力/等待时间）；队列分片 + 跨 shard 配对。",
            "热更差分、资源补丁链、CDN 预热。",
        ],
    )

    doc.add_heading("3.7 AI 辅助面", level=2)
    add_para(
        doc,
        "AI 是「参谋」不是「裁判」。玩家顾问默认 rule-coach，LLM 关闭时回退规则；"
        "智能 NPC 可带短期记忆/情绪，超时回退规则；Admin 可生成活动/商城/任务草案与投诉分类、战斗周报。"
        "真正改数值、发道具仍走正式业务链路。"
    )

    # ========== 第4章 ==========
    doc.add_heading("第四章 技术全景", level=1)

    doc.add_heading("4.1 技术选型", level=2)
    add_table(
        doc,
        ["技术", "通俗理解", "作用"],
        [
            ["Java 17", "主流服务端语言版本", "全部业务代码基础"],
            ["Spring Boot 3.2", "应用脚手架", "各服务启动与配置"],
            ["Spring Cloud Alibaba", "微服务全家桶（可选）", "发现、远程调用等"],
            ["Netty", "高性能网络引擎", "TCP / UDP / KCP 长连接"],
            ["Protobuf", "紧凑消息格式", "游戏协议"],
            ["MySQL 8", "可靠账本", "账号、邮件、订单、进度等"],
            ["Redis", "超快便签与计分板", "在线、匹配、排行、票据、热数据"],
            ["RocketMQ（可关）", "邮局", "战斗结束、订单支付等事件"],
            ["OpenFeign", "服务间打电话", "微服务命令转发"],
            ["Docker Compose", "一键依赖环境", "MySQL/Redis/Nacos"],
            ["TestNG / CI / JaCoCo / SpotBugs", "自动考试", "质量门禁"],
            ["Prometheus", "仪表盘指标", "调度/MQ/商城等计数"],
        ],
        col_widths=[4.0, 4.5, 6.7],
    )

    doc.add_heading("4.2 仓库模块一览", level=2)
    add_table(
        doc,
        ["模块", "角色说明"],
        [
            ["mmorpg-protocol", "协议词典：*.proto、MessageId、RetCode"],
            ["mmorpg-domain-api", "跨服务 Port 合同（如背包命令/发货）"],
            ["mmorpg-common", "公共工具箱：网络、RPC、远程 Port 等"],
            ["mmorpg-gateway", "可选总门岗：区域、功能号段、粘滞"],
            ["player-service", "主入口；单体时嵌入多域"],
            ["scene-service", "场景 / AOI / 大世界"],
            ["battle-service", "战斗与 AI 队友 / Boss BT"],
            ["activity-service", "活动 / 世界事件 / 签到"],
            ["bag-service", "背包 / 体力 / 词条"],
            ["shop-service", "商城 / 支付 / 战令"],
            ["gacha-service", "抽卡（8994，可嵌入）"],
            ["hall-service", "好友邮件队伍助战家园"],
            ["quest-service", "任务"],
            ["matchmaking-service", "匹配（含跨服副本）"],
            ["chat-service", "聊天与举报"],
            ["skill-service", "技能"],
            ["admin-service", "运营后台与热更"],
            ["update-service", "版本与资源更新"],
            ["mmorpg-cli", "命令行导入工具"],
        ],
        col_widths=[4.5, 10.7],
    )

    doc.add_heading("4.3 两种部署模式", level=2)
    doc.add_heading("4.3.1 单体（默认）", level=3)
    add_para(
        doc,
        "只启动 player-service（HTTP 8989），多数域以 Maven 依赖「嵌」在同一进程。适合本地开发、集成测试、小规模演示。"
    )
    add_numbered(
        doc,
        [
            "docker compose up -d mysql redis",
            "复制 .env.example 为 .env",
            "mvn -pl player-service -am spring-boot:run -Dspring-boot.run.profiles=dev",
            "客户端连玩家服 HTTP / 游戏端口",
        ],
    )

    doc.add_heading("4.3.2 微服务", level=3)
    add_para(
        doc,
        "各域独立 JAR；player 作网关 Feign 转发；域服务通过 RestTemplate 远程 Port 回调。"
        "需统一 INTERNAL_API_SECRET；独立 activity 须开远程 Port 并配置 BAG_SERVICE_URL，否则发奖 NoOp，生产 fail-fast。"
    )
    add_table(
        doc,
        ["服务", "端口", "职责"],
        [
            ["player-service", "8989", "客户端连接与玩家主入口"],
            ["scene-service", "8981", "场景与大世界"],
            ["chat-service", "8982", "聊天"],
            ["bag-service", "8983", "背包"],
            ["skill-service", "8984", "技能"],
            ["admin-service", "8985", "后台"],
            ["hall-service", "8986", "大厅社交"],
            ["quest-service", "8987", "任务"],
            ["matchmaking-service", "8988", "匹配"],
            ["shop-service", "8990", "商城支付"],
            ["battle-service", "8991", "战斗"],
            ["activity-service", "8992", "活动"],
            ["update-service", "8993", "更新"],
            ["gacha-service", "8994", "抽卡"],
            ["mmorpg-gateway", "8443（可配）", "可选统一网关"],
        ],
        col_widths=[4.5, 3.5, 7.2],
    )

    doc.add_heading("4.4 逻辑拓扑", level=2)
    add_para(
        doc,
        "Client（TCP/KCP/WS）→ mmorpg-gateway（可选，区域/功能号段/粘滞）→ player 及各域服务；"
        "底层依赖 MySQL 多库、Redis、可选 MQ Outbox/Inbox。详细 ASCII 图见仓库 docs/troubleshooting.md。",
        first_line_indent=True,
    )

    # ========== 第5章 ==========
    doc.add_heading("第五章 核心子系统详解", level=1)

    doc.add_heading("5.1 玩家入口与 Center 路由", level=2)
    add_para(
        doc,
        "多场景节点时需要「总调度」。支持 local / remote 中心模式。"
        "跨节点传送签发一次性 session_ticket；目标节点验票进入；断线可用 ResumeScene；"
        "重复登录先踢旧连接。场景节点可向 Center 注册/心跳；支持单线人数上限与空线回收。"
    )

    doc.add_heading("5.2 场景、AOI 与大世界顶尖能力", level=2)
    add_para(
        doc,
        "AOI 用空间哈希切格子再距离精筛。WorldZoneManager 按密度管理 Zone；SceneInstancePool 管理实例。"
        "P0/P1 能力包括：WorldState、WorldLevel、BossRespawnTimer、LootOwnershipPolicy、PartyEntityOwnership、"
        "PortalPreload、ReconnectProtection、AoiBroadcastStrategy 等。业务大盘指标："
        "GET /internal/scene/open-world/metrics/business。"
    )
    add_para(
        doc,
        "P4 二游产品玩法层由 OpenWorldGameplayFacade 统一门面，覆盖探索-奖励循环、功能奇观、"
        "区域世界状态、叙事偶遇、探索技能、移动手段、环境交互、惊喜彩蛋、家园/捉宠/休闲/搜打撤等。"
        "GET /internal/scene/open-world/gameplay/status 可带 playerId、regionId 返回 region_progress。"
    )
    doc.add_heading("5.2.1 P5 解谜引擎与物理层", level=3)
    add_table(
        doc,
        ["能力", "实现类", "主要 API"],
        [
            ["ECA 规则触发器", "RuleTriggerService", "POST .../rule/fire、GET .../rule/list"],
            ["物理状态层（AOI 广播）", "PhysicsLayerService", "POST .../physics/apply、GET .../physics/aoi"],
            ["解谜库（≥10 模板）", "PuzzleTemplateService", "GET .../puzzle/templates、POST .../instantiate|advance"],
            ["环境交互联动规则", "OpenWorldGameplayFacade.envInteractWithRules", "POST .../env/interact（自动 fire 规则+物理）"],
        ],
        col_widths=[4.0, 5.0, 5.2],
    )
    doc.add_heading("5.2.2 P5 分层探索与角色生态位", level=3)
    add_table(
        doc,
        ["能力", "实现类", "主要 API"],
        [
            ["收集物分级/神瞳", "CollectibleService", "POST .../collectible/collect、GET .../collectible/progress"],
            ["区域探索度/声望", "RegionProgressService", "GET .../region/progress"],
            ["奇观入口条件+LOD", "LandmarkWonderService", "POST .../landmark/enter、GET .../landmark/lod"],
            ["角色绑定世界技", "WorldSkillService", "POST .../world-skill/unlock-mode"],
            ["队伍生态位天赋", "EnvironmentInteractionService", "POST .../env/gather-talent"],
            ["野外篝火烹饪", "WorldCookingService", "POST .../cook"],
        ],
        col_widths=[4.0, 5.0, 5.2],
    )
    doc.add_heading("5.2.3 P5 区域潮汐、异步社交与量产管线", level=3)
    add_table(
        doc,
        ["能力", "实现类", "主要 API"],
        [
            ["区域 HOSTILE/CONTESTED/SAFE/CHAOS", "RegionImpactService", "POST .../region/clear-camp|anchor|tick-decay"],
            ["Boss 联动与事件链", "RegionImpactService", "GET .../region/boss-link、POST .../region/event-chain/advance"],
            ["公共地图标记", "PublicMarkService", "POST .../mark/place|thank"],
            ["幻影残影回放", "WorldEncounterService", "POST .../phantom/leave|play"],
            ["区域频道蹭怪", "RegionWorldChannelService", "POST .../channel/broadcast-elite|join-battle"],
            ["程序化洒点", "ProceduralPlacementService", "POST .../placement/fill"],
            ["grid_cell 配置热更", "OpenWorldConfigPatchService", "POST .../config/patch|reload"],
        ],
        col_widths=[4.5, 4.5, 4.2],
    )
    add_para(
        doc,
        "Scene 定时 tick（game.world.tick-ms）除 NPC/天气外，会调用 gameplay.tickGameplay 推进区域潮汐衰退。"
        "GM 命令 puzzle_instantiate、region_enter_chaos、config_reload_cells 已注册。"
        "完整对照与测试类见 docs/open-world-top-tier.md（P5 章节）、OpenWorldP5EnhancementFlowTest、OpenWorldP5BusinessFlowTest。"
    )

    doc.add_heading("5.2.4 P6 二游手感纵深（垂直移动 / DoT / 引导 / 养成 / Boss 贡献）", level=3)
    add_table(
        doc,
        ["能力", "实现类", "主要 API"],
        [
            ["移动状态机", "SceneMoveCmd + MovementType", "POST .../move/admit"],
            ["物理准入", "MovementAdmissionService", "攀爬 mesh / 滑翔风场校验"],
            ["体力阶梯", "StaminaConsumeService", "GET .../stamina"],
            ["持续伤害区", "ZoneLifecycleManager", "POST .../physics/zone/tick"],
            ["引导链", "GuidanceChainService", "POST .../guidance/start|advance|particles"],
            ["动态宝箱品质", "CollectibleService.DynamicLootTier", "collect 带 worldLevel/探索度"],
            ["CHAOS 诅咒/净化", "RegionImpactService", "POST .../region/affliction/tick|cleanse"],
            ["元素精通/量", "ElementReactionEngine", "resolve(..., gauge, mastery)"],
            ["装备强化", "EquipEnhanceService", "POST /internal/bag/equip/enhance"],
            ["Boss 贡献分档", "WorldEventService", "POST .../world-events/assist|settle"],
            ["MVP 广播", "RegionWorldChannelService", "POST .../channel/broadcast-mvp"],
        ],
        col_widths=[4.0, 5.0, 5.2],
    )
    add_para(
        doc,
        "完整对照与测试见 docs/open-world-top-tier.md（P6 章节）、OpenWorldP6DepthFlowTest、"
        "OpenWorldP6BusinessFlowTest、ElementReactionMasteryGaugeTest、EquipEnhanceServiceTest、"
        "WorldBossRewardTierTest、P6CrossServiceBusinessFlowTest。"
    )

    doc.add_heading("5.3 战斗域与 AI", level=2)
    add_para(
        doc,
        "战斗负责开战、行动、结算，并发布 BattleEnded 等事件供活动/任务订阅。"
        "状态常在 Redis；活跃战斗用 SET 索引，避免 KEYS 全库扫描。"
        "元素反应（含 GaugeUnit 与精通倍率、残留）、客户端预测与 rollback、"
        "Boss BT 调试接口与 ThreatTable 仇恨表均已落地。"
        "红线：实时战斗不引入 LLM。"
    )

    doc.add_heading("5.4 背包、体力、词条与发奖幂等", level=2)
    add_para(
        doc,
        "发奖强调幂等（同一凭证只成功一次）。邮件附件经 MailItemGrantPort 进背包。"
        "玩家与背包可冷热分离：Redis 热数据 + 定时刷 MySQL。"
        "体力 /internal/bag/resin*；装备随机词条 EquipRandomizer；"
        "装备强化 /internal/bag/equip/enhance（种子伪随机 + Redis/本地锁）。"
    )

    doc.add_heading("5.5 商城、支付、战令与对账", level=2)
    add_para(
        doc,
        "P0 经济硬化：订单双写 shop_order、支付事件 mq_outbox、发奖/累充幂等（grant_idempotency / mq_inbox）、"
        "流水账本（wallet_ledger / item_ledger）、对账 POST /internal/shop/orders/reconcile。"
        "开发可用 MOCK；生产禁用 MOCK 并配置渠道密钥。战令/月卡/首充走 PassService。"
    )

    doc.add_heading("5.6 活动、签到与世界事件", level=2)
    add_para(
        doc,
        "活动配置模型较完整（条件、阶段、商店货品、展示、导入校验）。"
        "可消费商店已支付与战斗结束消息。SignInService 提供签到/月卡日领。"
        "世界事件提供排期、伤害与 grantPlans 结算钩子。"
    )

    doc.add_heading("5.7 抽卡服务", level=2)
    add_para(
        doc,
        "独立 gacha-service（8994，可嵌入）。协议约 14xx；保底/十连/天井/历史；"
        "GET /internal/gacha/probability 概率公示；GET /internal/gacha/audit 审计；"
        "GAME_GACHA_REMOTE_ENABLED 切换远程。"
    )

    doc.add_heading("5.8 匹配系统", level=2)
    add_para(
        doc,
        "队列在 Redis，SET NX + Lua 降低错配；可分片与跨 shard 配对；互补评分照顾久等玩家。"
        "跨服副本匹配 type=3，场景 9100+modeId。"
    )

    doc.add_heading("5.9 管理后台与热更新", level=2)
    add_para(
        doc,
        "Admin：配置导入、投诉、操作日志、AI 草稿、分阶段 reload（activity→update→quest→player）、"
        "发布审计、配置 Diff/回滚/灰度/引用完整性、运营控制台 ops-console.html。"
        "RBAC 预设：OPS / PLANNER / CS / SUPERADMIN。鉴权：IP 白名单 + HMAC/API Key；生产强制开启。"
    )

    doc.add_heading("5.10 网关、就近接入与功能号段路由", level=2)
    add_para(
        doc,
        "Gateway 可按 X-Account-Id 等粘滞到同一 player 实例；GeoIP/区域头注入就近信息；"
        "FunctionNumberRoutingFilter 按 X-Msg-Id 把移动路由到 Scene、背包到 Bag 等，削弱中转瓶颈。"
    )

    doc.add_heading("5.11 反作弊与客户端完整性", level=2)
    add_para(
        doc,
        "AntiCheatService + AntiCheatRuleEngine + ClientIntegrityChecker（Root/模拟器等）；"
        "场景移动与伤害校验 API，累计违规可临时封禁。"
    )

    # ========== 第6章 ==========
    doc.add_heading("第六章 数据、消息与一致性", level=1)

    doc.add_heading("6.1 MySQL 与 Redis 分工", level=2)
    add_para(
        doc,
        "MySQL 像保险箱账本；Redis 像前台便签与计分板。重要结果落库，高频读写用 Redis。"
    )

    doc.add_heading("6.2 多库与迁移", level=2)
    add_para(
        doc,
        "四库：player_db / battle_db / activity_db / update_db。权威 DDL 在各服务 schema/*.sql；"
        "Compose 首次启动自动建库挂载。Flyway 已在 player/activity/shop 等启用。"
        "活动进度以 Redis JSON 为主。单体模式通常只需 player_db；微服务需四库齐全。"
    )

    doc.add_heading("6.3 消息队列与最终一致", level=2)
    add_para(
        doc,
        "开发默认关闭 RocketMQ（NoOp）；生产依赖事件链路须 ROCKETMQ_ENABLED=true。"
        "必须可靠的事件：商店已支付→活动充值进度；战斗结束→活动/任务进度（Inbox 幂等，另有 HTTP 投影兜底）。"
        "Outbox/Inbox、死信队列、MqLagMonitor 支撑可靠投递与积压观测。"
    )
    add_para(
        doc,
        "「最终一致」不是允许错账，而是允许极短时间各处不完全一样，最终通过重试、幂等与补偿达成正确状态。"
    )

    doc.add_heading("6.4 TLog 与日志治理", level=2)
    add_para(
        doc,
        "TLogEventPublisher 默认结构化日志；可选 Kafka。LogDesensitizer 脱敏；LogSamplingFilter 采样；"
        "TraceContext 跨线程 TraceId，与网关 X-Trace-Id 贯通。"
    )

    # ========== 第7章 ==========
    doc.add_heading("第七章 安全、风控与生产意识", level=1)

    doc.add_heading("7.1 内部 API HMAC", level=2)
    add_para(
        doc,
        "/internal/** 使用 HMAC-SHA256（时间戳 + 签名 + 玩家头等）。生产强制开启并拒绝弱密钥。"
        "各服务密钥必须一致；时钟偏差过大会导致验签失败，需 NTP。"
    )

    doc.add_heading("7.2 后台安全", level=2)
    add_bullets(
        doc,
        [
            "IP 白名单、API Key / HMAC、RBAC 权限点、操作日志。",
            "生产样例：ADMIN_AUTH_ENABLED=true，强随机 ADMIN_HMAC_SECRET / ADMIN_API_KEY。",
        ],
    )

    doc.add_heading("7.3 会话、TLS 与密钥", level=2)
    add_para(
        doc,
        "会话 RSA 生产必须固定密钥；开发未配置可临时生成。可配置服务/网关 SSL。"
        "Nacos 本机可关认证，共享环境务必开认证。"
    )

    doc.add_heading("7.4 支付与反作弊", level=2)
    add_para(
        doc,
        "支付验签、对账、幂等履约；反作弊降低外挂收益。不能指望客户端诚实上报结果。"
    )

    # ========== 第8章 ==========
    doc.add_heading("第八章 环境变量与配置矩阵（精要）", level=1)
    add_para(doc, "完整样例见仓库 .env.example。下表帮助理解「开发随便、生产严格」。")
    add_table(
        doc,
        ["变量/项", "dev", "prod 要点"],
        [
            ["MYSQL_PASSWORD", "可弱默认", "强密码必填"],
            ["INTERNAL_API_SECRET", "可关校验", "必填强密钥，开启校验"],
            ["ADMIN_AUTH_ENABLED", "常 false", "必须 true"],
            ["SHOP_PAYMENT_MOCK_ENABLED", "true", "必须 false"],
            ["SESSION_RSA_*", "可空临时生成", "必填固定密钥"],
            ["ROCKETMQ_ENABLED", "false", "事件链路需 true"],
            ["GAME_*_REMOTE_ENABLED", "false", "拆分时 true"],
            ["GAME_TLOG_KAFKA_ENABLED", "false", "需要时 true"],
        ],
        col_widths=[5.0, 4.5, 5.7],
    )

    # ========== 第9章 ==========
    doc.add_heading("第九章 本地运行、测试、CI 与压测", level=1)

    doc.add_heading("9.1 最小启动", level=2)
    add_numbered(
        doc,
        [
            "安装 JDK 17、Maven、Docker。",
            "docker compose up -d mysql redis",
            "复制 .env.example 为 .env",
            "启动 player-service（dev profile）",
            "用客户端或协议工具验证",
        ],
    )

    doc.add_heading("9.2 测试分层", level=2)
    add_table(
        doc,
        ["类型", "目的", "命令"],
        [
            ["单元测试", "快速验证逻辑", "mvn -B test"],
            ["集成测试", "带 MySQL/Redis 联动", "mvn -B test -Pintegration-tests"],
            ["质量门禁", "JaCoCo≥10% + SpotBugs", "mvn -B verify -Pquality-gate"],
        ],
        col_widths=[3.2, 5.5, 6.5],
    )
    add_para(
        doc,
        "各模块有 testng-unit.xml / testng-integration.xml；根目录 test-suites 有聚合套件。"
        "GitHub Actions CI：unit-tests + 带 MySQL/Redis service 的 integration-tests。"
    )

    doc.add_heading("9.3 可观测性", level=2)
    add_para(
        doc,
        "Actuator 暴露 health/info/metrics/prometheus；告警骨架在 deploy/observability/。"
        "Feign FallbackFactory 覆盖 scene/hall/battle；熔断示例见 circuitbreaker-example.yml。"
        "建议接入 ELK/Loki 以 TraceId 检索；生产可接 Zipkin/OpenTelemetry。"
    )

    doc.add_heading("9.4 压测", level=2)
    add_para(
        doc,
        "perf/ 目录含 JMeter 与 Gatling 脚本；AOI 有单进程压力接口。"
        "注意：单进程基准 ≠ 集群千人同屏已验收。"
    )

    doc.add_heading("9.5 Redis Sentinel", level=2)
    add_para(
        doc,
        "高可用可用 docker-compose.redis-sentinel.yml，profile 叠加 redis-sentinel 并配置哨兵节点。"
    )

    # ========== 第10章 ==========
    doc.add_heading("第十章 能力成熟度对照表", level=1)

    doc.add_heading("10.1 已实现（工程可用）", level=2)
    add_bullets(
        doc,
        [
            "单体/微服务双模式；内部 HMAC 与生产密钥校验；Admin 导入/热更/审计/Diff/回滚/RBAC。",
            "Center 路由、迁移票据、分线、重连保护、AOI、动态域、无缝交接、WorldState 等大世界底座。",
            "元素反应、体力、词条、战令/月卡/签到、CoopRoom、队伍、助战、家园；"
            "P6：垂直移动/DoT Zone/引导链/CHAOS 诅咒/精通反应/装备强化/Boss 贡献分档与 MVP。",
            "抽卡、挑战、肉鸽、匹配分片、排行榜、邮件附件幂等发奖。",
            "支付对账骨架、反作弊、Feign 熔断、冷热分离、TLog、CI 与质量门禁。",
            "智能 NPC / AI 顾问 / Boss BT / Admin AI 内容草稿。",
        ],
    )

    doc.add_heading("10.2 演示级", level=2)
    add_bullets(
        doc,
        [
            "部分奖励数值为演示常量；世界 BOSS grantPlans 需继续接完整履约。",
            "跨服场景已接 K8sSceneAllocator/HPA/Agones 清单与 SceneBackgroundPrecreator 秒切预创建。",
            "支付渠道为适配器骨架，官方 SDK 需替换 HTTP 占位。",
            "GeoIP 为粗解析；Admin Web 为静态控制台而非完整表编辑器。",
            "RocketMQ / TLog Kafka 开发默认关闭。",
        ],
    )

    doc.add_heading("10.3 计划中", level=2)
    add_bullets(
        doc,
        [
            "player 去除剩余域 Maven 依赖，真正 gateway-only。",
            "完整 RedLock、官方支付 SDK、KCP 更深对齐。",
            "真实 GSLB、跨区域多活、集群千人同屏压测进 CI、mTLS + Vault 轮转、可选实时语音等。",
        ],
    )

    # ========== 第11章 ==========
    doc.add_heading("第十一章 典型业务故事", level=1)

    doc.add_heading("11.1 登录并走进大世界", level=2)
    add_numbered(
        doc,
        [
            "登录拿到会话；玩家服记在线并规划场景节点/分线。",
            "移动时 AOI 只同步附近对象；过密可能换线或换 Zone。",
            "跨节点传送签发票据，目标验票；接近 Portal 可预加载。",
            "短时掉线可 Resume，并有短暂保护期。",
        ],
    )

    doc.add_heading("11.2 打完挑战，活动进度 +1", level=2)
    add_numbered(
        doc,
        [
            "开战转到战斗域；结算后发 BattleEnded。",
            "活动/任务幂等消费，更新胜利次数等。",
            "玩家领奖进入背包（幂等）。",
        ],
    )

    doc.add_heading("11.3 充值礼包到账", level=2)
    add_numbered(
        doc,
        [
            "下单落库 → 渠道回调验签 → Outbox 可靠投递。",
            "发货与累充幂等；活动充值进度更新。",
            "对账发现差异则人工/自动补单。",
        ],
    )

    doc.add_heading("11.4 策划上线新活动", level=2)
    add_numbered(
        doc,
        [
            "JSON/CSV 校验后经 Admin 或 CLI 导入。",
            "分阶段 reload；发布审计留痕。",
            "玩家侧看到新活动（视资源与配置而定）。",
        ],
    )

    doc.add_heading("11.5 十连抽卡", level=2)
    add_numbered(
        doc,
        [
            "客户端发 14xx 抽卡命令（或 HTTP 内部接口联调）。",
            "抽卡服计算保底/天井，写历史。",
            "道具经背包幂等发放；概率与审计可对外公示。",
        ],
    )

    doc.add_heading("11.6 大世界探索纵深旅程（P5）", level=2)
    add_numbered(
        doc,
        [
            "玩家靠近环境物体 POST /env/interact：服务端判定元素反应，并 fire ECA 规则（如火+藤蔓+夜晚→隐藏路径），"
            "必要时写入 physics_layer 供 AOI 广播。",
            "POST /puzzle/advance 逐步破解 JSON 配置的解谜实例；完成后 markPuzzleSolved，解锁 sealed-sanctum 等奇观入口。",
            "POST /collectible/collect 收集分级宝箱/神瞳；同步 markCollectible 计入 region_progress。",
            "GET /region/progress 达阈值解锁声望 grantPlans（翅膀、名片、食谱等）。",
            "char_venti 绑定解锁 GLIDE → 进入 floating-ruin；清营变 SAFE → Boss 可召唤 + 商队事件链启动。",
            "POST /mark/place 分享 BOSS 提醒；POST /phantom/leave 留下残影；区域频道 broadcast-elite 后其他玩家 JoinBattle 独立结算。",
            "策划 POST /placement/fill 程序化洒点；改配置后 POST /config/patch + /config/reload 只刷新指定 grid_cell。",
        ],
    )

    doc.add_heading("11.7 二游手感纵深旅程（P6）", level=2)
    add_numbered(
        doc,
        [
            "解锁 CLIMB/GLIDE 后 POST /move/admit：攀爬校验 mesh 与体力阶梯扣除；风场内滑翔减耗。",
            "POST /physics/apply kind=FIRE → 自动注册燃烧 DoT Zone；POST /physics/zone/tick 对进入实体结算伤害。",
            "POST /guidance/start → advance 仙灵节点 → particles AOI 广播 → 终点 chestCollectibleId。",
            "POST /collectible/collect 携带 worldLevel 与 regionExplorationRate，动态抬升宝箱品质。",
            "区域 enterChaos 后 /region/affliction/tick 叠 CurseLayer；POST /region/cleanse 交互净化锚点。",
            "战斗侧 resolve(..., gauge, mastery) 结算精通蒸发/残留；背包 POST /equip/enhance 每 4 级成长副词条。",
            "世界事件 damage + assist → settle 百分位发奖；POST /channel/broadcast-mvp 广播 MVP 与最高伤害。",
        ],
    )

    # ========== 第12章 ==========
    doc.add_heading("第十二章 团队角色如何协作", level=1)
    add_table(
        doc,
        ["角色", "关心什么", "主要模块"],
        [
            ["客户端", "协议、同步、表现", "protocol、player 入口"],
            ["服务端", "正确性、性能、边界", "各 *-service、common"],
            ["策划", "数值与配置", "config、Admin 导入"],
            ["运营/客服", "投诉、补单、封禁", "admin、对账"],
            ["运维/SRE", "部署、监控、密钥", "Docker、网关、Prometheus"],
            ["测试", "回归、幂等、压测", "TestNG、CI、perf"],
        ],
        col_widths=[3.0, 5.5, 6.7],
    )

    # ========== 第13章 ==========
    doc.add_heading("第十三章 故障排查入门", level=1)
    add_table(
        doc,
        ["现象", "先查什么"],
        [
            ["Redis 连不上", "compose ps、REDIS_HOST、actuator health"],
            ["支付后没发奖 / MQ 积压", "mq_outbox 状态、死信、lag 指标、MQ 开关"],
            ["内部 API 401", "各服务密钥是否一致、时钟、生产校验"],
            ["对账 mismatch", "渠道账单格式、自动对账开关、人工确认后再补单"],
            ["传送黑屏 / 票据失效", "Portal 预加载 TTL、迁移票据 Redis 键、资源预载"],
            ["scene/battle 降级", "熔断阈值、下游延迟、TraceId 拉链路"],
        ],
        col_widths=[5.5, 9.7],
    )
    add_para(doc, "更完整步骤见 docs/troubleshooting.md；OpenAPI/Postman 见 docs/openapi.md。")

    # ========== 第14章 ==========
    doc.add_heading("第十四章 协议号段与内部 API 速览", level=1)
    add_para(
        doc,
        "游戏消息用数字消息号区分玩法。常见号段印象：皮肤约 12xx、挑战约 13xx、抽卡约 14xx、肉鸽约 15xx；"
        "场景传送重定向 retcode=27；ResumeScene 约 120/121。精确以 mmorpg-protocol 为准。"
    )
    add_para(
        doc,
        "内部 HTTP 前缀 /internal/**：战斗 start/action/end、活动世界事件、签到、背包 resin、"
        "大厅 coop/party/assist/home、商城 pass/对账、抽卡 probability/audit、场景 open-world、Admin /admin/** 等。"
        "命令型可为 Protobuf 流；Port 型多为 JSON。"
    )

    # ========== 第15章 ==========
    doc.add_heading("第十五章 常见问题 FAQ", level=1)
    faqs = [
        (
            "Q1：没有编程基础能看懂吗？",
            "可以。先读第 1～3 章与业务故事，把服务当城市部门；建立直觉后再深入技术章。",
        ),
        (
            "Q2：这是已上线的商业游戏吗？",
            "是完整度较高的服务端工程底座与示例。部分数值与结算仍是演示级；商业上线还需内容、合规、运维与更严压测。",
        ),
        (
            "Q3：单体还是微服务？",
            "学习与本机演示优先单体；需要独立扩容与多人分工再开远程开关。",
        ),
        (
            "Q4：没有客户端怎么办？",
            "可起服务看健康检查与日志；完整玩法通常需客户端或协议调试工具。仓库侧重服务端。",
        ),
        (
            "Q5：目录和页码不显示？",
            "用 Word 打开，Ctrl+A 后按 F9，或右键目录「更新域」。WPS 一般也支持。",
        ),
        (
            "Q6：数据会丢吗？",
            "持久数据在 MySQL，热点在 Redis；发奖与支付强调幂等与账本。仍需正规备份，示例不能替代完整容灾。",
        ),
        (
            "Q7：如何导入活动配置？",
            "可用 Admin 接口或 mmorpg-cli：import activity / import manifest，配合 API Key 或 HMAC。",
        ),
    ]
    for q, a in faqs:
        doc.add_heading(q, level=2)
        add_para(doc, a)

    # ========== 第16章 ==========
    doc.add_heading("第十六章 名词小词典", level=1)
    add_table(
        doc,
        ["名词", "一句话解释"],
        [
            ["MMORPG", "大型多人在线角色扮演游戏"],
            ["客户端 / 服务端", "玩家程序 / 机房裁判与账本"],
            ["单体 / 微服务", "合署办公 / 分局独立部署"],
            ["网关", "统一入口与路由"],
            ["Protobuf / 长连接", "紧凑消息格式 / 保持通道高频交互"],
            ["AOI / 分线", "只同步附近对象 / 平行地图实例"],
            ["幂等", "同一请求多次执行不重复发奖"],
            ["HMAC", "带密钥的防篡改签名"],
            ["Outbox / Inbox", "可靠发信与防重复收信"],
            ["Feign / 熔断", "服务间调用 / 下游挂了快速失败"],
            ["热更新 / 对账", "尽量不停服刷新配置 / 账单核对"],
            ["CI / Prometheus", "自动构建测试 / 指标监控"],
            ["Resin / 战令", "体力 / 通行证成长付费线"],
            ["行为树 BT", "Boss AI 的规则流程图，非 LLM"],
            ["CoopRoom", "最多 4 人联机房间"],
        ],
        col_widths=[4.0, 11.2],
    )

    # ========== 第17章 ==========
    doc.add_heading("第十七章 仓库文档地图", level=1)
    add_table(
        doc,
        ["文件", "内容"],
        [
            ["README.md", "快速开始与模块一览"],
            ["DEPLOYMENT.md", "部署、环境变量、能力分层"],
            ["TESTING.md", "TestNG 套件说明"],
            [".env.example", "全量环境变量样例"],
            ["docs/troubleshooting.md", "拓扑与排障"],
            ["docs/openapi.md", "OpenAPI / 协议 / Postman"],
            ["docs/open-world-top-tier.md", "大世界顶尖标准（含 P0～P6）"],
            ["docs/service-boundaries.md", "服务边界与契约"],
            ["docs/architecture-gateway.md", "网关化拆分路线"],
            ["docs/db-migration.md", "数据库拆分迁移"],
            ["docs/resilience-observability.md", "韧性与可观测"],
            ["perf/README.md", "压测脚本说明"],
        ],
        col_widths=[5.5, 9.7],
    )

    # ========== 第18章 ==========
    doc.add_heading("第十八章 总结", level=1)
    add_para(
        doc,
        "MyMmorpg 用现代 Java 技术栈，搭出一套既适合学习演示、又具备微服务扩展潜力的 MMORPG 服务端。"
        "它覆盖玩家入口、大世界（含 P5 解谜/探索纵深/异步社交与 P6 垂直移动/DoT/引导链/精通反应/装备强化/Boss 贡献）、战斗、社交、任务、活动、抽卡、商城支付、后台运营、安全与工程化测试等板块，"
        "并吸收二游常见体验（元素反应与精通附着量、攀爬/滑翔体力、装备强化、战令、签到、联机房间、"
        "行为树 Boss、神瞳式收集与引导链、区域声望、野外烹饪、世界 Boss 贡献分档等）。"
    )
    add_para(
        doc,
        "请记住三句话：第一，客户端负责表现，服务器负责裁决与记账；"
        "第二，可以先合署办公（单体），再分局扩容（微服务）；"
        "第三，发奖与支付必须以幂等、验签、对账为底线——这是信任的基础。"
    )
    add_para(
        doc,
        "学生可用本文作课程设计说明骨架；转行者可作读源码导航；管理者可据此评估完整度与后续投入重点。"
    )

    # 附录
    doc.add_heading("附录 A 服务端口速查", level=1)
    add_para(
        doc,
        "player 8989；scene 8981；chat 8982；bag 8983；skill 8984；admin 8985；hall 8986；"
        "quest 8987；matchmaking 8988；shop 8990；battle 8991；activity 8992；update 8993；"
        "gacha 8994；gateway 8443（可配置）。",
    )

    doc.add_heading("附录 B 微服务远程开关速查", level=1)
    add_para(
        doc,
        "GAME_SCENE/CHAT/BAG/SKILL/HALL/QUEST/MATCH/BATTLE/ACTIVITY/UPDATE/SHOP/GACHA_REMOTE_ENABLED=true；"
        "域服务 GAME_PORT_REMOTE_ENABLED=true 与 PLAYER_SERVICE_URL 等基址。详见 DEPLOYMENT.md。",
    )

    doc.add_heading("附录 C 文档信息", level=1)
    add_bullets(
        doc,
        [
            f"生成日期：{today}",
            "语言：简体中文",
            "目标：面向所有人的全方位介绍，兼顾技术准确性",
            "排版：自动目录域 + 页脚「第 X 页 / 共 Y 页」",
            f"输出路径：{OUT_PATH}",
        ],
    )

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_PATH))
    print(f"OK: {OUT_PATH}")
    return OUT_PATH


if __name__ == "__main__":
    build()
