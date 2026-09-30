# -*- coding: utf-8 -*-
"""原地修补桌面《MyMmorpg项目全方位详细介绍文档.docx》，写入 P6 新功能（不另存新文件）。"""

from __future__ import annotations

import datetime
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

DESKTOP = Path.home() / "Desktop"
DOC_PATH = DESKTOP / "MyMmorpg项目全方位详细介绍文档.docx"


def set_run_font(run, name_cn="微软雅黑", name_en="Calibri", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    run.font.name = name_en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name_en)
    rFonts.set(qn("w:hAnsi"), name_en)
    rFonts.set(qn("w:eastAsia"), name_cn)


def insert_paragraph_after(anchor: Paragraph, text: str = "", style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    para = Paragraph(new_p, anchor._parent)
    if style:
        para.style = style
    if text:
        run = para.add_run(text)
        if style and str(style).startswith("Heading"):
            set_run_font(run, size=14 if style == "Heading 2" else 12, bold=True, color=RGBColor(0x2E, 0x75, 0xB6))
        else:
            set_run_font(run, size=12)
    return para


def insert_runs_paragraph_after(anchor: Paragraph, text: str, style: str | None = None, *, size=12, bold=False) -> Paragraph:
    para = insert_paragraph_after(anchor, "", style=style)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return para


def insert_table_after(doc: Document, anchor: Paragraph, headers: list[str], rows: list[list[str]], col_widths=None) -> Paragraph:
    """在 anchor 后插入表格，返回表格后的空段落（便于继续插）。"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "2E75B6")
        shading.set(qn("w:val"), "clear")
        hdr[i]._tc.get_or_add_tcPr().append(shading)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)

    tbl = table._tbl
    # 表格默认加在文档末尾，挪到锚点后
    anchor._p.addnext(tbl)

    after = OxmlElement("w:p")
    tbl.addnext(after)
    return Paragraph(after, anchor._parent)


def find_para(doc: Document, contains: str) -> Paragraph | None:
    for p in doc.paragraphs:
        if contains in (p.text or ""):
            return p
    return None


def find_exact(doc: Document, text: str) -> Paragraph | None:
    for p in doc.paragraphs:
        if (p.text or "").strip() == text:
            return p
    return None


def replace_in_paragraph(p: Paragraph, old: str, new: str) -> bool:
    full = p.text or ""
    if old not in full:
        return False
    # 简单策略：清空 runs 后写入整段
    new_text = full.replace(old, new)
    for r in list(p.runs):
        r.text = ""
    if p.runs:
        p.runs[0].text = new_text
        set_run_font(p.runs[0], size=12)
    else:
        run = p.add_run(new_text)
        set_run_font(run, size=12)
    return True


def already_patched(doc: Document) -> bool:
    for p in doc.paragraphs:
        t = p.text or ""
        if "P6（二游手感纵深）" in t or "3.2.2 二游手感纵深（P6）" in t or "5.2.4 P6" in t:
            return True
    return False


def patch(doc: Document) -> list[str]:
    log: list[str] = []

    # --- 3.2 大世界列表：追加 P6 要点 ---
    p5_cap = find_para(doc, "【P5 产品纵深】")
    if p5_cap and "【P6 二游手感纵深】" not in (p5_cap.text or ""):
        bullet = insert_paragraph_after(
            p5_cap,
            "【P6 二游手感纵深】移动状态机 SceneMoveCmd/MovementType（WALK/SWIM/CLIMB/GLIDE/SWING）+ "
            "MovementAdmissionService 物理准入与 StaminaConsumeService 阶梯体力；"
            "PersistentEffectZone/ZoneLifecycleManager 持续伤害区（火烧草地 DoT）；"
            "GuidanceChainService 仙灵引导链；CollectibleService.DynamicLootTier 动态宝箱品质；"
            "RegionImpactService CurseLayer + CleanseAnchor（CHAOS 硬约束）；"
            "ElementReactionEngine GaugeUnit/精通倍率与残留；EquipEnhanceService 装备强化；"
            "WorldEventService 贡献百分位分档 + RegionWorldChannelService MVP 广播。",
            style="List Bullet",
        )
        log.append("3.2 追加 P6 能力 bullet")

    # --- 3.2.2 新小节（插在 3.3 之前）---
    h321 = find_exact(doc, "3.2.1 大世界探索纵深循环（P5）")
    if h321 and not find_exact(doc, "3.2.2 二游手感纵深（P6）"):
        # 找到 3.3 前一段（P5 最后 bullet）
        anchor = None
        paras = list(doc.paragraphs)
        for i, p in enumerate(paras):
            if (p.text or "").strip() == "3.3 战斗与成长" and i > 0:
                anchor = paras[i - 1]
                break
        if anchor is None:
            anchor = find_para(doc, "ProceduralPlacementService 程序化洒点")
        if anchor:
            cur = insert_paragraph_after(anchor, "3.2.2 二游手感纵深（P6）", style="Heading 3")
            items = [
                "所见即所达：攀爬须校验 ClimbableMeshId 与体力；滑翔校验风场/体力阈值（IsGlidingAvailable）。",
                "体力按动作类型 + 持续时间阶梯扣除（攀爬/游泳/滑翔/摆荡），不再只是上限加成。",
                "元素改世界：火焰 apply 注册 PersistentEffectZone，每 0.5s Tick 对进入实体造成持续伤害。",
                "收集引导链：起点→仙灵→风圈→终点宝箱；AOI 广播 GuideParticle。",
                "动态掉落：宝箱 GrantPlan 按 WorldLevel 与区域探索度抬升品质权重。",
                "CHAOS 硬约束：停留超 30 秒叠加 CurseLayer（攻降 + DoT + 禁传送）；交互 CleanseAnchor 临时净化。",
            ]
            for it in items:
                cur = insert_paragraph_after(cur, it, style="List Bullet")
            log.append("新增 3.2.2 P6 小节")

    # --- 3.3 元素反应 ---
    for p in doc.paragraphs:
        if "元素反应引擎：附着→触发→反应→倍率" in (p.text or "") and "精通" not in (p.text or ""):
            replace_in_paragraph(
                p,
                "元素反应引擎：附着→触发→反应→倍率；支持客户端预测与 rollback 标记。",
                "元素反应引擎：附着(GaugeUnit 弱1/强2/超强4)→触发→反应→倍率×(1+精通转化)；"
                "蒸发/融化支持残留减半；客户端预测与 rollback 标记。",
            )
            log.append("更新 3.3 元素反应描述")
            break

    # --- 3.5 世界 Boss / 装备 ---
    for p in doc.paragraphs:
        t = p.text or ""
        if "世界 BOSS / 动态事件：排期、伤害、结算 grantPlans" in t and "百分位" not in t:
            replace_in_paragraph(
                p,
                t.strip(),
                "世界 BOSS / 动态事件：排期、伤害与治疗/护盾辅助折算、按贡献百分位（前10%/30%/50%）分档结算 grantPlans，并广播 MVP。",
            )
            log.append("更新 3.5 世界 Boss 描述")
        if t.strip() == "装备随机词条写入 affix_blob。" or (
            "装备随机词条写入 affix_blob" in t and "强化" not in t
        ):
            replace_in_paragraph(
                p,
                t.strip(),
                "装备随机词条写入 affix_blob；EquipEnhanceService 支持强化升级（每 4 级副词条成长，服务端种子 + 分布式锁）。",
            )
            log.append("更新 3.5 装备强化描述")

    # --- 5.2.4 P6 表格（插在 5.3 之前）---
    if not find_para(doc, "5.2.4 P6 二游手感纵深"):
        anchor = None
        paras = list(doc.paragraphs)
        for i, p in enumerate(paras):
            if (p.text or "").strip() == "5.3 战斗域与 AI" and i > 0:
                anchor = paras[i - 1]
                break
        if anchor is None:
            anchor = find_para(doc, "OpenWorldP5BusinessFlowTest")
        if anchor:
            cur = insert_paragraph_after(anchor, "5.2.4 P6 二游手感纵深（垂直移动 / DoT / 引导 / 养成 / Boss 贡献）", style="Heading 3")
            rows = [
                ["移动状态机", "SceneMoveCmd + MovementType", "POST .../move/admit"],
                ["物理准入", "MovementAdmissionService", "攀爬 mesh / 滑翔风场校验"],
                ["体力阶梯", "StaminaConsumeService", "GET .../stamina"],
                ["持续伤害区", "ZoneLifecycleManager", "POST .../physics/zone/tick"],
                ["引导链", "GuidanceChainService", "POST .../guidance/start|advance|particles"],
                ["动态宝箱品质", "CollectibleService.DynamicLootTier", "collect 带 worldLevel/探索度"],
                ["CHAOS 诅咒/净化", "RegionImpactService", "POST .../region/affliction/tick|cleanse"],
                ["元素精通/量", "ElementReactionEngine", "resolve(..., gauge, mastery)"],
                ["装备强化", "EquipEnhanceService", "POST /internal/bag/equip/enhance"],
                ["Boss 贡献分档", "WorldEventService", "POST .../world-events/assist|settle"],
                ["MVP 广播", "RegionWorldChannelService", "POST .../channel/broadcast-mvp"],
            ]
            after_tbl = insert_table_after(
                doc,
                cur,
                ["能力", "实现类", "主要 API"],
                rows,
                col_widths=[4.0, 5.0, 5.2],
            )
            insert_runs_paragraph_after(
                after_tbl,
                "完整对照与测试见 docs/open-world-top-tier.md（P6 章节）、OpenWorldP6DepthFlowTest、"
                "OpenWorldP6BusinessFlowTest、ElementReactionMasteryGaugeTest、EquipEnhanceServiceTest、"
                "WorldBossRewardTierTest、P6CrossServiceBusinessFlowTest。",
                size=12,
            )
            log.append("新增 5.2.4 P6 表格")

    # --- 5.3 / 5.4 段落补丁 ---
    for p in doc.paragraphs:
        t = p.text or ""
        if "元素反应、客户端预测与 rollback、Boss BT 调试接口" in t and "精通" not in t:
            replace_in_paragraph(
                p,
                "元素反应、客户端预测与 rollback、Boss BT 调试接口与 ThreatTable 仇恨表均已落地。",
                "元素反应（含 GaugeUnit 与精通倍率、残留）、客户端预测与 rollback、"
                "Boss BT 调试接口与 ThreatTable 仇恨表均已落地。",
            )
            log.append("更新 5.3 战斗域描述")
        if "体力 /internal/bag/resin*；装备随机词条 EquipRandomizer" in t and "Enhance" not in t:
            replace_in_paragraph(
                p,
                "体力 /internal/bag/resin*；装备随机词条 EquipRandomizer。",
                "体力 /internal/bag/resin*；装备随机词条 EquipRandomizer；"
                "装备强化 /internal/bag/equip/enhance（种子伪随机 + Redis/本地锁）。",
            )
            log.append("更新 5.4 背包描述")

    # --- 5.6 活动若提到世界事件 ---
    p56 = find_para(doc, "5.6 活动、签到与世界事件")
    # 找 5.6 后第一段
    if p56:
        paras = list(doc.paragraphs)
        for i, p in enumerate(paras):
            if p is p56 and i + 1 < len(paras):
                nxt = paras[i + 1]
                if "百分位" not in (nxt.text or "") and "世界" in (nxt.text or ""):
                    # append note via new para after nxt
                    insert_runs_paragraph_after(
                        nxt,
                        "P6：世界 Boss 结算按有效贡献（伤害+辅助折算）百分位分档（S/A/B/C），"
                        "settle 返回 mvp，并可经区域频道 broadcast-mvp 炫耀最高伤害。",
                        size=12,
                    )
                    log.append("补充 5.6 Boss 贡献说明")
                break

    # --- 11.7 业务流程 ---
    if not find_exact(doc, "11.7 二游手感纵深旅程（P6）"):
        anchor = None
        paras = list(doc.paragraphs)
        for i, p in enumerate(paras):
            if (p.text or "").strip() == "第十二章 团队角色如何协作" and i > 0:
                anchor = paras[i - 1]
                break
        if anchor is None:
            # 找 11.6 最后一步
            for p in reversed(paras):
                if "config/patch" in (p.text or "") or "placement/fill" in (p.text or ""):
                    anchor = p
                    break
        if anchor:
            cur = insert_paragraph_after(anchor, "11.7 二游手感纵深旅程（P6）", style="Heading 2")
            steps = [
                "解锁 CLIMB/GLIDE 后 POST /move/admit：攀爬校验 mesh 与体力阶梯扣除；风场内滑翔减耗。",
                "POST /physics/apply kind=FIRE → 自动注册燃烧 DoT Zone；POST /physics/zone/tick 对进入实体结算伤害。",
                "POST /guidance/start → advance 仙灵节点 → particles AOI 广播 → 终点 chestCollectibleId。",
                "POST /collectible/collect 携带 worldLevel 与 regionExplorationRate，动态抬升宝箱品质。",
                "区域 enterChaos 后 /region/affliction/tick 叠 CurseLayer；POST /region/cleanse 交互净化锚点。",
                "战斗侧 resolve(..., gauge, mastery) 结算精通蒸发/残留；背包 POST /equip/enhance 每 4 级成长副词条。",
                "世界事件 damage + assist → settle 百分位发奖；POST /channel/broadcast-mvp 广播 MVP 与最高伤害。",
            ]
            for s in steps:
                cur = insert_paragraph_after(cur, s, style="List Number")
            log.append("新增 11.7 P6 业务流程")

    # --- 第十章已实现列表 ---
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("元素反应、体力、词条、战令") and "P6" not in t:
            replace_in_paragraph(
                p,
                t.strip(),
                t.strip().rstrip("。")
                + "；P6：垂直移动/DoT Zone/引导链/CHAOS 诅咒/精通反应/装备强化/Boss 贡献分档与 MVP。",
            )
            log.append("更新第十章已实现列表")
            break

    # --- 文档地图 ---
    for p in doc.paragraphs:
        t = p.text or ""
        if "open-world-top-tier.md" in t and "P0～P5" in t:
            replace_in_paragraph(p, "含 P0～P5", "含 P0～P6")
            log.append("更新文档地图 P0～P6")
        if "大世界顶尖标准（含 P0～P5）" in t:
            replace_in_paragraph(p, "含 P0～P5", "含 P0～P6")

    # --- 总结 ---
    for p in doc.paragraphs:
        t = p.text or ""
        if "含 P5 解谜/探索纵深/异步社交" in t and "P6" not in t:
            replace_in_paragraph(
                p,
                "含 P5 解谜/探索纵深/异步社交",
                "含 P5 解谜/探索纵深/异步社交与 P6 垂直移动/DoT/引导链/精通反应/装备强化/Boss 贡献",
            )
            log.append("更新总结段落")

    # --- 附录 C 修订信息 ---
    for p in doc.paragraphs:
        t = p.text or ""
        if "修订日期" in t or "生成日期" in t or "文档日期" in t:
            # 若整段是日期说明，追加修订注记
            if "P6" not in t:
                insert_runs_paragraph_after(
                    p,
                    f"修订说明：已原地增补 P6 二游手感纵深（{datetime.date.today().isoformat()}）。"
                    "请用 Word 打开后全选按 F9 更新目录域。",
                    size=11,
                )
                log.append("追加修订说明")
            break

    # 封面/副标题旁：若有「全方位」介绍句含 P5 则补 P6
    for p in doc.paragraphs:
        t = p.text or ""
        if "并吸收二游常见体验" in t and "装备强化" not in t:
            replace_in_paragraph(
                p,
                "并吸收二游常见体验（元素反应、体力、战令、签到、联机房间、行为树 Boss、神瞳式收集、区域声望、野外烹饪等）。",
                "并吸收二游常见体验（元素反应与精通附着量、攀爬/滑翔体力、装备强化、战令、签到、联机房间、"
                "行为树 Boss、神瞳式收集与引导链、区域声望、野外烹饪、世界 Boss 贡献分档等）。",
            )
            log.append("更新总结二游体验列举")
            break

    return log


def main():
    if not DOC_PATH.exists():
        raise SystemExit(f"未找到原文档: {DOC_PATH}")

    # 备份到同目录，仍只改原文件名输出
    bak = DOC_PATH.with_suffix(".docx.bak")
    shutil.copy2(DOC_PATH, bak)

    doc = Document(str(DOC_PATH))
    if already_patched(doc):
        print("文档似乎已含 P6 内容，仍执行幂等补丁（跳过已存在小节）…")

    log = patch(doc)
    doc.save(str(DOC_PATH))
    print(f"已原地更新: {DOC_PATH}")
    print(f"备份: {bak}")
    if log:
        print("变更项:")
        for line in log:
            print(" -", line)
    else:
        print("未检测到可写入的锚点（可能结构与预期不符）。")


if __name__ == "__main__":
    main()
