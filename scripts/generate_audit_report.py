# -*- coding: utf-8 -*-
"""生成 MyMmorpg 项目审计报告 Word 文档"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os
from datetime import date


def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color_hex,
        qn("w:val"): "clear",
    })
    shading.append(shd)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    return h


def add_para(doc, text, bold=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(11)
    run.bold = bold
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
    run = p.add_run(text)
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(11)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_shading(hdr_cells[i], "4472C4")
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for run in p.runs:
                    run.font.name = "微软雅黑"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                    run.font.size = Pt(10)
    doc.add_paragraph()
    return table


def main():
    doc = Document()

    # 页边距
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    # 封面标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MyMmorpg 项目审计报告")
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"生成日期：{date.today().strftime('%Y年%m月%d日')}")
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_paragraph()

    # ========== 一、项目概览 ==========
    add_heading(doc, "一、项目概览", 1)
    add_para(doc, "MyMmorpg 是一个基于 Spring Boot 3.2.5 + Java 17 的多模块 MMORPG 后端项目，采用 Protobuf + Netty/WebSocket 游戏协议，并预留 Spring Cloud Gateway + Nacos 微服务演进路径。")

    add_table(doc,
              ["模块", "职责"],
              [
                  ["mmorpg-common", "协议（Protobuf）、实体、Port 接口、RPC/MQ 基础设施"],
                  ["player-service", "主入口：Netty + WebSocket + HTTP，默认嵌入全部域服务"],
                  ["mmorpg-gateway", "网关：鉴权、限流、可观测性"],
                  ["scene-service", "场景/AOI/分线/移动"],
                  ["chat-service", "聊天（世界/私聊/队伍/公会）"],
                  ["bag-service", "背包（使用/出售/整理）"],
                  ["skill-service", "技能（学习/施法/CD）"],
                  ["battle-service", "战斗 + Buff"],
                  ["activity-service", "活动（列表/详情/领奖）"],
                  ["admin-service", "GM 后台 RBAC 数据层（REST 层不完整）"],
                  ["mmorpg-cli", "本地 CLI 流程模拟器"],
              ])

    add_para(doc, "默认运行模式：单体部署", bold=True)
    add_para(doc, "player-service 通过 Maven 依赖嵌入各域模块；game.battle.remote.enabled 与 game.activity.remote.enabled 默认为 false，battle/activity 远程调用未启用。")
    add_para(doc, "技术栈：Spring Boot 3.2.5 | Java 17 | Spring Cloud 2023.0.5 | Nacos | Protobuf 3.25.3 | Netty/WebSocket | Redis | MySQL | RocketMQ（默认关闭）")
    add_para(doc, "测试状态：全模块 mvn test 已通过。")

    # ========== 二、关键问题 ==========
    add_heading(doc, "二、关键问题（P0 — 必须修复）", 1)

    add_heading(doc, "2.1 网关鉴权形同虚设", 2)
    add_para(doc, "AuthGlobalFilter 已实现 Token 校验，但白名单放行了几乎所有游戏流量：")
    add_bullet(doc, "/ws/** — WebSocket 全部放行")
    add_bullet(doc, "/player/** — 玩家 HTTP 全部放行")
    add_bullet(doc, "/actuator/** — 监控端点放行")
    add_para(doc, "配置文件：mmorpg-gateway/src/main/resources/application.yml")
    add_para(doc, "风险：网关层对核心游戏通道不起作用；客户端直连 player-service:8989 同样绕过网关。")
    add_para(doc, "改进方案：", bold=True)
    add_bullet(doc, "白名单仅保留登录、公钥交换等必要路径（如 /player/login、/player/register）")
    add_bullet(doc, "WebSocket 握手携带 Token，或在首帧协议内校验")
    add_bullet(doc, "生产环境关闭或保护 /actuator/**")

    add_heading(doc, "2.2 内部微服务 API 无鉴权", 2)
    add_para(doc, "battle-service / activity-service 的内部接口仅信任 X-Player-Id 请求头，无 mTLS、无内部 Token、无 IP 限制。")
    add_para(doc, "涉及文件：battle-service/.../InternalBattleController.java、activity-service/.../InternalActivityController.java")
    add_para(doc, "风险：攻击者若可访问 8991/8992 端口，可伪造任意 playerId 发起战斗/领奖。")
    add_para(doc, "改进方案：", bold=True)
    add_bullet(doc, "内部 API 增加服务间 JWT 或 HMAC 签名校验")
    add_bullet(doc, "网络层限制仅 player-service 可访问（K8s NetworkPolicy / 安全组）")
    add_bullet(doc, "敏感操作增加幂等与审计日志")

    add_heading(doc, "2.3 硬编码密钥与弱默认凭证", 2)
    add_table(doc,
              ["位置", "问题"],
              [
                  ["config/common.yml", "RPC signKey 明文写死在仓库：5eb4bfea5d024e1d92c19cb6e4baab7d"],
                  ["player-service/application.yml", "MySQL 默认密码 123456"],
                  ["activity-service/application.yml", "MySQL 默认密码 123456"],
                  ["mmorpg-gateway/application.yml", "SSL keystore 密码默认 changeit"],
                  ["DevDataLoader.java", "@Profile(\"!test\")，非 test 环境可能写入 testuser/123456"],
              ])
    add_para(doc, "改进方案：", bold=True)
    add_bullet(doc, "DevDataLoader 改为 @Profile(\"dev\")")
    add_bullet(doc, "所有密钥走环境变量 / Secret Manager")
    add_bullet(doc, "RPC 签名升级为 HMAC-SHA256（当前 RpcSignUtil 使用 MD5，强度不足）")
    add_bullet(doc, "生产禁用默认密码，启动时校验必填环境变量")

    add_heading(doc, "2.4 Redis KEYS 命令用于在线统计", 2)
    add_para(doc, "AuthTokenService.countOnlineAccounts() 与 PlayerSessionService 使用 redisTemplate.keys() 扫描在线账号/玩家，在生产环境会阻塞 Redis。")
    add_para(doc, "涉及文件：player-service/.../AuthTokenService.java、PlayerSessionService.java")
    add_para(doc, "改进方案：", bold=True)
    add_bullet(doc, "登录/登出时用 SADD/SREM 维护 auth:online_accounts Set")
    add_bullet(doc, "或使用 INCR/DECR 计数器")
    add_bullet(doc, "需要扫描时用 SCAN 替代 KEYS")

    # ========== 三、中等问题 ==========
    add_heading(doc, "三、中等问题（P1 — 架构与可运维性）", 1)

    add_heading(doc, "3.1 微服务拆分未完成", 2)
    add_para(doc, "各域服务独立部署时注册 NoOp 占位实现，功能严重降级：")
    add_table(doc,
              ["服务", "NoOp 影响"],
              [
                  ["scene-service", "无玩家缓存 → 进场景失败；无 WebSocket 推送"],
                  ["chat-service", "私聊目标离线；世界频道推送无效"],
                  ["skill-service", "施法后 CD/学习通知不下发"],
                  ["battle-service", "无场景联动、无进度更新、无推送"],
                  ["activity-service", "无法发道具、无推送、预加载门控失效"],
              ])
    add_para(doc, "目前仅 battle/activity 有 Feign 远程调用，scene/chat/bag/skill 无内部 HTTP API，无法真正独立部署。")
    add_para(doc, "改进方案：", bold=True)
    add_bullet(doc, "明确两种部署模式文档：单体（默认）vs 微服务")
    add_bullet(doc, "为每个 *Port 补 Feign/gRPC 远程实现")
    add_bullet(doc, "为 scene/chat/bag/skill 补 application.yml 与独立启动入口")
    add_bullet(doc, "拆分 mmorpg-common：protocol / domain-api / infra 三层")

    add_heading(doc, "3.2 player-service 超级单体 + mmorpg-common 过重", 2)
    add_para(doc, "player-service Maven 依赖全部 7 个域模块，编译单元过大；mmorpg-common 同时引入 JPA、Netty、WebSocket、RocketMQ、Groovy，所有服务被迫引入完整栈。")
    add_para(doc, "改进方案：短期保持单体默认，用 Port + Gateway 模式隔离边界；中期域服务独立 JAR + Feign 调用。")

    add_heading(doc, "3.3 异常静默吞掉，故障难排查", 2)
    add_para(doc, "MessageDispatchPipeline 中 RPC 转发失败无日志，直接降级本地处理（catch Exception ignore）。")
    add_para(doc, "WebSocket 推送失败同样吞异常（PlayerPushRegistry）；MQ 发送失败仅 warn 日志。")
    add_para(doc, "改进方案：至少 log.warn + Prometheus 计数；关键路径考虑熔断而非无限静默降级。")

    add_heading(doc, "3.4 多个域服务缺少独立配置", 2)
    add_para(doc, "scene-service、chat-service、bag-service、skill-service 无 application.yml，独立启动需依赖 Nacos 或外部配置。")

    add_heading(doc, "3.5 过度注释，维护成本高", 2)
    add_para(doc, "150+ 文件含统一「文件维护说明」头注释；大量逐行 import 注释（如 AuthGlobalFilter.java），代码可读性反而下降。")
    add_para(doc, "改进方案：保留类级 Javadoc 与非 obvious 业务注释；删除模板化/重复注释。")

    # ========== 四、次要问题 ==========
    add_heading(doc, "四、次要问题（P2 — 质量与工程化）", 1)

    add_heading(doc, "4.1 测试覆盖不均衡", 2)
    add_table(doc,
              ["模块", "测试情况"],
              [
                  ["mmorpg-common", "~20 个，RPC/协议/工具覆盖较好"],
                  ["battle/activity/player-service", "有单元测试 + 部分集成测试"],
                  ["bag/chat/skill/scene", "各仅 1~2 个 Policy 测试"],
                  ["admin-service / mmorpg-cli", "0 个测试"],
              ])
    add_para(doc, "其他问题：admin-service testng-unit.xml 复制自 battle-service；test-suites 未聚合全部模块；无 WebSocket/Netty 端到端协议测试。")

    add_heading(doc, "4.2 DevOps / 文档缺失", 2)
    add_bullet(doc, "无 README、无架构图、无本地启动指南")
    add_bullet(doc, "无 CI/CD（无 .github/workflows）")
    add_bullet(doc, "docker-compose.yml 仅有 Redis + Nacos，无 MySQL")
    add_bullet(doc, "Nacos 关闭认证（NACOS_AUTH_ENABLE: false）")
    add_bullet(doc, "无各服务 Dockerfile")

    add_heading(doc, "4.3 未完成功能（TODO）", 2)
    add_table(doc,
              ["项", "状态"],
              [
                  ["SecurityController", "TODO：加密上传无持久化（OSS、分片合并）"],
                  ["AdminComplaintService", "TODO：投诉处理为骨架，无真实业务"],
                  ["admin-service", "无 REST Controller；admin.http.ips 白名单未 enforcement"],
                  ["RocketMQ", "默认 enabled: false，事件链路透 NoOp"],
                  ["mmorpg-cli", "不连真实服务器，Token 为假字符串"],
              ])

    add_heading(doc, "4.4 其他代码质量细节", 2)
    add_bullet(doc, "SessionCryptoService：RSA 密钥每次重启重新生成；会话 AES 密钥仅存内存，多节点不可用")
    add_bullet(doc, "RedisTestContainerHolder 在 player/battle/activity 三处各一份（重复代码）")
    add_bullet(doc, "网关路由仅指向 player-service，未为独立域服务配置路由")
    add_bullet(doc, "缺少 JaCoCo/SpotBugs/Sonar 等质量门禁")

    # ========== 五、设计亮点 ==========
    add_heading(doc, "五、值得保留的设计亮点", 1)
    add_bullet(doc, "Port 抽象模式（PlayerNotificationPort、BattleScenePort 等）为微服务演进预留了正确方向")
    add_bullet(doc, "BattleCommandGateway / ActivityCommandGateway 本地/远程切换设计清晰")
    add_bullet(doc, "Protobuf + msgId 路由（GameMessageFactory + @MessageRoute）协议层结构合理")
    add_bullet(doc, "幂等去重（IdempotencyService）与玩家分片串行（DispatchThreadModel）对 MMORPG 并发场景考虑周到")
    add_bullet(doc, "Groovy Policy 热替换策略（Battle/Chat/Skill/Scene/Item）灵活可扩展")
    add_bullet(doc, "Testcontainers + embedded-redis 双 fallback 的集成测试基础设施已有雏形")
    add_bullet(doc, "当前全模块单元测试可通过（mvn test 成功）")

    # ========== 六、实施路线图 ==========
    add_heading(doc, "六、优先行动建议与实施路线图", 1)

    add_heading(doc, "P0 — 安全（建议 1~2 周）", 2)
    add_bullet(doc, "收紧网关白名单；WebSocket 握手/首包校验 Token")
    add_bullet(doc, "内部 API 加服务间认证；/internal/** 限制网络访问")
    add_bullet(doc, "密钥/密码全部外置；DevDataLoader 限制 @Profile(\"dev\")")
    add_bullet(doc, "替换 Redis KEYS 为 Set/Counter")
    add_bullet(doc, "RPC 签名升级 HMAC-SHA256")

    add_heading(doc, "P1 — 架构（建议 2~4 周）", 2)
    add_bullet(doc, "明确部署模式文档：单体 vs 微服务切换步骤")
    add_bullet(doc, "为 scene/chat/bag/skill 补 Feign 内部 API + Port 远程实现")
    add_bullet(doc, "拆分 mmorpg-common；减少 player-service 直接 Maven 依赖")
    add_bullet(doc, "为各域服务补 application.yml 与 Docker Compose 全栈编排")

    add_heading(doc, "P2 — 质量与工程化（持续）", 2)
    add_bullet(doc, "补 admin/bag/chat/skill/scene 单元与集成测试；修 testng 配置")
    add_bullet(doc, "添加 GitHub Actions：mvn test + 集成测试 profile")
    add_bullet(doc, "编写 README（架构、端口、启动顺序、环境变量）")
    add_bullet(doc, "清理过度注释；RPC/MQ 失败路径加结构化日志与指标")
    add_bullet(doc, "完成 TODO：SecurityController 持久化、Admin REST + 投诉流程")

    # ========== 七、总结 ==========
    add_heading(doc, "七、总结", 1)
    add_para(doc, "MyMmorpg 项目在游戏核心逻辑与协议层设计方面表现较好，单体模式下可运行且测试通过。主要短板集中在三个方面：")
    add_bullet(doc, "安全边界未收紧 — 网关鉴权白名单过宽、内部 API 无保护、密钥硬编码")
    add_bullet(doc, "微服务拆分停留在 Port 抽象层 — 大量 NoOp 占位，scene/chat/bag/skill 无法独立部署")
    add_bullet(doc, "工程化配套不足 — 缺少 README、CI/CD、完整 docker-compose 与均衡测试覆盖")
    add_para(doc, "建议按 P0 → P1 → P2 优先级逐步推进改进，优先解决安全与生产风险问题。")

    # 保存
    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    filename = "MyMmorpg项目审计报告.docx"
    filepath = os.path.join(desktop, filename)
    doc.save(filepath)
    print(f"文档已保存至: {filepath}")
    return filepath


if __name__ == "__main__":
    main()
