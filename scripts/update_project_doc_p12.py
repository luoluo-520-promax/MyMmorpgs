# -*- coding: utf-8 -*-
"""就地更新桌面 Word 项目介绍文档，补充 P12 性能加固章节。"""
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC_PATH = r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx"


def insert_paragraph_after(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def replace_if_contains(doc, needle, new_text):
    for p in doc.paragraphs:
        if needle in p.text:
            p.text = new_text
            return True
    return False


def main():
    doc = Document(DOC_PATH)

    replace_if_contains(
        doc,
        "修订说明",
        "修订说明：在原有 P6–P11（元素共鸣/传说状态机/图鉴/指挥标记/地形冷却/肉鸽命运卡等）基础上，"
        "新增 P12 性能瓶颈加固（Actor 信箱、对象池、AOI 增量、Tick 分片、写流水线、Lua 原子锁、冷热分离）；"
        "技术栈截至 2026-08-22。请在 Word 中全选后按 F9 更新目录。",
    )

    replace_if_contains(
        doc,
        "perf/ 目录",
        "perf/ 目录含 JMeter/Gatling 脚本；"
        "GET /internal/scene/open-world/metrics/performance 观测 P12 指标；"
        "POST .../mixed-load-benchmark 混合压测；"
        "PerformanceHardeningFlowTest + PerformanceHardeningIntegrationTest 回归。"
        "注意：单进程基准 ≠ 集群千人同屏验收。",
    )

    # 在 P11 段落后插入 P12
    anchor = None
    for p in doc.paragraphs:
        if "5.2.9 P11" in p.text and p.style.name.startswith("Heading"):
            anchor = p
            break
    if anchor is None:
        for p in doc.paragraphs:
            if "TeamCompositionService" in p.text and "StoryStateMachine" in p.text:
                anchor = p
                break

    if anchor is not None:
        cur = anchor
        for line in [
            ("Heading 3", "5.2.10 P12 性能瓶颈加固（并发 / GC / AOI / Tick / DB / 锁 / 冷热分离）"),
            ("Normal", "代码详见 docs/open-world-top-tier.md P12 节；测试：PerformanceHardeningFlowTest、PerformanceHardeningIntegrationTest。"),
            ("List Bullet", "并发：SceneActorMailbox（Actor 信箱）+ BusinessBatchConsumer（无锁批量消费）+ DispatchThreadModel stripe；game.dispatch.use-batch-consumer 可开关。"),
            ("List Bullet", "GC：MoveCmdPool 复用移动指令槽位；PrimitiveGridStore 用 long[] 存网格坐标，SceneActorService.handleMove 已接入。"),
            ("List Bullet", "AOI：AoiDeltaEncoder 字段级增量 + AoiBroadcastStrategy 三级频次（10m/30m/50m）+ AoiUpdateBatcher 合并；game.scene.aoi-near/mid/far-distance 可配。"),
            ("List Bullet", "Tick：GameplayTickSlicer 20 片分帧（budget 5ms）；生态 AI 异步 runAsyncFireAndForget，不阻塞主场景线程。"),
            ("List Bullet", "DB：PlayerWritePipeline（50 条/1s batch）+ AssetWalWriter（WAL）；game.player.write-pipeline-ms 可配。"),
            ("List Bullet", "锁：RedisAtomicScriptService Lua 原子升级/首杀；POST /internal/scene/open-world/world-level/upgrade。"),
            ("List Bullet", "冷热：HistoricalDataRetention（热 3 天/冷 7 天 OSS）+ SimpleBloomFilter（图鉴/神瞳）+ BattleReplayService 强制归档。"),
            ("List Bullet", "观测：GET /internal/scene/open-world/metrics/performance → performance.concurrency/gcPooling/aoi/tick/persistence。"),
        ]:
            style, text = line
            cur = insert_paragraph_after(cur, text, style)

    # 10.1 已实现列表追加一条
    for p in doc.paragraphs:
        if "P10" in p.text and "生态状态机" in p.text and p.style.name.startswith("List"):
            insert_paragraph_after(
                p,
                "P12 性能加固：Actor 信箱、MoveCmd 对象池、AOI 增量与三级频次、Tick 分片、写流水线/WAL、Lua 原子世界等级升级、布隆过滤器与战斗重播冷热归档；已接入 SceneActorService、DispatchThreadModel、OpenWorldRuntimeService。",
                "List Bullet",
            )
            break

    # FAQ Q12
    for p in doc.paragraphs:
        if p.text.startswith("Q11") and "P11" in p.text:
            cur = insert_paragraph_after(
                p,
                "Q12：P12 性能指标与压测怎么看？",
                "Heading 2",
            )
            insert_paragraph_after(
                cur,
                "答：GET /internal/scene/open-world/metrics/performance 查看聚合指标；"
                "移动热路径已接 MoveCmdPool + AoiDeltaEncoder；"
                "mvn -pl mmorpg-common,scene-service test -Dtest=PerformanceHardening* 跑回归。",
                "Normal",
            )
            break

    # 附录日期
    for p in doc.paragraphs:
        if "文档日期" in p.text or "2026年08月21日" in p.text:
            if "P12" not in p.text:
                p.text = p.text.replace("2026年08月21日", "2026年08月22日（P12 性能加固）")
                if "2026年08月22日" not in p.text:
                    p.add_run("（P12 性能加固：2026-08-22）")

    doc.save(DOC_PATH)
    print("Updated:", DOC_PATH)


if __name__ == "__main__":
    main()
