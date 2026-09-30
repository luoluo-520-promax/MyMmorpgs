# -*- coding: utf-8 -*-
"""生成 MyMmorpg 项目健康检查报告（Word）并保存到桌面。"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os
from datetime import date


def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shd = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color_hex,
        qn("w:val"): "clear",
    })
    shading.append(shd)


def style_run(run, size=11, bold=False, color=None):
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        style_run(run, size=16 if level == 1 else 13, bold=True)
    return h


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    style_run(run, bold=bold)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.25)
    run = p.add_run(text)
    style_run(run)
    return p


def add_table(doc, headers, rows, header_color="4472C4"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        set_cell_shading(cell, header_color)
        for p in cell.paragraphs:
            for run in p.runs:
                style_run(run, size=10, bold=True, color=RGBColor(255, 255, 255))
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    style_run(run, size=10)
    doc.add_paragraph()
    return table


def add_issue(doc, title, severity, desc, evidence, solution):
    add_heading(doc, title, 2)
    add_para(doc, f"严重级别：{severity}", bold=True)
    add_para(doc, "问题描述：", bold=True)
    add_para(doc, desc)
    add_para(doc, "证据/涉及位置：", bold=True)
    for item in evidence:
        add_bullet(doc, item)
    add_para(doc, "解决方案：", bold=True)
    for item in solution:
        add_bullet(doc, item)


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MyMmorpg 项目健康检查报告")
    style_run(run, size=22, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"检查日期：{date.today().strftime('%Y年%m月%d日')}  |  基于当前代码库静态审查 + 编译验证")
    style_run(run, size=11, color=RGBColor(0x66, 0x66, 0x66))

    doc.add_paragraph()

    # ========== 一、总体结论 ==========
    add_heading(doc, "一、总体结论", 1)
    add_para(doc, "项目已从早期 demo 演进为多模块微服务骨架，单体编译可通过（mvn -DskipTests compile 成功）。"
             "相对历史审计，网关白名单、内部 API HMAC、DevDataLoader 限制、CI 基础流水线、DEPLOYMENT.md、"
             "域服务 application.yml 等已明显改善。")
    add_para(doc, "但仍处于“可演示 / 可本地开发”到“可稳定生产部署”的过渡期：配置治理未统一、"
             "部分核心玩法仍是内存态、CI 质量门禁未真正落地、文档与实现存在落差。")

    add_table(doc,
              ["维度", "评级", "说明"],
              [
                  ["编译可用性", "良好", "全模块 compile 通过"],
                  ["本地/演示可用性", "中等", "单体模式可运行，依赖本地 MySQL/Redis"],
                  ["生产部署可用性", "偏低", "prod 配置与密钥校验未收口，多服务状态不持久"],
                  ["测试可信度", "偏低", "CI 仅跑 test，未强制 quality-gate，缺基础设施编排"],
                  ["工程债务", "较高", "配置分裂、内存态服务、文档与实现不一致"],
              ])

    # ========== 二、已改善项 ==========
    add_heading(doc, "二、相较历史问题，已改善的项（无需再作为 P0）", 1)
    add_table(doc,
              ["历史问题", "当前状态"],
              [
                  ["网关鉴权白名单过宽（/ws、/player 全放行）", "已收紧为仅 /api/security/public-key 与 session-key"],
                  ["内部 API 无鉴权", "已有 InternalApiAuthFilter + HMAC 签名（默认启用）"],
                  ["DevDataLoader 非 test 环境乱写数据", "已改为 @Profile(\"dev\")"],
                  ["RPC signKey 明文硬编码", "config/common.yml 已改为 ${RPC_SIGN_KEY:}"],
                  ["无 CI", "已有 .github/workflows/ci.yml（unit + integration）"],
                  ["无部署说明", "已有 DEPLOYMENT.md（单体/微服务切换说明）"],
                  ["scene/chat/bag/skill 无独立配置", "多数域服务已补 application.yml / application-dev.yml"],
                  ["RedisTestContainerHolder 多处重复", "已收敛到 mmorpg-common/.../test/support"],
                  ["admin-service 无 REST", "已有 AdminComplaintController / AdminImportController 等"],
              ])

    # ========== 三、Critical ==========
    add_heading(doc, "三、Critical（会直接影响生产启动/部署稳定性）", 1)

    add_issue(
        doc,
        "3.1 生产密钥校验与各服务 prod 配置不一致",
        "Critical",
        "mmorpg-common 的 ProductionSecretsValidator 在 prod/production profile 下强制检查 "
        "MYSQL_PASSWORD、RPC_SIGN_KEY、INTERNAL_API_SECRET（及 SSL 密码）。但多数服务的 "
        "application-prod.yml 几乎为空或仅含局部开关，例如 player-service 的 prod 配置只打开 RocketMQ。"
        "这会导致：要么生产启动直接失败，要么依赖“隐式环境变量约定”却未文档化到每个服务。",
        [
            "mmorpg-common/.../ProductionSecretsValidator.java",
            "player-service/.../application-prod.yml（仅 rocketmq.enabled）",
            "各服务 application-prod.yml 内容不统一",
        ],
        [
            "统一制定“按服务类型”的必填密钥清单（player/admin/gateway/域服务可分层）",
            "为每个服务补齐 application-prod.yml 中的占位与环境变量映射，并在 DEPLOYMENT.md 列出矩阵",
            "将 ProductionSecretsValidator 改为按服务能力条件校验，或允许声明豁免项，避免误杀无 DB/无 RPC 的服务",
            "生产启动采用 fail-fast：缺密钥即失败，并输出清晰错误信息（当前已有基础，需与配置对齐）",
        ],
    )

    add_issue(
        doc,
        "3.2 内部 API 开关默认策略不统一，易出现“本地能跑、联调失败”",
        "Critical",
        "InternalApiAuthAutoConfiguration 默认 game.internal-api.enabled=true（matchIfMissing=true）。"
        "部分服务的 application-dev.yml 显式关闭，部分服务基础 application.yml 未声明。"
        "联调/CI/单体嵌入模式下，签名头缺失会导致 /internal/** 401，排查成本高。",
        [
            "mmorpg-common/.../InternalApiAuthAutoConfiguration.java",
            "各服务 application-dev.yml 中 game.internal-api.enabled 声明不一致",
            "InternalApiRestClient / Feign Interceptor 依赖同一 secret",
        ],
        [
            "在所有服务统一声明：dev=可关闭或使用固定开发密钥；test=与 IT 对齐；prod=强制开启且拒绝弱密钥",
            "把“启用状态 + secret 是否注入”写入各服务启动日志一行摘要，便于排障",
            "为内部调用链路补一条最小联调 smoke（player → battle/activity/hall）",
        ],
    )

    add_issue(
        doc,
        "3.3 CI 未覆盖质量门禁与真实依赖，回归防护不足",
        "Critical",
        "当前 CI 仅执行 mvn test 与 mvn test -Pintegration-tests，未执行 DEPLOYMENT.md 提到的 "
        "mvn verify -Pquality-gate（JaCoCo/SpotBugs），也未启动 MySQL/Redis/Nacos。"
        "大量“配置错误、跨服务调用失败、生产密钥策略”不会在合并前暴露。",
        [
            ".github/workflows/ci.yml",
            "根 pom.xml 中 quality-gate profile",
            "docker-compose.yml 仅有 Redis + Nacos（无 MySQL）",
        ],
        [
            "CI 增加 quality-gate job（或 nightly），失败即阻断",
            "为集成测试引入 service containers / Testcontainers（至少 MySQL + Redis）",
            "明确 unit / IT / quality 三类流水线职责，避免“假绿”",
        ],
    )

    # ========== 四、High ==========
    add_heading(doc, "四、High（高优先级：生产数据与运维风险）", 1)

    add_issue(
        doc,
        "4.1 hall / matchmaking / quest 等核心玩法仍为内存态",
        "High",
        "好友、邮件、匹配队列、任务进度使用 ConcurrentHashMap 保存在进程内存中。"
        "服务重启数据丢失；多实例部署状态不共享。文档写“可独立部署”，但这些实现仅适合演示。",
        [
            "hall-service/.../HallService.java（friendsByPlayer / mailsByPlayer）",
            "matchmaking-service/.../MatchmakingService.java（queues / matchStatuses）",
            "quest-service/.../QuestService.java（progressByPlayer）",
        ],
        [
            "短期：在 DEPLOYMENT.md / 接口说明明确标注“演示实现，不可生产多实例”",
            "中期：好友/邮件/任务落库（MySQL），匹配状态落 Redis（TTL + 队列键）",
            "多实例前先做会话粘性或共享存储，否则禁止水平扩容",
        ],
    )

    add_issue(
        doc,
        "4.2 DEPLOYMENT.md 与代码完成度存在落差",
        "High",
        "部署文档描述了完整微服务切换步骤、端口矩阵、环境变量。"
        "但部分服务仍是骨架/内存实现，远程 Port 开关与真实能力不完全对等，容易误导运维按“已生产就绪”部署。",
        [
            "DEPLOYMENT.md",
            "RemotePortAutoConfiguration / 各服务 remote.enabled 开关",
            "hall/matchmaking/quest 内存实现与文档“独立部署”表述冲突",
        ],
        [
            "文档按模块标注：已实现 / 演示实现 / 计划中",
            "每个服务列出“必须环境变量 / 可选环境变量 / 依赖组件”",
            "补充一页“已知限制”避免把规划当现状",
        ],
    )

    add_issue(
        doc,
        "4.3 docker-compose 缺少 MySQL，本地一键依赖不完整",
        "High",
        "本地依赖编排仅有 Redis 与关闭认证的 Nacos，无 MySQL。"
        "新成员按 compose 启动后仍需自行准备数据库，且 Nacos 无认证不适合任何共享环境。",
        [
            "docker-compose.yml（redis + nacos，NACOS_AUTH_ENABLE=false）",
        ],
        [
            "compose 增加 MySQL 服务与初始化 SQL/脚本挂载",
            "提供 .env.example（端口、账号、密钥占位）",
            "共享/预发环境启用 Nacos 认证，禁止默认关闭认证镜像配置外泄",
        ],
    )

    add_issue(
        doc,
        "4.4 battle-service 仍使用 Redis KEYS 扫描",
        "High",
        "BattleService 使用 stringRedisTemplate.keys(REDIS_ACTIVE + \"*\") 扫描活跃战斗键。"
        "生产 Redis 上 KEYS 会阻塞，在线人数升高时风险显著。player-service 侧历史 KEYS 问题已清理。",
        [
            "battle-service/.../BattleService.java（keys 扫描）",
        ],
        [
            "维护 battle:active 的 Set/Hash 索引，登录开战时写入、结束时删除",
            "必须扫描时改用 SCAN 游标，并限制频率",
            "补充对应单元/集成测试，避免回归回 KEYS",
        ],
    )

    add_issue(
        doc,
        "4.5 admin-service 安全组合复杂，误配置风险高",
        "High",
        "同时存在 IP 白名单、HMAC、API Key、内部 API secret。"
        "dev 默认关闭鉴权；prod 依赖环境变量。漏配时可能启动失败或接口不可用，排障路径长。",
        [
            "admin-service/.../AdminApiAuthFilter.java / AdminIpWhitelistFilter.java",
            "AdminProductionSecretsValidator.java",
            "admin-service application-dev.yml / application-prod.yml",
        ],
        [
            "文档化认证优先级与最小必填组合（例如：prod 至少 HMAC+IP 或 API Key+IP）",
            "给出一份可直接复制的生产环境变量样例",
            "启动失败时错误信息明确指出缺哪一类凭证",
        ],
    )

    # ========== 五、Medium / Low ==========
    add_heading(doc, "五、Medium / Low（工程化与可维护性）", 1)

    add_issue(
        doc,
        "5.1 缺少顶层 README.md",
        "Medium",
        "仓库无 README，新人通常无法快速找到模块结构、启动方式与文档入口。"
        "DEPLOYMENT.md 不能替代 README。",
        ["仓库根目录无 README* 文件"],
        [
            "新增 README：简介、模块表、单体启动、测试命令、环境变量入口、链接到 DEPLOYMENT.md",
        ],
    )

    add_issue(
        doc,
        "5.2 配置默认值大量指向 127.0.0.1 / 空密码，易产生“假可运行”",
        "Medium",
        "application.yml 中普遍默认本机 URL、空密码、开发密钥。"
        "本地方便，但错误 profile 或漏配环境变量时可能静默连到错误目标。",
        [
            "player-service 及其他域服务 application.yml",
        ],
        [
            "prod profile 禁止兜底到 localhost/空密钥（与 ProductionSecretsValidator 对齐）",
            "dev 保留默认值，但日志明确提示“开发默认配置”",
        ],
    )

    add_issue(
        doc,
        "5.3 测试套件入口多层，易漏测",
        "Medium",
        "同时存在模块内 testng-unit/integration.xml、根工程 suite 变量、test-suites 聚合套件。"
        "新增测试可能未进入统一入口，导致根执行与模块执行结果不一致。",
        [
            "根 pom.xml 的 ${testng.suiteXmlFile}",
            "各模块 src/test/resources/testng-*.xml",
            "test-suites/ 聚合套件",
        ],
        [
            "明确唯一主入口；聚合套件自动覆盖新增模块",
            "CI 打印实际使用的 suite 文件路径，便于核对",
        ],
    )

    add_issue(
        doc,
        "5.4 会话加密/密钥多节点一致性风险（历史债务遗留）",
        "Medium",
        "若 SessionCrypto 仍在进程内生成/保存会话密钥，多实例网关或 player 节点下会话不可共享。"
        "需确认当前是否已外置到 Redis；若仍内存，则禁止无粘性多节点。",
        [
            "mmorpg-common/.../SessionCryptoConfiguration.java",
            "相关 SessionCrypto 服务实现",
        ],
        [
            "将会话密钥与在线状态统一落到 Redis，并设定 TTL",
            "多节点部署前做 sticky session 或共享存储验证",
        ],
    )

    add_issue(
        doc,
        "5.5 过度模板化注释仍较重",
        "Low",
        "大量文件含“文件维护说明”与逐行 import 注释，增加噪声，真实业务意图反而被淹没。",
        ["多个服务源文件头注释与逐行注释风格"],
        [
            "保留类级职责与非 obvious 业务注释；逐步清理模板化/重复注释",
        ],
    )

    # ========== 六、验证结果 ==========
    add_heading(doc, "六、本次检查已执行的验证", 1)
    add_table(doc,
              ["检查项", "结果"],
              [
                  ["mvn -DskipTests compile（全模块）", "通过（exit code 0）"],
                  ["关键安全组件存在性", "InternalApiAuth* / ProductionSecretsValidator / Admin 过滤器均存在"],
                  ["网关白名单", "已收紧，不再放行全部 /ws 与 /player"],
                  ["内存态服务代码审查", "hall / matchmaking / quest 确认为 ConcurrentHashMap"],
                  ["CI 工作流", "存在 unit + integration，无 quality-gate / 无基础设施服务"],
                  ["README", "缺失"],
                  ["docker-compose", "仅 Redis + Nacos，无 MySQL"],
              ])

    # ========== 七、整改路线图 ==========
    add_heading(doc, "七、建议整改路线图", 1)

    add_heading(doc, "第一阶段（1~2 周）— 先消 Critical", 2)
    add_bullet(doc, "统一 prod 密钥校验策略与各服务 application-prod.yml / 环境变量矩阵")
    add_bullet(doc, "统一 game.internal-api.enabled 与 secret 在 dev/test/prod 的行为")
    add_bullet(doc, "CI 增加 quality-gate；至少补 MySQL+Redis 到集成测试或 compose")
    add_bullet(doc, "替换 battle-service 的 Redis KEYS")

    add_heading(doc, "第二阶段（2~4 周）— 消 High", 2)
    add_bullet(doc, "hall/邮件/好友、quest 进度落库；matchmaking 状态进 Redis")
    add_bullet(doc, "修订 DEPLOYMENT.md：已实现 vs 演示 vs 计划中")
    add_bullet(doc, "完善 admin 生产鉴权样例与启动校验提示")
    add_bullet(doc, "补顶层 README 与 .env.example")

    add_heading(doc, "第三阶段（持续）— 工程化", 2)
    add_bullet(doc, "统一 TestNG suite 入口，补薄弱模块测试")
    add_bullet(doc, "收敛配置默认值策略（prod fail-fast）")
    add_bullet(doc, "清理过度注释；补跨服务 smoke 与可观测性指标")

    # ========== 八、总结 ==========
    add_heading(doc, "八、总结", 1)
    add_para(doc, "结论：项目“有明显进展，但仍有必须处理的问题”。")
    add_para(doc, "最值得优先处理的四件事：")
    add_bullet(doc, "生产密钥与 profile 配置统一治理")
    add_bullet(doc, "CI / 集成测试基础设施补齐（避免假绿）")
    add_bullet(doc, "hall / matchmaking / quest 状态持久化（或明确仅演示）")
    add_bullet(doc, "补 README + 修正部署文档与实现的一致性")
    add_para(doc, "在完成上述事项前，建议生产环境优先使用单体模式小流量验证，"
             "不要按“全微服务多实例”方式直接上线内存态服务。")

    desktop = os.path.join(os.path.expanduser("~"), "Desktop")
    filename = "MyMmorpg项目健康检查报告.docx"
    filepath = os.path.join(desktop, filename)
    doc.save(filepath)
    print(f"文档已保存至: {filepath}")
    return filepath


if __name__ == "__main__":
    main()
