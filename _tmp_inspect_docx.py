# -*- coding: utf-8 -*-
from docx import Document
path = r'c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx'
doc = Document(path)
out = r'c:\Users\ASUS\IdeaProjects\test\MyMmorpg\_tmp_docx_outline.txt'
with open(out, 'w', encoding='utf-8') as f:
    f.write(f'paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}\n')
    for i, p in enumerate(doc.paragraphs):
        text = (p.text or '').strip()
        if not text:
            continue
        style = p.style.name if p.style else ''
        f.write(f'{i}|{style}|{text}\n')
    f.write('\n=== TABLES ===\n')
    for ti, t in enumerate(doc.tables):
        f.write(f'\n--- table {ti} rows={len(t.rows)} cols={len(t.columns)} ---\n')
        for ri, row in enumerate(t.rows[:8]):
            cells = [c.text.strip().replace('\n',' ')[:40] for c in row.cells]
            f.write(f'  r{ri}: ' + ' | '.join(cells) + '\n')
print('wrote', out)