# -*- coding: utf-8 -*-
"""Generate MyMmorpg project introduction Word document on Desktop."""
from __future__ import annotations

import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, name="宋体", size=12, bold=False, color=None, east_asia=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    if east_asia is None:
        east_asia = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if color:
        run.font.color.rgb = color


def set_paragraph_spacing(p, before=0, after=6, line=1.5):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    sizes = {1: 18, 2: 15, 3: 13}
    for run in p.runs:
        set_run_font(
            run,
            name="黑体",
            size=sizes.get(level, 12),
            bold=True,
            color=RGBColor(0x1F, 0x4E, 0x79),
            east_asia="黑体",
        )
    return p


def add_para(doc, text, first_indent=True, bold=False, size=12, align="justify"):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    set_paragraph_spacing(p, before=0, after=6, line=1.5)
    run = p.add_run(text)
    set_run_font(run, name="宋体", size=size, bold=bold, east_asia="宋体")
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Cm(0.75 + level * 0.5)
    p.paragraph_format.first_line_indent = Cm(0)
    set_paragraph_spacing(p, before=0, after=3, line=1.35)
    run = p.add_run(text)
    set_run_font(run, name="宋体", size=12, east_asia="宋体")
    return p


def shade_cell(cell, fill="D6E3F0"):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_run_font(run, name="黑体", size=10.5, bold=True, east_asia="黑体")
        shade_cell(cell)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, name="宋体", size=10.5, east_asia="宋体")
    doc.add_paragraph()
    return table


def add_code_block(doc, text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=9, east_asia="宋体")
    doc.add_paragraph()


def add_page_field(paragraph):
    """Append PAGE field to paragraph."""
    run = paragraph.add_run()
    run._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run2 = paragraph.add_run()
    run2._r.append(
        parse_xml(
            f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>'
        )
    )
    run3 = paragraph.add_run()
    run3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


def add_numpages_field(paragraph):
    run = paragraph.add_run()
    run._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run2 = paragraph.add_run()
    run2._r.append(
        parse_xml(
            f'<w:instrText {nsdecls("w")} xml:space="preserve"> NUMPAGES </w:instrText>'
        )
    )
    run3 = paragraph.add_run()
    run3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


def add_toc_field(paragraph):
    run = paragraph.add_run()
    run._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    run2 = paragraph.add_run()
    run2._r.append(
        parse_xml(
            f'<w:instrText {nsdecls("w")} xml:space="preserve">'
            r' TOC \o "1-3" \h \z \u '
            "</w:instrText>"
        )
    )
    run3 = paragraph.add_run()
    run3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
    run4 = paragraph.add_run("（打开 Word 后请右键此处并选择「更新域」以生成带页码的目录）")
    set_run_font(run4, name="宋体", size=11, color=RGBColor(0x80, 0x80, 0x80), east_asia="宋体")
    run5 = paragraph.add_run()
    run5._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)

    for i, (size, ea) in enumerate([(18, "黑体"), (15, "黑体"), (13, "黑体")], 1):
        hs = doc.styles[f"Heading {i}"]
        hs.font.name = ea
        hs.font.size = Pt(size)
        hs.font.bold = True
        hs.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        hs._element.rPr.rFonts.set(qn("w:eastAsia"), ea)
        hs.paragraph_format.space_before = Pt(14 if i == 1 else 10)
        hs.paragraph_format.space_after = Pt(8)

    # Header / Footer
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("MyMmorpg 项目详细介绍")
    set_run_font(hr, name="宋体", size=9, color=RGBColor(0x66, 0x66, 0x66), east_asia="宋体")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = fp.add_run("第 ")
    set_run_font(r1, name="宋体", size=9, east_asia="宋体")
    add_page_field(fp)
    r3 = fp.add_run(" 页 / 共 ")
    set_run_font(r3, name="宋体", size=9, east_asia="宋体")
    add_numpages_field(fp)
    r5 = fp.add_run(" 页")
    set_run_font(r5, name="宋体", size=9, east_asia="宋体")

    # ===== Cover =====
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        p.add_run("MyMmorpg"),
        name="黑体",
        size=36,
        bold=True,
        color=RGBColor(0x1F, 0x4E, 0x79),
        east_asia="黑体",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("大型多人在线角色扮演游戏（MMORPG）"), name="黑体", size=18, bold=True, east_asia="黑体")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        p.add_run("服务端项目详细介绍文档"),
        name="黑体",
        size=22,
        bold=True,
        color=RGBColor(0x2E, 0x75, 0xB6),
        east_asia="黑体",
    )

    for _ in range(2):
        doc.add_paragraph()

    today = datetime.date.today().strftime("%Y年%m月%d日")
    meta_items = [
        ("技术栈", "Java 17 · Spring Boot 3.2 · Spring Cloud Alibaba"),
        ("工程类型", "Maven 多模块微服务 / 可单体部署"),
        ("文档版本", "1.0"),
        ("编写日期", today),
        ("适用读者", "产品、策划、开发、测试、运维及对游戏服务端感兴趣的读者"),
    ]
    for k, v in meta_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(f"{k}：{v}"), name="宋体", size=12, east_asia="宋体")

    doc.add_page_break()

    # ===== TOC =====
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("目  录"), name="黑体", size=22, bold=True, east_asia="黑体")

    add_para(
        doc,
        "说明：请用 Microsoft Word 或 WPS 打开本文档后，在下方目录区域右键选择「更新域」→「更新整个目录」，"
        "即可自动生成带页码的完整目录。页脚「第 X 页 / 共 Y 页」同样会在更新域或打印预览后正确显示。",
        first_indent=False,
        size=10.5,
    )

    toc_p = doc.add_paragraph()
    add_toc_field(toc_p)

    doc.add_paragraph()
    add_para(doc, "【章节结构预览】", first_indent=False, bold=True)
    toc_preview = [
        "一、项目概述与定位",
        "二、适合谁阅读本文档",
        "三、用大白话理解：这套系统在做什么",
        "四、技术架构总览",
        "五、两种部署模式：单体与微服务",
        "六、各业务模块详解",
        "七、客户端如何与服务器通信",
        "八、核心玩法能力说明",
        "九、数据存储与中间件",
        "十、安全、支付与反作弊",
        "十一、运维、热更新与可观测性",
        "十二、如何本地快速跑起来",
        "十三、测试与质量保障",
        "十四、能力成熟度：已实现 / 演示 / 计划中",
        "十五、与 MyLunarCore 的对齐要点",
        "十六、术语表（小白友好）",
        "十七、常见问题 FAQ",
        "十八、总结与后续演进",
    ]
    for t in toc_preview:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.3
        set_run_font(p.add_run(t), name="宋体", size=12, east_asia="宋体")

    doc.add_page_break()

    # ===== Body =====
    add_heading_cn(doc, "一、项目概述与定位", 1)
    add_para(
        doc,
        "MyMmorpg 是一套用 Java 编写的大型多人在线角色扮演游戏（MMORPG）服务端示例工程。"
        "简单说：它就是游戏「背后那台脑子」——负责处理玩家登录、走路、打架、聊天、背包、任务、抽卡、"
        "商城充值、活动奖励等几乎所有服务器逻辑。",
    )
    add_para(doc, "项目基于 Java 17 与 Spring Boot 3，并引入 Spring Cloud Alibaba 相关能力，支持两种运行方式：")
    add_bullet(
        doc,
        "单体模式（默认）：主要启动一个 player-service，把场景、大厅、任务、匹配、背包等能力「嵌」在同一个进程里，"
        "适合本地开发、学习和小规模演示。",
    )
    add_bullet(
        doc,
        "微服务模式：把战斗、活动、场景、聊天等拆成独立服务，通过 Feign 远程调用与消息队列协作，"
        "适合进一步学习分布式架构与线上扩展。",
    )
    add_para(
        doc,
        "本项目不是「只能看不能跑」的空壳：仓库内包含数据库脚本、Docker Compose 依赖、内部 API、后台管理、"
        "测试套件与部署说明，目标是让读者既能理解 MMORPG 服务端怎么切分领域，也能真正把服务启动起来体验。",
    )

    add_heading_cn(doc, "二、适合谁阅读本文档", 1)
    add_para(doc, "本文刻意用通俗中文写成，尽量少堆砌黑话。不同角色可以这样使用：")
    add_table(
        doc,
        ["读者角色", "建议关注章节", "能得到什么"],
        [
            ["完全新手 / 产品同学", "一、三、六、八、十六", "理解游戏服在干什么、有哪些玩法模块"],
            ["策划同学", "六、八、十一、十四", "活动导入、热更、奖励与配置驱动能力"],
            ["后端开发", "四、五、六、七、九、十", "模块边界、协议、存储、安全与拆分路径"],
            ["测试同学", "十二、十三、十七", "如何启动、如何跑测试、常见坑"],
            ["运维同学", "五、九、十、十一、十二", "端口、环境变量、监控、生产校验"],
        ],
    )

    add_heading_cn(doc, "三、用大白话理解：这套系统在做什么", 1)
    add_heading_cn(doc, "3.1 把游戏拆成「前台」和「后台」", 2)
    add_para(
        doc,
        "玩家手机或电脑上的客户端，负责画面、按键、特效；服务器负责「权威真相」：你到底有没有这件装备、"
        "伤害算不算暴击、活动进度有没有到、钱有没有真付成功。客户端可以「申请」，服务器说了才算。",
    )
    add_heading_cn(doc, "3.2 一台电脑也能玩转的「游戏世界」", 2)
    add_para(
        doc,
        "开发时，你可以用 Docker 拉起 MySQL 和 Redis，再启动 player-service。"
        "这时你已经拥有：账号登录、角色数据、场景移动、战斗、大厅好友邮件、任务、匹配、背包、商城等一整条链路的骨架。"
        "想更接近线上，再把各域服务拆开单独启动即可。",
    )
    add_heading_cn(doc, "3.3 为什么要做成很多「小服务」", 2)
    add_para(
        doc,
        "MMORPG 功能极多。如果全部揉在一个巨大程序里，改抽卡可能不小心弄坏战斗，扩容场景服务器时也得把整坨一起扩。"
        "拆成 scene（场景）、battle（战斗）、activity（活动）等服务后，边界更清晰，也方便按压力分别扩容——"
        "例如场景人多就多开场景节点，战斗高峰就加战斗池。",
    )

    add_heading_cn(doc, "四、技术架构总览", 1)
    add_heading_cn(doc, "4.1 技术选型一览", 2)
    add_table(
        doc,
        ["层次", "技术", "作用（人话）"],
        [
            ["语言与运行时", "Java 17", "主流企业级语言，性能与生态成熟"],
            ["应用框架", "Spring Boot 3.2.5", "快速搭建 HTTP 服务、依赖注入、配置管理"],
            ["微服务套件", "Spring Cloud / Alibaba", "服务发现、Feign 调用、网关、可选 Nacos"],
            ["实时通信", "Netty（TCP）+ 可选 UDP/KCP", "游戏高频消息通道，比纯 HTTP 更适合实时玩法"],
            ["协议", "Protobuf", "紧凑的二进制消息格式，省流量、结构清晰"],
            ["关系库", "MySQL", "账号、角色、好友、邮件、订单等持久数据"],
            ["缓存/状态", "Redis", "在线态、匹配队列、排行榜、战斗临时状态等"],
            ["消息队列", "RocketMQ（可关）", "最终一致事件：战斗结束、订单支付等"],
            ["构建", "Maven 多模块", "协议、公共库、各业务服务分模块管理"],
            ["容器依赖", "Docker Compose", "一键起 MySQL/Redis（及可选 Nacos）"],
            ["质量", "TestNG、JaCoCo、SpotBugs、CI", "测试、覆盖率、静态检查、持续集成"],
        ],
    )

    add_heading_cn(doc, "4.2 工程模块结构", 2)
    add_para(
        doc,
        "仓库是一个 Maven 父工程 mmorpg-parent，下面挂着多个子模块。可以粗分为「公共底座」和「业务服务」两类。",
    )
    add_para(doc, "公共底座：", first_indent=False, bold=True)
    add_bullet(doc, "mmorpg-protocol：游戏协议（.proto）与消息号、返回码等。")
    add_bullet(
        doc,
        "mmorpg-domain-api：跨服务的 Port 接口（例如背包发货命令），让模块之间依赖「契约」而不是硬绑实现。",
    )
    add_bullet(doc, "mmorpg-common：公共适配、远程 Port、过渡聚合层。")
    add_bullet(doc, "mmorpg-gateway：可选 API 网关（默认端口 8443），做路由、区域头、粘滞等。")
    add_bullet(doc, "mmorpg-cli：命令行工具，例如导入活动配置、更新清单。")
    add_para(doc, "业务服务（节选）：", first_indent=False, bold=True)
    add_bullet(doc, "player-service：玩家入口、会话、长连接接入，单体模式下会嵌入多个域。")
    add_bullet(
        doc,
        "scene / chat / bag / skill / hall / quest / matchmaking：场景、聊天、背包、技能、大厅、任务、匹配。",
    )
    add_bullet(
        doc,
        "battle / activity / shop / gacha / update / admin：战斗、活动、商城、抽卡、客户端更新、后台管理。",
    )

    add_heading_cn(doc, "4.3 逻辑架构（一句话版）", 2)
    add_para(
        doc,
        "客户端 →（可选 Gateway）→ player-service（鉴权与长连接入口）→ 各域服务（场景/战斗/活动/…）→ "
        "MySQL / Redis / RocketMQ。后台运营人员则通过 admin-service 做配置导入、热更与投诉处理。",
    )

    add_heading_cn(doc, "4.4 领域拆分原则（给开发看）", 2)
    add_para(
        doc,
        "第一批拆分通常以 player（入口）、battle（战斗计算）、activity（活动进度）为最小可落地集合。"
        "跨服务一致性优先用「事件驱动 + 幂等处理 + 最终一致」，而不是处处上分布式事务。"
        "player-service 尽量成为唯一对外入口（含长连接），其它服务只做领域能力。"
        "更细的 Feign 清单与事件清单见仓库 docs/service-boundaries.md。",
    )

    add_heading_cn(doc, "五、两种部署模式：单体与微服务", 1)
    add_heading_cn(doc, "5.1 单体模式（默认，推荐入门）", 2)
    add_para(
        doc,
        "只启动 player-service（HTTP 端口 8989）。场景、大厅、任务、匹配、背包等通过 Maven 嵌入同一进程，"
        "远程开关 game.*.remote.enabled 默认为 false。",
    )
    add_para(doc, "典型启动步骤：", first_indent=False)
    add_code_block(
        doc,
        """# 1. 启动依赖
docker compose up -d mysql redis

# 2. 复制环境变量
cp .env.example .env

# 3. 启动玩家服（dev 配置）
mvn -pl player-service -am spring-boot:run -Dspring-boot.run.profiles=dev""",
    )
    add_para(doc, "适合：本地开发、集成测试、课程演示、功能联调。")

    add_heading_cn(doc, "5.2 微服务模式", 2)
    add_para(
        doc,
        "各域以独立 JAR 进程运行；player-service 更像「游戏网关」，用 Feign 把命令转到对应服务；"
        "域服务需要回调时，通过远程 Port 回到玩家服或其它服务。",
    )
    add_para(doc, "各服务 HTTP 端口一览：", first_indent=False)
    add_table(
        doc,
        ["服务", "端口", "一句话职责"],
        [
            ["player-service", "8989", "客户端连接、会话、玩家数据聚合入口"],
            ["scene-service", "8981", "场景、AOI、怪物、大世界域"],
            ["chat-service", "8982", "聊天"],
            ["bag-service", "8983", "背包与道具"],
            ["skill-service", "8984", "技能"],
            ["admin-service", "8985", "运营后台"],
            ["hall-service", "8986", "好友、邮件、排行、助战、家园"],
            ["quest-service", "8987", "任务"],
            ["matchmaking-service", "8988", "匹配（含跨服副本）"],
            ["shop-service", "8990", "商城货架与订单支付"],
            ["battle-service", "8991", "战斗与 AI 队友"],
            ["activity-service", "8992", "活动与世界 BOSS"],
            ["update-service", "8993", "客户端版本与资源更新"],
            ["gacha-service", "8994", "抽卡（保底/概率/历史）"],
            ["mmorpg-gateway", "8443", "可选 HTTP API 网关"],
        ],
    )
    add_para(
        doc,
        "切换微服务时，需要统一配置 INTERNAL_API_SECRET，打开各 GAME_*_REMOTE_ENABLED，并正确填写各 *_SERVICE_URL。"
        "生产环境还会强制校验弱密钥、MOCK 支付等危险配置。",
    )

    add_heading_cn(doc, "5.3 微服务开关速查（概念）", 2)
    add_para(
        doc,
        "当 GAME_SCENE_REMOTE_ENABLED、GAME_BATTLE_REMOTE_ENABLED、GAME_ACTIVITY_REMOTE_ENABLED 等打开后，"
        "player-service 会把对应命令转发到独立服务。独立 activity-service 还需要 GAME_PORT_REMOTE_ENABLED "
        "与 BAG_SERVICE_URL，否则发奖可能是空操作；生产对此会 fail-fast。",
    )

    add_heading_cn(doc, "六、各业务模块详解", 1)
    add_heading_cn(doc, "6.1 玩家服务（player-service）", 2)
    add_para(
        doc,
        "这是玩家最先碰到的服务器。它负责账号与角色、Token 会话、在线管理，以及 Netty/WebSocket 长连接入口："
        "客户端发来的二进制游戏包，大多先到这里再路由到内部能力。单体模式下，它还会「带着」其它域一起跑，"
        "所以入门只起这一个服务就够。",
    )

    add_heading_cn(doc, "6.2 场景服务（scene-service）", 2)
    add_para(doc, "场景服务管「你在哪个地图、附近能看见谁、怪物在哪」。核心能力包括：")
    add_bullet(
        doc,
        "AOI（Area of Interest，兴趣区域）：用空间哈希网格（AoiGrid）快速算出附近玩家，避免把全图所有人广播给你。",
    )
    add_bullet(doc, "动态域（WorldZoneManager）：按玩家密度拆分/合并 Zone，为大规模同屏做底座。")
    add_bullet(doc, "分线：单图人数过多时自动开线，空线超时回收。")
    add_bullet(doc, "无缝交接票据：跨节点迁移时带上速度、朝向、zone 等信息，尽量让切换不那么「闪一下」。")
    add_bullet(doc, "反作弊接入：对移动超速、瞬移等做校验。")

    add_heading_cn(doc, "6.3 战斗服务（battle-service）", 2)
    add_para(
        doc,
        "负责战斗开始、行动、结算等核心流程；可发布战斗领域事件（如 BattleEnded），供活动、任务去更新进度。"
        "还提供 AI 队友指挥接口（集火、治疗、散开、跟随等）。挑战关卡（协议段 13xx）会联动战斗开始逻辑。",
    )

    add_heading_cn(doc, "6.4 活动服务（activity-service）", 2)
    add_para(
        doc,
        "管理活动配置、玩家活动进度、领奖；支持配置导入与校验；可消费战斗结束、商店支付等事件更新进度。"
        "世界 BOSS 玩法在此沉淀：伤害榜、结算发奖计划（grantPlans）等。"
        "生产环境下若发奖 Port 未正确接到背包，会 fail-fast，避免「看起来发奖了其实是空操作」。",
    )

    add_heading_cn(doc, "6.5 大厅服务（hall-service）", 2)
    add_para(
        doc,
        "偏社交与外围系统：好友在线（读 Redis 在线键）、邮件（含附件道具发奖）、排行榜、事件临时小队、"
        "好友助战借用、家园拜访等。部分社交能力目前是 Redis 快照级演示，尚未做成完整养成/装修产品。",
    )

    add_heading_cn(doc, "6.6 匹配服务（matchmaking-service）", 2)
    add_para(
        doc,
        "把玩家从「想打」变成「开好一局」。支持双人副本/PVP，以及跨服 4 人副本匹配（match type=3）。"
        "队列可分片，并可用互补评分（等级、战力、等待时间）提升组队体验；"
        "跨实例用 Redis SET NX + Lua 保证一致性相关逻辑。",
    )

    add_heading_cn(doc, "6.7 背包与技能（bag / skill）", 2)
    add_para(
        doc,
        "背包管道具增减与幂等发奖；技能管技能相关逻辑。"
        "邮件附件发奖会走 MailItemGrantPort → bag，保证重复投递也不轻易重复给道具。",
    )

    add_heading_cn(doc, "6.8 任务服务（quest-service）", 2)
    add_para(
        doc,
        "任务配置可由 JSON 驱动（如 config/quest/Quests.json），支持 Admin/Internal 导入与 reload。"
        "战斗结束等事件可推进日常任务进度。",
    )

    add_heading_cn(doc, "6.9 商城与支付（shop-service）", 2)
    add_para(
        doc,
        "负责氪金货架与订单。经济链路做了硬化：订单双写、支付事件走 Outbox、发奖/累充幂等、钱包/道具流水账本；"
        "提供对账入口对比渠道账单差异。开发可用 MOCK 支付，生产必须关闭 MOCK 并配置渠道密钥。",
    )

    add_heading_cn(doc, "6.10 抽卡服务（gacha-service）", 2)
    add_para(
        doc,
        "独立抽卡微服务（8994，也可嵌入）。协议段约 14xx：保底、十连、天井、历史；"
        "提供概率公示与审计摘要接口，方便合规与运营核对。",
    )

    add_heading_cn(doc, "6.11 聊天、更新、后台（chat / update / admin）", 2)
    add_bullet(doc, "chat-service：聊天通道。")
    add_bullet(doc, "update-service：客户端版本与资源更新清单。")
    add_bullet(
        doc,
        "admin-service：运营后台。含 IP 白名单、HMAC/API Key、配置导入（活动/更新清单）、投诉处理、操作日志；"
        "以及分阶段热更与发布审计。还提供 AI 内容生成（活动/商城/投诉/战斗周报/任务草案等，LLM 可选，失败回退模板）。",
    )

    add_heading_cn(doc, "七、客户端如何与服务器通信", 1)
    add_heading_cn(doc, "7.1 通道类型", 2)
    add_bullet(doc, "HTTP：登录、部分 REST/内部接口、网关路由。")
    add_bullet(doc, "Netty TCP：游戏主实时通道（二进制协议）。")
    add_bullet(
        doc,
        "可选 UDP/KCP：降低延迟或做特定推送；带写缓冲水位与背压策略（不可写时丢弃部分推送）。",
    )

    add_heading_cn(doc, "7.2 协议与消息号", 2)
    add_para(
        doc,
        "业务消息使用 Protobuf。不同玩法占用不同消息号段，例如皮肤约 12xx、挑战 13xx、抽卡 14xx、肉鸽 15xx。"
        "返回码统一管理，便于客户端提示错误原因。",
    )

    add_heading_cn(doc, "7.3 内部 API", 2)
    add_para(
        doc,
        "服务与服务之间走 /internal/**。生产环境用 HMAC-SHA256 签名（时间戳 + 签名 + 玩家 ID 等头），"
        "防止内网接口被任意调用。命令型接口可走 Protobuf 流；Port 型可用 JSON。",
    )

    add_heading_cn(doc, "7.4 中心路由与跨节点迁移", 2)
    add_para(
        doc,
        "Center 可配置 local 或 remote。玩家跨场景节点时，会签发一次性 session_ticket；"
        "客户端收到重定向（如 SCENE_TRANSFER_REDIRECT）后连到目标节点并消费票据。"
        "重复登录会踢旧连接。断线可用 ResumeScene 一类能力在宽限期内恢复。",
    )

    add_heading_cn(doc, "八、核心玩法能力说明", 1)
    add_heading_cn(doc, "8.1 大世界底座", 2)
    add_para(
        doc,
        "动态域网格 + AOI + 无缝交接，是「很多人在同一张大地图」的基础。"
        "仓库还提供 AOI 压测接口，用于单进程千人级查询基准（注意：这不等于真实集群千人同屏压测结论）。",
    )
    add_heading_cn(doc, "8.2 世界 BOSS 与事件小队", 2)
    add_para(doc, "世界事件可排期、累计伤害、结算产出发奖计划。大厅侧支持玩法驱动的临时小队。")
    add_heading_cn(doc, "8.3 轻社交：助战与家园", 2)
    add_para(doc, "好友助战允许借用好友角色快照；家园可发布布局供好友参观。当前偏演示与异步社交骨架。")
    add_heading_cn(doc, "8.4 皮肤、挑战、肉鸽", 2)
    add_bullet(doc, "皮肤衣柜/穿戴：配置驱动。")
    add_bullet(doc, "挑战关卡：与战斗服务联动开局。")
    add_bullet(doc, "肉鸽：局内天赋加点，战斗结算回写。")
    add_heading_cn(doc, "8.5 AI 相关能力", 2)
    add_bullet(doc, "玩家顾问（rule-coach）：规则建议为主，可选 LLM。")
    add_bullet(doc, "智能 NPC 对话：默认规则线索回退。")
    add_bullet(doc, "战斗 AI 队友：听指挥走位。")
    add_bullet(doc, "Admin 内容生成：辅助策划产出草案。")
    add_para(doc, "设计原则：LLM 挂了也不应阻断游戏主循环，规则回退保底。")

    add_heading_cn(doc, "九、数据存储与中间件", 1)
    add_heading_cn(doc, "9.1 MySQL 多库", 2)
    add_para(
        doc,
        "Compose 会初始化多套库（如 player / battle / activity / update），对应领域拆分。"
        "玩家侧落库包括好友、邮件、任务进度等表；商店有订单、账本、幂等表；活动有进度与 MQ Inbox 等。",
    )
    add_heading_cn(doc, "9.2 Redis 的典型用途", 2)
    add_table(
        doc,
        ["键/用途示例", "说明"],
        [
            ["player:online:{id}", "好友在线态"],
            ["match:queue:* / match:player:*", "匹配队列与玩家匹配状态"],
            ["rank:level / rank:power", "排行榜 ZSET；miss 可回退 DB Top50 回填"],
            ["battle:stats:daily|weekly:*", "战斗日/周统计 Hash"],
            ["center:migration:ticket:*", "跨节点迁移票据"],
            ["活跃战斗索引 SET", "避免使用危险的 KEYS 扫描"],
        ],
    )
    add_para(doc, "生产可切换 Redis Sentinel 高可用（见 docker-compose.redis-sentinel.yml）。")

    add_heading_cn(doc, "9.3 RocketMQ", 2)
    add_para(
        doc,
        "开发默认关闭（NoOp），避免本机强依赖。生产若走最终一致事件，需显式 ROCKETMQ_ENABLED=true。"
        "必须走 MQ 的典型链路包括：商店订单已支付 → 活动充值进度；战斗结束 → 活动/任务进度（另有 HTTP 投影兜底）。"
        "配合 Outbox/Inbox 做可靠投递与幂等消费。",
    )

    add_heading_cn(doc, "十、安全、支付与反作弊", 1)
    add_heading_cn(doc, "10.1 内部与后台鉴权", 2)
    add_bullet(doc, "内部 API：HMAC；生产强制强密钥，拒绝弱默认。")
    add_bullet(
        doc,
        "Admin：HMAC 或 API Key + 用户 ID；可开 IP 白名单；RBAC 权限如 import:activity、complaint:handle。",
    )
    add_bullet(doc, "会话 RSA：生产要求固定密钥；未配置时开发可临时生成，但生产会 fail-fast。")

    add_heading_cn(doc, "10.2 支付安全", 2)
    add_para(
        doc,
        "支持渠道沙箱 HMAC 与 HTTP 验签适配。生产禁用 MOCK；建议开启生产验签。"
        "对账 SPI（ChannelBillProvider + reconcile）用于发现差异单与重试履约候选。",
    )

    add_heading_cn(doc, "10.3 反作弊底座", 2)
    add_para(
        doc,
        "AntiCheatService 覆盖移动超速/瞬移、伤害溢出、违规计数与临时封禁等，并已接入场景移动校验。"
        "后续可继续规则引擎生产化。",
    )

    add_heading_cn(doc, "十一、运维、热更新与可观测性", 1)
    add_heading_cn(doc, "11.1 分阶段热更与发布审计", 2)
    add_para(
        doc,
        "运营导入活动或更新清单后，可调用 POST /admin/ops/reload 做分阶段热更："
        "activity → update → player（清空策划表 Redis 缓存等）。"
        "GET /admin/ops/publish-history 可查看发布审计，方便回溯「谁在什么时候发了什么」。",
    )

    add_heading_cn(doc, "11.2 网关能力", 2)
    add_bullet(doc, "按账号粘滞到同一 player-service 实例，减少会话漂移。")
    add_bullet(
        doc,
        "RegionAwareRoutingFilter 注入区域相关头，配合基础设施 GSLB/智能 DNS 做就近接入。",
    )
    add_bullet(doc, "TraceId 透传，便于全链路排查。")

    add_heading_cn(doc, "11.3 监控指标与告警", 2)
    add_para(
        doc,
        "各服务暴露 Actuator：health/info/metrics/prometheus。关键业务有 Prometheus 计数器（如 dispatch、mq、shop）。"
        "仓库提供 Prometheus 配置与告警骨架（验签失败、履约停滞、宕机、堆内存、Outbox 堆积等）。",
    )

    add_heading_cn(doc, "11.4 容错", 2)
    add_para(
        doc,
        "scene/hall/battle 等 Feign 客户端配有 FallbackFactory 与 circuit-breaker 开关，"
        "下游抖动时尽量降级而不是整链崩溃。",
    )

    add_heading_cn(doc, "十二、如何本地快速跑起来", 1)
    add_heading_cn(doc, "12.1 环境准备", 2)
    add_bullet(doc, "安装 JDK 17、Maven、Docker Desktop（或兼容环境）。")
    add_bullet(doc, "克隆本仓库，进入根目录。")

    add_heading_cn(doc, "12.2 最小可行路径（单体）", 2)
    add_code_block(
        doc,
        """docker compose up -d mysql redis
cp .env.example .env
mvn -pl player-service -am spring-boot:run -Dspring-boot.run.profiles=dev""",
    )
    add_para(doc, "客户端默认连 player-service：HTTP 8989；游戏 Netty 端口见 config/ 配置。")

    add_heading_cn(doc, "12.3 常用命令", 2)
    add_code_block(
        doc,
        """mvn -B test                          # 单元测试
mvn -B test -Pintegration-tests      # 集成测试（需 Docker）
mvn -B verify -Pquality-gate         # JaCoCo + SpotBugs 质量门禁
docker compose up -d                 # MySQL + Redis + 可选 Nacos""",
    )

    add_heading_cn(doc, "12.4 配置入口", 2)
    add_para(
        doc,
        "完整环境变量样例见 .env.example；部署矩阵与生产注意项见 DEPLOYMENT.md；"
        "数据库迁移说明见 docs/db-migration.md。",
    )

    add_heading_cn(doc, "十三、测试与质量保障", 1)
    add_para(doc, "项目使用 TestNG。每个模块通常有：")
    add_bullet(doc, "testng-unit.xml：快速单元测试（默认）。")
    add_bullet(doc, "testng-integration.xml：功能集成测试（可能依赖 Docker 中的 MySQL/Redis）。")
    add_para(
        doc,
        "根目录 test-suites/ 提供全仓库聚合套件。CI（GitHub Actions）会跑单元/集成与质量门禁："
        "SpotBugs 失败即失败，JaCoCo 行覆盖率设有下限。建议日常用 mvn test，合并前用集成测试与 quality-gate。",
    )

    add_heading_cn(doc, "十四、能力成熟度：已实现 / 演示 / 计划中", 1)
    add_heading_cn(doc, "14.1 已实现（可认真使用）", 2)
    add_para(
        doc,
        "包括但不限于：双模式部署、内部 HMAC、Admin 导入与热更审计、AI 顾问/NPC/队友/内容生成骨架、"
        "Center 路由与迁移票据、大世界 AOI/动态域、世界 BOSS 结算计划、跨服匹配、抽卡服、反作弊底座、"
        "支付对账 SPI、Feign 降级、好友邮件任务落库、匹配与排行 Redis、Netty+KCP、CI 与 Compose 等。"
        "详细清单以仓库 DEPLOYMENT.md 为准。",
    )
    add_heading_cn(doc, "14.2 演示级（能跑但别当完整产品）", 2)
    add_bullet(doc, "部分奖励数值为演示常量。")
    add_bullet(doc, "助战/家园为快照级社交。")
    add_bullet(doc, "AI 以规则为主。")
    add_bullet(doc, "AOI 压测是单进程基准。")
    add_bullet(doc, "RocketMQ 开发默认关闭。")
    add_heading_cn(doc, "14.3 计划中", 2)
    add_bullet(
        doc,
        "继续拆 CommandPort，去掉 player 对域实现的硬依赖，走向更纯粹的 gateway-only。",
    )
    add_bullet(doc, "更完整的分布式锁、官方支付 SDK 直连、KCP 与参考实现全量对齐。")
    add_bullet(doc, "真实 GeoIP/GSLB、跨区域多活、反作弊规则引擎生产化、集群级千人同屏压测。")

    add_heading_cn(doc, "十五、与 MyLunarCore 的对齐要点", 1)
    add_para(doc, "项目在设计上参考并对齐 MyLunarCore 一类成熟服务端思路，重点包括：")
    add_bullet(doc, "运维面优先外拆：Admin 导入 + 分阶段热更与发布审计。")
    add_bullet(doc, "AI 智能面：顾问、NPC、AI 队友、内容生成（可关 LLM）。")
    add_bullet(
        doc,
        "游戏实时面：Netty + 可选 KCP；Center 本地/远程；大厅好友；匹配评分；皮肤/挑战/抽卡/肉鸽协议段。",
    )
    add_bullet(
        doc,
        "大世界与轻社交：动态域、AOI、世界 BOSS、临时小队、跨服副本、助战、家园、区域就近、反作弊与支付对账。",
    )

    add_heading_cn(doc, "十六、术语表（小白友好）", 1)
    add_table(
        doc,
        ["术语", "通俗解释"],
        [
            ["MMORPG", "很多人同时在线的角色扮演游戏"],
            ["微服务", "把大系统拆成多个可独立部署的小服务"],
            ["单体", "很多功能跑在同一个进程里，开发简单"],
            ["Feign", "写接口就能发起 HTTP 远程调用的客户端"],
            ["AOI", "只同步你附近的人/物，省流量省算力"],
            ["Protobuf", "一种紧凑的结构化二进制数据格式"],
            ["HMAC", "用共享密钥做消息完整性与身份校验"],
            ["幂等", "同一请求做多次，结果与做一次一样（防重复发奖）"],
            ["Outbox/Inbox", "发消息前先落库、收消息先记已处理，防丢防重"],
            ["热更", "尽量不停机刷新配置或缓存"],
            ["Sentinel（Redis）", "Redis 高可用：主从+哨兵自动故障转移"],
            ["GSLB", "全球负载均衡，让玩家连更近的接入点"],
            ["Port（六边形架构）", "对外依赖的抽象接口，便于替换实现"],
        ],
    )

    add_heading_cn(doc, "十七、常见问题 FAQ", 1)
    add_para(doc, "Q1：我只是想看看能不能跑，最少做什么？", first_indent=False, bold=True)
    add_para(doc, "起 MySQL+Redis，再 spring-boot:run player-service 的 dev 配置即可。")
    add_para(doc, "Q2：为什么活动发奖没效果？", first_indent=False, bold=True)
    add_para(
        doc,
        "独立 activity-service 时需要打开远程 Port 并配置 BAG_SERVICE_URL；否则可能是 NoOp。"
        "生产会直接拒绝危险空配置。",
    )
    add_para(doc, "Q3：生产启动直接失败？", first_indent=False, bold=True)
    add_para(
        doc,
        "检查是否仍使用弱默认密钥、是否未关 MOCK 支付、会话 RSA 是否未固定等。"
        "生产校验是故意「严」的，避免带病上线。",
    )
    add_para(doc, "Q4：目录页码不显示？", first_indent=False, bold=True)
    add_para(
        doc,
        "用 Word/WPS 打开后，在目录处右键「更新域」。页脚「第 X 页 / 共 Y 页」也会在打印预览或更新域后正确显示。",
    )
    add_para(doc, "Q5：和完整商业网游差在哪？", first_indent=False, bold=True)
    add_para(
        doc,
        "这是教学与工程示范向的完整骨架：很多玩法是可运行的底座或演示级实现，"
        "数值、运营后台体验、跨区多活、官方支付直连等仍需按产品继续打磨。",
    )

    add_heading_cn(doc, "十八、总结与后续演进", 1)
    add_para(
        doc,
        "MyMmorpg 用一套可读、可跑、可测的代码，展示了现代 MMORPG 服务端的关键拼图："
        "实时接入、领域拆分、配置驱动运营、经济幂等、大世界底座、轻社交与 AI 辅助，"
        "以及从单体平滑走向微服务的路径。",
    )
    add_para(
        doc,
        "若你是学习者：先从单体跑通登录与一场战斗，再对照模块边界文档理解拆分。"
        "若你是工程实践者：优先关注内部鉴权、支付对账、Outbox 幂等、热更审计与可观测性——"
        "这些往往比「再堆一个玩法」更能决定系统是否可上线。",
    )
    add_para(
        doc,
        "后续演进方向已在仓库文档写明：继续 Port 化与 gateway-only、补齐 infra 模块、"
        "强化锁与支付 SDK、推进集群级压测与多活。欢迎在理解本文后，结合 README、DEPLOYMENT.md 与 docs/ 目录深入源码。",
    )

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        p.add_run("—— 文档结束 ——"),
        name="黑体",
        size=12,
        color=RGBColor(0x66, 0x66, 0x66),
        east_asia="黑体",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(
        p.add_run("更多细节请参阅仓库内 README.md、DEPLOYMENT.md、TESTING.md 与 docs/ 目录。"),
        name="宋体",
        size=10.5,
        color=RGBColor(0x66, 0x66, 0x66),
        east_asia="宋体",
    )

    return doc


def main():
    out = Path.home() / "Desktop" / "MyMmorpg项目详细介绍.docx"
    doc = build_document()
    doc.save(str(out))
    print(f"SAVED: {out}")
    print(f"EXISTS: {out.exists()} SIZE: {out.stat().st_size if out.exists() else 0}")


if __name__ == "__main__":
    main()
