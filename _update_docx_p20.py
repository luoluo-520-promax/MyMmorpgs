# -*- coding: utf-8 -*-
"""在原 Word 文档中增补 P20 人群拥挤与战斗性能加固（不新建文档）。"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC_PATH = Path(r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx")
BACKUP = DOC_PATH.with_suffix(".docx.bak-before-P20")


def set_paragraph_text(p, text: str) -> None:
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for run in p.runs[1:]:
        run.text = ""


def insert_paragraph_after(paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    if text:
        new_para.add_run(text)
    return new_para


def find_paragraph(doc, predicate):
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if predicate(t, p):
            return p
    return None


def already_has(doc, needle: str) -> bool:
    return any(needle in (p.text or "") for p in doc.paragraphs)


def main() -> None:
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    shutil.copy2(DOC_PATH, BACKUP)
    doc = Document(str(DOC_PATH))

    # ── 封面：生成日期 / 修订说明 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：") and "P20" not in t and "附录" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月27日（P20 人群拥挤与战斗性能加固；P19 大世界沉浸纵深补齐；"
                "P18 八大手感短板补齐；P17 联机一致性加固；P16 权威纵深加固）",
            )
            break

    rev = find_paragraph(doc, lambda t, _: t.startswith("修订说明："))
    if rev and "P20" not in (rev.text or ""):
        set_paragraph_text(
            rev,
            "修订说明：在原有 P6–P19（含性能加固、体验优化、体验纵深补齐、七大短板加固、"
            "生成式 AI 沉浸、权威纵深加固、联机一致性加固、八大手感短板、大世界沉浸纵深）基础上，"
            "增补 P20 人群拥挤与战斗性能加固：SceneTickMicroPipeline 微流水线（Movement 10ms / "
            "Logic 50ms / Sync 100ms）、BroadcastImportanceFuseService 广播重要性熔断、"
            "LocalDamageCounter 公会伤害本地聚合 + RedisPositionBatchWriter 静止跳过、"
            "ReactionValidator 客户端时间戳背书 ±100ms、TerrainStateVector 读写分离与延迟碎岩、"
            "DualClockService 双时钟、CombatEventRingBuffer 零 GC 战斗事件、"
            "BattleScenePodAllocator 战斗 Pod 与大世界进程隔离、ZoneLoadBalancer 接入场景密度；"
            "技术栈截至 2026-08-27。请在 Word 中全选后按 F9 更新目录。",
        )

    # ── 3.2 列表 bullet（插在 P19 bullet 后）──
    if not already_has(doc, "【P20 人群拥挤与战斗性能加固】"):
        p19b = find_paragraph(
            doc,
            lambda t, _: t.startswith("【P19 大世界沉浸纵深补齐】")
            or ("EcologicalTableauService" in t and "P19" in t and t.startswith("【")),
        )
        if p19b is None:
            p19b = find_paragraph(doc, lambda t, _: "PerceptionModifierService" in t and "P19" in t)
        if p19b:
            insert_paragraph_after(
                p19b,
                "【P20 人群拥挤与战斗性能加固】SceneTickMicroPipeline（Movement 10ms / Logic 50ms / "
                "Sync 100ms 微流水线）+ BroadcastImportanceFuseService（超带宽熔断，非战斗 2Hz 仅 Position）；"
                "LocalDamageCounter 1s 合并 ZINCRBY + RedisPositionBatchWriter 静止 3s 跳过；"
                "ReactionValidator.CLIENT_TIMESTAMP_BUFFER_MS=100 闪避背书；"
                "TerrainStateVector StampedLock 乐观读 + scheduleBump 200ms 延迟碎岩；"
                "DualClockService Wall-Clock 技能 CD；CombatEventRingBuffer 零 GC（投射物上限 200）；"
                "BattleScenePodAllocator 战斗 Pod 隔离（≤4/≤20 人内存快照）；"
                "ZoneLoadBalancer 接入 SceneActorService.refreshZoneDensity。",
                style="List Bullet",
            )

    # ── 3.2.16 P20 产品说明（插在 P19 测试 bullet 之后）──
    if not already_has(doc, "3.2.16 人群拥挤与战斗性能加固"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("测试：OpenWorldP19FlowTest")
            and "OpenWorldP19BusinessFlowTest" in t,
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("3.2.15 大世界沉浸纵深补齐"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "3.2.16 人群拥挤与战斗性能加固（P20）",
                style="Heading 3",
            )
            sections = [
                (
                    "Normal",
                    "微流水线解耦：SceneTickMicroPipeline 将 SceneTickEngine 拆为三条预算流水线——"
                    "MovementPipeline（10ms，仅位置/碰撞）、LogicPipeline（50ms，Buff/AI 分片）、"
                    "SyncPipeline（100ms，PhysicsStateHash 审计占位）。"
                    "SceneActorService 以 sceneComponentTick / sceneLogicTick / sceneSyncTick 三调度执行，"
                    "密集人群时高频移动不阻塞低频战斗结算，目标移动回包 P50 < 20ms。",
                ),
                (
                    "Normal",
                    "广播重要性熔断：BroadcastImportanceFuseService 监测单 Scene 节点出带宽，"
                    "超 512KB/s 时非战斗玩家同步降至 2Hz（仅 Position，不发 Rotation），"
                    "战斗中玩家与 Boss 保持高频；已接入 handleMove 广播链路与 performanceSnapshot。",
                ),
                (
                    "Normal",
                    "Redis 热 Key 防护：guild-service LocalDamageCounter 内存累加公会 Boss 伤害，"
                    "每 1 秒 ZINCRBY 合并提交 guild:boss:damage:{guildId}；"
                    "RedisPositionBatchWriter 只写变更帧，静止站立超过 3 秒的玩家跳过 Redis 写入，"
                    "减少约 50% 位置缓存写压力。",
                ),
                (
                    "Normal",
                    "战斗手感与物理：ReactionValidator 引入 CLIENT_TIMESTAMP_BUFFER_MS=100，"
                    "校验 ClientActionTimestamp，网络晚到 ≤100ms 仍判定闪避成功；"
                    "TerrainStateVector 使用 StampedLock 乐观读处理移动 revision 校验，"
                    "地形破坏 scheduleBump 先下发视觉碎岩、200ms 后 flushPendingWrites 最终生效；"
                    "DualClockService 技能 CD / Buff 用墙上时钟，局部子弹时间不影响 CD 真实流逝；"
                    "HitFeedbackService 下发 skillCdUsesWallClock。",
                ),
                (
                    "Normal",
                    "零 GC 与进程隔离：CombatEventRingBuffer + DamageEvent（int 状态码）预分配环形槽位，"
                    "投射物上限 200 复用最旧实体；BattleScenePodAllocator 将 OPEN_WORLD 大进程与 "
                    "BATTLE_INSTANCE 轻量 Pod（Boss 房 ≤4 人 / 攻城 ≤20 人，memorySnapshotOnly）物理隔离；"
                    "ZoneLoadBalancer.rebalance 接入 refreshZoneDensity。",
                ),
                (
                    "List Bullet",
                    "配置：game.scene.ecs-tick-ms=10、game.scene.logic-tick-ms=50、"
                    "game.scene.sync-tick-ms=100、game.guild.damage-flush-ms=1000、"
                    "game.cache.position-flush-ms=500。",
                ),
                (
                    "List Bullet",
                    "观测：SceneActorService.performanceSnapshot() → tickMicroPipeline / broadcastFuse；"
                    "GET .../gameplay/status → p20 快照（movementBudgetMs、clientTimestampBufferMs、"
                    "dualClock、combatEventRingBuffer、battleScenePodAllocator 等）。",
                ),
                (
                    "List Bullet",
                    "测试：PerformanceCrowdOptimizationFlowTest（common）；"
                    "PerformanceCrowdOptimizationBusinessFlowTest（scene）；"
                    "LocalDamageCounterTest、GuildBusinessFlowTest（guild）；"
                    "权威 docs/open-world-top-tier.md P20。",
                ),
            ]
            cur = h
            for style, text in sections:
                cur = insert_paragraph_after(cur, text, style=style)

    # ── 5.2.18 P20 实现对照 ──
    if not already_has(doc, "5.2.18 P20"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("测试：OpenWorldP19FlowTest / OpenWorldP19EdgeCaseTest"),
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("5.2.17 P19"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "5.2.18 P20 人群拥挤与战斗性能加固（微流水线 / 广播熔断 / Redis 热 Key / "
                "时间戳背书 / 地形读写分离 / 双时钟 / 零 GC / 战斗 Pod 隔离）",
                style="Heading 3",
            )
            cur = insert_paragraph_after(
                h,
                "代码详见 docs/open-world-top-tier.md P20 节；门面 OpenWorldGameplayFacade；"
                "GET /gameplay/status 返回 p20 快照；SceneActorService.performanceSnapshot() "
                "返回 tickMicroPipeline / broadcastFuse。",
                style="Normal",
            )
            bullets = [
                "Tick：SceneTickMicroPipeline + SceneComponentStore.entityIds()；"
                "SceneActorService sceneComponentTick / sceneLogicTick / sceneSyncTick。",
                "广播：BroadcastImportanceFuseService + AoiBroadcastStrategy；handleMove 熔断决策。",
                "缓存：RedisPositionBatchWriter.shouldWrite（静止 3s 跳过）；"
                "LocalDamageCounter（GuildExpeditionService 本地聚合）。",
                "战斗手感：ReactionValidator.CLIENT_TIMESTAMP_BUFFER_MS；"
                "DualClockService；HitFeedbackService.skillCdUsesWallClock。",
                "物理：TerrainStateVector.scheduleBump / flushPendingWrites / StampedLock 乐观读。",
                "GC：DamageEvent + CombatEventRingBuffer（MAX_PROJECTILES=200）。",
                "架构：BattleScenePodAllocator；ZoneLoadBalancer 接入 refreshZoneDensity。",
                "测试：PerformanceCrowdOptimizationFlowTest / PerformanceCrowdOptimizationBusinessFlowTest；"
                "LocalDamageCounterTest；GuildBusinessFlowTest 回归。",
            ]
            for b in bullets:
                cur = insert_paragraph_after(cur, b, style="List Bullet")

    # ── 第十章成熟度 ──
    if not already_has(doc, "P20 人群拥挤与战斗性能加固：微流水线"):
        p19m = find_paragraph(doc, lambda t, _: t.startswith("P19 大世界沉浸纵深补齐："))
        if p19m:
            insert_paragraph_after(
                p19m,
                "P20 人群拥挤与战斗性能加固：微流水线 Tick、广播重要性熔断、Redis 伤害本地聚合与"
                "静止位置跳过、客户端时间戳闪避背书、地形读写分离、双时钟技能 CD、"
                "零 GC 战斗事件环、战斗 Pod 与大世界进程隔离；已接入 SceneActorService、"
                "OpenWorldGameplayFacade、GuildExpeditionService。",
                style="List Bullet",
            )

    # ── 11.22 业务故事 ──
    if not already_has(doc, "11.22 主城拥挤到 Boss 房丝滑"):
        anchor = find_paragraph(doc, lambda t, _: t.startswith("11.21 从脚印到文明节律"))
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: "11.21" in t and "P19" in t)
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "11.22 主城拥挤到 Boss 房丝滑（P20）",
                style="Heading 2",
            )
            stories = [
                "主城 50+ 人同屏移动 → sceneComponentTick 仅跑 MovementPipeline 10ms；"
                "Logic/Sync 分片不阻塞移动回包。",
                "出带宽超阈值 → BroadcastImportanceFuseService 熔断，远处非战斗玩家仅收 Position 2Hz；"
                "攻城战斗玩家 markCombat 保持 10Hz 全量。",
                "公会 Boss 20 人同时 reportDamage → LocalDamageCounter 本地累加，"
                "1s 后批量 ZINCRBY，Redis CPU 不再毛刺。",
                "高 RTT 玩家闪避：ClientActionTimestamp 在窗口内但包晚到 80ms → "
                "ReactionValidator 仍返回 ok=true、clientTimestampBacked=true。",
                "碎崖：scheduleBump 立即 visualImmediate → 200ms 后 flushPendingWrites；"
                "移动校验 revisionOf 走 StampedLock 乐观读不卡主帧。",
                "子弹时间触发 → DualClockService Wall CD 仍按真实时间缩短；"
                "HitFeedback 下发 skillCdUsesWallClock=true。",
                "深渊 Boss 房 allocateBattlePod → memorySnapshotOnly，≤4 人独立 Pod；"
                "主城拥挤不影响 Boss 战弹反帧率。",
            ]
            cur = h
            for s in stories:
                cur = insert_paragraph_after(cur, s, style="Normal")

    # ── 第十四章 API 速览追加 P20 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("内部 HTTP 前缀 /internal/**") and "P20：" not in t:
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；P20：SceneActorService.performanceSnapshot（tickMicroPipeline/broadcastFuse）；"
                "GET .../gameplay/status（p20）；guild 远征 reportDamage 经 LocalDamageCounter；"
                "game.scene.ecs-tick-ms / logic-tick-ms / sync-tick-ms。",
            )
            break

    # ── FAQ Q21 ──
    if not already_has(doc, "Q21：P20"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("答：场景侧 InternalOpenWorldController")
            and "/feel/terrain-detail" in t,
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("Q20：P19"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "Q21：P20 微流水线/广播熔断/Redis 热 Key/闪避背书/双时钟/战斗 Pod 如何联调？",
                style="Heading 2",
            )
            insert_paragraph_after(
                h,
                "答：场景侧 SceneActorService.handleMove + sceneComponentTick/sceneLogicTick/"
                "sceneSyncTick；performanceSnapshot 观测 tickMicroPipeline、broadcastFuse。"
                "GET .../gameplay/status 含 p20 快照。"
                "公会侧 GuildExpeditionService.reportDamage 经 LocalDamageCounter 批量刷 Redis。"
                "回归：mvn -pl mmorpg-common,scene-service,guild-service test "
                "\"-Dtest=PerformanceCrowdOptimization*\"。"
                "权威说明 docs/open-world-top-tier.md（P20）。",
                style="Normal",
            )

    # ── 总结：P5～P20 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if "P5～P19" in t and "P20" not in t:
            set_paragraph_text(
                p,
                t.replace("P5～P19", "P5～P20").replace(
                    "联机叙事隔膜、伙伴幽灵、断线快进、拍卖熔断、配置双缓冲等）",
                    "联机叙事隔膜、伙伴幽灵、断线快进、拍卖熔断、配置双缓冲、"
                    "微流水线 Tick、广播熔断、Redis 热 Key 防护、闪避时间戳背书、"
                    "双时钟技能 CD、零 GC 战斗事件、战斗 Pod 进程隔离等）",
                ),
            )
            break
        if "P12 性能深化（ECS dirty tick" in t and "微流水线" not in t:
            set_paragraph_text(
                p,
                t + "；P20 微流水线 Movement/Logic/Sync 三预算、广播重要性熔断、"
                "LocalDamageCounter、静止位置 Redis 跳过、DualClockService、CombatEventRingBuffer、"
                "BattleScenePodAllocator。",
            )
            break

    # ── P12 节增补（5.2.10 后或正文）──
    if not already_has(doc, "P20 微流水线第二轮"):
        p12 = find_paragraph(doc, lambda t, _: t.startswith("5.2.10 P12 性能瓶颈加固"))
        if p12:
            insert_paragraph_after(
                p12,
                "P12 性能深化第三轮（P20）：SceneTickMicroPipeline、BroadcastImportanceFuseService、"
                "RedisPositionBatchWriter 静止跳过、LocalDamageCounter、DualClockService、"
                "CombatEventRingBuffer、BattleScenePodAllocator；测试 PerformanceCrowdOptimization*。",
                style="Normal",
            )

    # ── 附录 C 生成日期 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：2026年08月27日（P19") and "P20" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月27日（P20 人群拥挤与战斗性能加固；P19 大世界沉浸纵深补齐；"
                "P18 八大手感短板补齐；P17 联机一致性加固；P16 权威纵深加固；P12 性能深化第三轮）",
            )

    # ── 附录表格 open-world-top-tier ──
    for table in doc.tables:
        for row in table.rows:
            row_text = " ".join(c.text for c in row.cells)
            if "docs/open-world-top-tier.md" not in row_text:
                continue
            for cell in row.cells:
                if "P0～P19" in cell.text:
                    cell.text = cell.text.replace("P0～P19", "P0～P20")
                elif "P0～P18" in cell.text:
                    cell.text = cell.text.replace("P0～P18", "P0～P20")
                elif "P0～P16" in cell.text:
                    cell.text = cell.text.replace("P0～P16", "P0～P20")

    doc.save(str(DOC_PATH))
    print("OK: updated", DOC_PATH)
    print("BACKUP:", BACKUP)


if __name__ == "__main__":
    main()
