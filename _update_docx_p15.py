# -*- coding: utf-8 -*-
"""在原 Word 文档中增补 P15 生成式 AI 沉浸补齐（不新建文档）。"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

DOC_PATH = Path(r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx")
BACKUP = DOC_PATH.with_suffix(".docx.bak-before-P15")


def set_paragraph_text(p, text: str) -> None:
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for run in p.runs[1:]:
        run.text = ""


def insert_paragraph_after(paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except Exception:
            pass
    if text:
        new_para.add_run(text)
    return new_para


def find_paragraph(doc, predicate):
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if predicate(t, p):
            return p
    return None


def replace_if_contains(doc, old_sub: str, new_text: str) -> bool:
    for p in doc.paragraphs:
        if old_sub in (p.text or ""):
            set_paragraph_text(p, (p.text or "").replace(old_sub, new_text) if old_sub != (p.text or "") else new_text)
            # 若整段替换意图更强：当 old_sub 是整段开头标记时直接写 new_text
            if (p.text or "").startswith(old_sub[:6]) and len(new_text) > 20 and old_sub in new_text:
                set_paragraph_text(p, new_text)
            return True
    return False


def replace_paragraph_startswith(doc, prefix: str, new_text: str) -> bool:
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith(prefix) and "\t" not in t:
            set_paragraph_text(p, new_text)
            return True
    return False


def already_has(doc, needle: str) -> bool:
    return any(needle in (p.text or "") for p in doc.paragraphs)


def main() -> None:
    if not DOC_PATH.exists():
        raise SystemExit(f"文档不存在: {DOC_PATH}")

    shutil.copy2(DOC_PATH, BACKUP)
    doc = Document(str(DOC_PATH))

    # ── 封面：生成日期 / 修订说明 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：") and "2026年08月24日" in t and "P15" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月25日（P15 生成式 AI 沉浸补齐；P14 体验纵深补齐；P13 体验优化；P12 性能加固）",
            )
            break

    rev = find_paragraph(doc, lambda t, _: t.startswith("修订说明："))
    if rev and "P15" not in (rev.text or ""):
        set_paragraph_text(
            rev,
            "修订说明：在原有 P6–P14（含性能加固、体验优化、体验纵深补齐与 AI 平台集成）基础上，"
            "增补 P15 生成式 AI 沉浸补齐：World-Context 动态叙事（DynamicNarrativeService）、"
            "常驻 AI 伙伴 CompanionBot（MsgId 2500）、战斗连招诊断 CombatAnalyst、"
            "截图路点 ScreenWaypointTranslator、动态人格 DynamicPersonaEngine；"
            "技术栈截至 2026-08-25。请在 Word 中全选后按 F9 更新目录。",
        )

    # ── 1.4 / 能力关键词补充 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if "P14 客户端预测软回滚" in t and "P15" not in t:
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；以及 P15 动态叙事、常驻 AI 伙伴、战斗分析师、截图路点与动态人格引擎。",
            )
            break

    # ── 3.2 大世界列表：P15 bullet（插在 P14 bullet 后）──
    if not already_has(doc, "【P15 生成式 AI 沉浸补齐】"):
        p14b = find_paragraph(doc, lambda t, _: t.startswith("【P14 体验纵深补齐"))
        if p14b:
            insert_paragraph_after(
                p14b,
                "【P15 生成式 AI 沉浸补齐】DynamicNarrativeService（World Snapshot→意图 JSON→情感台词，"
                "经 AiContentGuard + RuleTriggerService 防乱发任务）；CompanionBotService"
                "（companion:memory:{playerId}、战况注入 SquadCommanderService、好感解锁调查点/地图标记，"
                "MsgId 2500 COMPANION_DIALOGUE）；CombatAnalystService（技能空转/闪避浪费/反应覆盖→口语建议+失误时间轴，"
                "可深链 EquipEnhanceService）；ScreenWaypointTranslator（玩家主动截图求助，视觉标签对齐坐标，会话 TTL 5min）；"
                "DynamicPersonaEngine（CASUAL/HARDCORE/SOCIAL + 简洁/详细/感性滑块，注入 {{PLAYER_PERSONA}}）。",
                style="List Bullet",
            )

    # ── 3.2.10 P15 产品说明（插在 3.2.9 测试段之后、3.4 之前）──
    if not already_has(doc, "3.2.10 生成式 AI 沉浸补齐"):
        anchor = find_paragraph(doc, lambda t, _: t.startswith("测试：OpenWorldP14FlowTest"))
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("3.2.9 体验纵深补齐"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "3.2.10 生成式 AI 沉浸补齐（P15）",
                style="Heading 3",
            )
            sections = [
                (
                    "Normal",
                    "动态叙事（World-Context Narrative）：将天气/纪元/区域潮汐、PlayerChronicle、"
                    "NpcLongTermMemory 好感与背包特殊道具打包为 Structured World Snapshot；"
                    "情境压缩生成意图 JSON（如 intent=COMFORT/MOURN、emotion、hint），再渲染带情感标签的台词；"
                    "不生成游戏逻辑。出站经 AiContentGuard 过滤；RuleTriggerService 拦截 AI 乱发隐藏任务。",
                ),
                (
                    "Normal",
                    "常驻 AI 伙伴（CompanionBot）：独立记忆体 companion:memory:{playerId}"
                    "（坠崖次数、采花偏好、战斗风格）；气泡对话 MsgId 2500 + TTS 情绪标记；"
                    "战况指令由 CompanionBot 生成并注入 SquadCommanderService（非 LLM 执行层）；"
                    "好感达标触发 ExplorationVitalityService 隐藏调查点 / MapMarkerService「伙伴觉得有趣的地方」。",
                ),
                (
                    "Normal",
                    "战斗连招诊断（AI Combat Analyst）：读取 BattleReplayService 指令流与 HitFeedback 延迟/命中；"
                    "规则层计算技能空转率、闪避窗口浪费率、元素反应覆盖率；生成层口语化建议并高亮失误时间轴；"
                    "可一键深链 EquipEnhanceService 应用推荐词条。",
                ),
                (
                    "Normal",
                    "截图路点（Screen-to-Waypoint）：玩家主动「截图求助」上传脱敏低分视觉标签（不传人脸/UI 边框）；"
                    "轻量视觉层识别地标后与 RegionProgressService 坐标对齐，生成方向箭头与「前方 50m 巨石下」描述；"
                    "图片会话最多保留 5 分钟、不落盘持久化。",
                ),
                (
                    "Normal",
                    "动态人格（Dynamic Persona）：基于 Handbook/RegionProgress/Abyss 等信号聚类"
                    " CASUAL / HARDCORE / SOCIAL；设置页「AI 交流风格」滑块（简洁/详细/感性）；"
                    "调用 ai-service 时 System Prompt 注入 {{PLAYER_PERSONA}}，调整语气与建议深度。",
                ),
                (
                    "List Bullet",
                    "API（ai-service）：POST /internal/ai/narrative/generate、/companion/dialogue|remember|battle-coach|passive-explore、"
                    "/analyst/feedback、/vision/waypoint、/persona/resolve|style；"
                    "玩家侧 POST /api/player/ai/analyst/feedback、/companion/dialogue、/vision/waypoint、/persona/style。",
                ),
                (
                    "List Bullet",
                    "协议：MessageId.COMPANION_DIALOGUE_SC_NOTIFY = 2500（伙伴气泡 UI + TTS 情绪）。",
                ),
                (
                    "List Bullet",
                    "测试：AiP15CapabilitiesTest、AiP15BusinessFlowTest、AiP15EdgeCaseTest（common）；"
                    "AiP15BusinessFlowTest / AiPlatformBusinessFlowTest.p15*（ai-service）；"
                    "PlayerAiP15PlatformFlowTest（player-service）；权威 docs/ai-enhancement.md 八.五节 P15。",
                ),
            ]
            cur = h
            for style, text in sections:
                cur = insert_paragraph_after(cur, text, style=style)

    # ── 3.7 AI 辅助面：扩写 P15 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("3.7 AI") and "辅助面" in t and "\t" not in t:
            if "P15" not in t:
                set_paragraph_text(
                    p,
                    "3.7 AI 辅助面（AI 平台已集成 · 含 P15）\n"
                    "AI 是「参谋」不是「裁判」，分两条物理隔离管道：管道 A（战斗 Tick）仅用规则/行为树/DDA，禁止 LLM；"
                    "管道 B（交互/运营/生成式沉浸）走 ai-service 或 player-service 内嵌 AiPlatformFacade，超时/熔断后回退 rule-coach。\n"
                    "玩家侧：POST /api/player/ai/advise · /recommend · /support · /tactical · /profile；"
                    "P15 增补 /analyst/feedback · /companion/dialogue · /vision/waypoint · /persona/style。\n"
                    "NPC：短期会话记忆 + 长期情感向量；P15 DynamicNarrativeService 按世界状态生成情境台词。\n"
                    "伙伴：CompanionBot 常驻陪伴（记忆/吐槽/战况建议/MsgId 2500）。\n"
                    "战斗：DDA + BossStrategyTier；P15 CombatAnalyst 战后复盘与配装深链。\n"
                    "视觉：ScreenWaypointTranslator 主动截图寻路（TTL 5min）。\n"
                    "人格：DynamicPersonaEngine 按玩法画像调整语气。\n"
                    "运营：Admin 内容校验、流失干预、MLOps；微服务 ai-service:8995。",
                )
            break

    # 补充 3.7 正文段中若仍是旧短句
    for p in doc.paragraphs:
        t = p.text or ""
        if "AI 是「参谋」不是「裁判」。玩家顾问默认 rule-coach" in t and "P15" not in t and "管道 A" not in t:
            set_paragraph_text(
                p,
                "AI 是「参谋」不是「裁判」。玩家顾问默认 rule-coach，LLM 关闭时回退规则；"
                "P15 起顾问/战术/客服会按 DynamicPersona 调整语气。"
                "智能 NPC 可带短期记忆/情绪，并可由 DynamicNarrative 结合世界 Snapshot 生成独一无二台词。"
                "常驻伙伴 CompanionBot 提供探索陪伴与战况吐槽。Admin 可生成活动/商城/任务草案与投诉分类、战斗周报。"
                "真正改数值、发道具仍走正式业务链路。",
            )
            break

    # ── 5.3.1 AI 平台：补 P15 API ──
    for p in doc.paragraphs:
        t = p.text or ""
        if "5.3.1 AI 推理平台" in t and "P15" not in t:
            set_paragraph_text(
                p,
                "5.3.1 AI 推理平台（ai-service + AiPlatformFacade · 含 P15）\n"
                "模块：ai-service（8995）+ mmorpg-common/ai/*。"
                "单体嵌入 player-service（game.ai.remote=false）；微服务 Feign AiPlatformClient + Gateway Path=/internal/ai/**。\n"
                "核心 API：POST /internal/ai/infer · /recommend/items · /support/ask · /tactical/advise · "
                "/retention/evaluate · /content/generate|validate · /npc/bond · GET /mlops/drift。\n"
                "P15 API：/narrative/generate · /companion/* · /analyst/feedback · /vision/waypoint · /persona/resolve|style；"
                "控制器 InternalAiP15Controller；门面方法 narrativeGenerate / companion* / analystFeedback / screenWaypointHelp / persona*。\n"
                "玩家 API：/api/player/ai/advise|recommend|support|tactical|profile|"
                "analyst/feedback|companion/dialogue|vision/waypoint|persona/style（需 Token）。\n"
                "核心类：DynamicNarrativeService、CompanionBotService、CombatAnalystService、"
                "ScreenWaypointTranslator、DynamicPersonaEngine。\n"
                "测试：AiPlatformBusinessFlowTest、AiP15*、PlayerAiP15PlatformFlowTest、"
                "AiInferenceGatewayTest、NpcMemoryMoodFlowTest。",
            )
            break

    # ── 5.2.13 P15 技术章节（插在 5.2.12 测试段后）──
    if not already_has(doc, "5.2.13 P15"):
        anchor = find_paragraph(
            doc,
            lambda t, _: t.startswith("测试：OpenWorldP14FlowTest / OpenWorldP14EdgeCaseTest"),
        )
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: t.startswith("5.2.12 P14"))
        if anchor is not None:
            h = insert_paragraph_after(
                anchor,
                "5.2.13 P15 生成式 AI 沉浸补齐（动态叙事 / 常驻伙伴 / 战斗分析 / 截图路点 / 动态人格）",
                style="Heading 3",
            )
            cur = insert_paragraph_after(
                h,
                "代码详见 docs/ai-enhancement.md 八.五节；统一门面 AiPlatformFacade；"
                "HTTP InternalAiP15Controller（/internal/ai/**）与 PlayerAiController（/api/player/ai/**）；"
                "协议 MessageId 2500 COMPANION_DIALOGUE_SC_NOTIFY。管道 B 异步生成，禁止进入战斗 Tick。",
                style="Normal",
            )
            bullets = [
                "叙事：DynamicNarrativeService + WorldSnapshot；AiContentGuard；RuleTriggerService 防 GRANT_QUEST 泄漏。",
                "伙伴：CompanionBotService（Redis 语义 companion:memory:{playerId}）；"
                "bindSquad → SquadCommanderService；bindExploration → ExplorationVitalityService / MapMarkerService。",
                "分析：CombatAnalystService 读 BattleReplayService.playback；失误时间轴 highlight；equipHints → EquipEnhanceService 深链。",
                "视觉：ScreenWaypointTranslator.uploadScreenshot / translate；IMAGE_TTL_MS=5min；隐私脱敏标记。",
                "人格：DynamicPersonaEngine.cluster + StyleSlider；adaptAdvice 改写战术/客服/分析师语气。",
                "配额：narrativeGenerate 走 LlmDailyQuota.tryAcquire；超限返回 llm_daily_quota_exceeded + degraded。",
                "测试：AiP15CapabilitiesTest / AiP15BusinessFlowTest / AiP15EdgeCaseTest；"
                "ai-service AiP15BusinessFlowTest；player-service PlayerAiP15PlatformFlowTest。",
            ]
            for b in bullets:
                cur = insert_paragraph_after(cur, b, style="List Bullet")

    # ── 第十章成熟度 ──
    if not already_has(doc, "P15 生成式 AI 沉浸补齐："):
        p14m = find_paragraph(doc, lambda t, _: t.startswith("P14 体验纵深补齐："))
        if p14m:
            insert_paragraph_after(
                p14m,
                "P15 生成式 AI 沉浸补齐：动态叙事、常驻 AI 伙伴（MsgId 2500）、战斗连招诊断、"
                "截图路点、动态人格；已接入 AiPlatformFacade、InternalAiP15Controller、PlayerAiController。",
                style="List Bullet",
            )

    # 更新智能 NPC 那一行
    for p in doc.paragraphs:
        t = p.text or ""
        if t.startswith("智能 NPC（短期+长期记忆）") and "P15" not in t:
            set_paragraph_text(
                p,
                "智能 NPC（短期+长期记忆）/ AI 顾问+推荐+客服+战术 / Boss BT+DDA / "
                "Admin AI 草稿+自动校验+流失干预 / ai-service 8995 / "
                "P15 动态叙事+CompanionBot+CombatAnalyst+截图路点+DynamicPersona。",
            )
            break

    # ── 11.16 业务故事 ──
    if not already_has(doc, "11.16"):
        anchor = find_paragraph(doc, lambda t, _: t.startswith("11.15 客户端预测到命运残响"))
        if anchor is None:
            anchor = find_paragraph(doc, lambda t, _: "命运残响" in t and t.startswith("GET /rogue"))
        # 找到 11.15 小节末尾
        if anchor is not None:
            cur = anchor
            for p in doc.paragraphs:
                if (p.text or "").startswith("11.15"):
                    cur = p
            # 向后推进到该小节最后一段
            walk = cur
            for _ in range(20):
                nxt = walk._p.getnext()
                if nxt is None:
                    break
                cand = Paragraph(nxt, walk._parent)
                ct = (cand.text or "").strip()
                if cand.style and cand.style.name.startswith("Heading") and not ct.startswith("11.15"):
                    break
                if ct.startswith("Q") or ct.startswith("12.") or ct.startswith("第"):
                    break
                walk = cand
            cur = walk
            h = insert_paragraph_after(cur, "11.16 世界共鸣叙事到人格化建议（P15）", style="Heading 2")
            stories = [
                "击败巨龙后 POST /internal/ai/narrative/generate（triggerEventId=BOSS_DRAGON_DOWN）→ 返回 MOURN/SORROW 台词，提及龙尸；ContentGuard 与 RuleTrigger 校验通过。",
                "探索途中伙伴 remember 坠崖/采花 → POST /api/player/ai/companion/dialogue 下发 MsgId 2500 气泡与 TTS 情绪；低血量 battle-coach 建议盾墙并注入小队。",
                "好感达标 → companion/passive-explore 刷新隐藏调查点，并在地图标出「伙伴觉得有趣的地方」。",
                "通关后 POST /api/player/ai/analyst/feedback?replayId=… → 口语建议「大招晚开 0.8s」+ 失误时间轴高亮 + equipHints 深链强化。",
                "迷路时玩家主动截图求助 POST /api/player/ai/vision/waypoint（visualTags=巨石）→ 前方距离与方向箭头；5 分钟后会话过期。",
                "设置页选择「感性」风格 POST /persona/style → 深渊竞速党解析为 HARDCORE 时战术建议带帧数/期望；休闲党 CASUAL 少提数值、语气温柔。",
            ]
            cur = h
            for s in stories:
                cur = insert_paragraph_after(cur, s, style="Normal")

    # ── FAQ Q15 ──
    if not already_has(doc, "Q15：P15"):
        q14 = find_paragraph(doc, lambda t, _: t.startswith("Q14：P14"))
        anchor = q14
        if q14 is not None:
            # 找到 Q14 的答段
            for p in doc.paragraphs:
                if (p.text or "").startswith("答：") and "OpenWorldP14" in (p.text or ""):
                    anchor = p
            h = insert_paragraph_after(
                anchor,
                "Q15：P15 动态叙事/伙伴/战斗分析/截图路点/人格如何联调？",
                style="Heading 2",
            )
            insert_paragraph_after(
                h,
                "答：ai-service InternalAiP15Controller：/narrative/generate、/companion/*、/analyst/feedback、"
                "/vision/waypoint、/persona/resolve|style。玩家侧 PlayerAiController："
                "/api/player/ai/analyst/feedback、/companion/dialogue、/vision/waypoint、/persona/style。"
                "协议 MsgId 2500。回归：mvn -pl mmorpg-common,ai-service,player-service test "
                "\"-Dtest=AiP15*,AiPlatformBusinessFlowTest,PlayerAiP15*\"。"
                "权威说明 docs/ai-enhancement.md（八.五 P15）。",
                style="Normal",
            )

    # ── 第十四章 API 速览 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if ("P14：" in t or "/predict/action" in t) and "P15：" not in t and ("P13" in t or "P14" in t):
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；P15：/internal/ai/narrative/generate、/companion/*、/analyst/feedback、/vision/waypoint、/persona/*；"
                "/api/player/ai/analyst/feedback|companion/dialogue|vision/waypoint|persona/style；MsgId 2500。",
            )
            break

    # ── 总结段 ──
    for p in doc.paragraphs:
        t = p.text or ""
        if "它覆盖玩家入口、大世界（含 P5～P14" in t and "P15" not in t:
            set_paragraph_text(
                p,
                t.replace("含 P5～P14", "含 P5～P15").replace(
                    "周词缀与命运残响等）",
                    "周词缀与命运残响等）与生成式 AI 沉浸（动态叙事/常驻伙伴/战斗分析/截图路点/动态人格）",
                ),
            )
            break
        if "它覆盖玩家入口、大世界（含 P5～P14" in t:
            break

    for p in doc.paragraphs:
        t = p.text or ""
        if "MyMmorpg 用现代 Java" in t and "P14" in t and "P15" not in t:
            set_paragraph_text(
                p,
                t.rstrip("。")
                + "；P15 动态叙事与世界状态共鸣、常驻 AI 伙伴、战斗连招诊断、截图路点、动态人格引擎。",
            )
            break

    # ── 附录日期 / 文档地图 ──
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("生成日期：") and "P14 体验纵深补齐 + 七大" in t and "P15" not in t:
            set_paragraph_text(
                p,
                "生成日期：2026年08月25日（P15 生成式 AI 沉浸补齐；P14 体验纵深补齐 + 七大体验短板加固；"
                "P13 体验优化；P12 性能加固；P10/P11 原地增补）",
            )

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if "docs/ai-enhancement.md" in cell.text and "P15" not in cell.text:
                    cell.text = cell.text.rstrip("。") + "（含 P15 八.五节）。"
                if "docs/open-world-top-tier.md" in cell.text and "P0～P14" in cell.text:
                    cell.text = cell.text.replace("P0～P14", "P0～P14（AI 沉浸见 ai-enhancement P15）")

    # ai-service 模块描述补 P15
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith("统一 AI 推理平台：") and "P15" not in t:
            set_paragraph_text(
                p,
                "统一 AI 推理平台：推荐/客服/NPC 长期记忆/内容校验/流失干预/MLOps；"
                "P15 动态叙事/伙伴/战斗分析/截图路点/人格（8995）",
            )
            break

    doc.save(str(DOC_PATH))
    print("OK: updated", DOC_PATH)
    print("BACKUP:", BACKUP)


if __name__ == "__main__":
    main()
