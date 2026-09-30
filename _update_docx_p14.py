# -*- coding: utf-8 -*-
"""在原 Word 文档中增补 P14 体验纵深补齐内容（不新建文档）。"""
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC_PATH = r"c:\Users\ASUS\Desktop\MyMmorpg项目全方位详细介绍文档.docx"


def set_paragraph_text(p, text):
    p.text = text


def insert_paragraph_after(paragraph, text, style=None):
  new_p = OxmlElement("w:p")
  paragraph._p.addnext(new_p)
  from docx.text.paragraph import Paragraph
  new_para = Paragraph(new_p, paragraph._parent)
  if style:
    try:
      new_para.style = style
    except Exception:
      pass
  new_para.add_run(text)
  return new_para


def find_paragraph(doc, predicate):
  for p in doc.paragraphs:
    t = (p.text or "").strip()
    if predicate(t, p):
      return p
  return None


def replace_if_contains(doc, old_sub, new_text):
  for p in doc.paragraphs:
    if old_sub in (p.text or ""):
      p.text = (p.text or "").replace(old_sub, new_text)


def update_table_cell(table, row_idx, col_idx, old_sub, new_sub):
  try:
    cell = table.rows[row_idx].cells[col_idx]
    if old_sub in cell.text:
      cell.text = cell.text.replace(old_sub, new_sub)
  except Exception:
    pass


def main():
  doc = Document(DOC_PATH)

  # ── 封面元信息 ──
  replace_if_contains(
      doc,
      "生成日期：2026年08月22日（P13 体验优化；P12 性能加固同日；P10/P11 原地增补同日）",
      "生成日期：2026年08月24日（P14 体验纵深补齐；P13 体验优化；P12 性能加固）",
  )
  replace_if_contains(
      doc,
      "修订说明：在原有 P6–P12",
      "修订说明：在原有 P6–P12（含性能加固 Actor 信箱/对象池/AOI 增量等）基础上，新增 P13 体验优化，并增补 P14 体验纵深补齐（客户端预测与软回滚、立体移动动量继承、确定性破坏、探索活力/神瞳共鸣、野外营地/幻影借用/助战印记、多端辅助与连招宏、周词缀轮换/命运残响）；技术栈截至 2026-08-24。请在 Word 中全选后按 F9 更新目录。",
  )
  # 若上面整段替换未命中（已是旧格式），直接改第 6 段
  p6 = find_paragraph(doc, lambda t, _: t.startswith("修订说明："))
  if p6 and "P14" not in (p6.text or ""):
    set_paragraph_text(
        p6,
        "修订说明：在原有 P6–P12（含性能加固 Actor 信箱/对象池/AOI 增量等）基础上，新增 P13 体验优化，并增补 P14 体验纵深补齐（客户端预测与软回滚、立体移动动量继承、确定性破坏、探索活力/神瞳共鸣、野外营地/幻影借用/助战印记、多端辅助与连招宏、周词缀轮换/命运残响）；技术栈截至 2026-08-24。请在 Word 中全选后按 F9 更新目录。",
    )

  # ── 1.4 项目定位：补充 P14 能力关键词 ──
  p14_kw = "客户端预测软回滚、动量继承、确定性破坏、探索活力、神瞳共鸣、野外营地、幻影借用、助战印记、多端辅助、周词缀轮换、命运残响"
  p132 = find_paragraph(doc, lambda t, _: "肉鸽命运卡与大世界探索闭环" in t)
  if p132 and p14_kw not in (p132.text or ""):
    p132.text = (p132.text or "").rstrip("。") + "，以及 P14 客户端预测软回滚、动量继承、确定性破坏、探索活力、神瞳共鸣、野外营地、多端辅助与周词缀轮换等。"

  # ── 3.2 大世界列表：补充 P14 bullet ──
  p3bullet = find_paragraph(doc, lambda t, _: t.startswith("【P10 活物生态"))
  if p3bullet:
    anchor = p3bullet
    for p in doc.paragraphs:
      if (p.text or "").startswith("【P10"):
        anchor = p
    # 在 P13 heading 之前插入 P14 bullet（若尚未存在）
    p13h = find_paragraph(doc, lambda t, _: t.startswith("3.2.8 探索便利"))
    if p13h and not find_paragraph(doc, lambda t, _: "【P14 体验纵深补齐】" in t):
      insert_paragraph_after(
          anchor,
          "【P14 体验纵深补齐】ClientPredictedActionService + ServerShadowService（500ms 影子快照）+ ReactionValidator.dynamicDodgeWindowMs（Base+RTT×0.5）+ PrePlaybackService 软回滚；SceneMoveCmd.inheritedVelocity + TraverseMomentumService 动量继承 + ClimbRestPointService 歇脚攀爬跳+30% + PoiseService 韧性霸体/闪避无敌帧；DeterministicMutationService 确定性破坏 + TerrainInteractionTracker.clientCacheCheck；ExplorationVitalityService / EcoMigrationScheduler / OculiResonanceService；CoopCampService / PhantomBorrowService / SocialTokenService；CombatAssistService.DeviceType + 软锁敌 + BuildRecommendationService 连招宏；AffixShuffleService + RogueFateCardService.exchangeFateEcho。",
          style="List Bullet",
      )

  # ── 第三章：3.2.9 P14 完整小节（插在 3.2.8 与 3.4 之间）──
  if not find_paragraph(doc, lambda t, _: t.startswith("3.2.9 体验纵深补齐")):
    p38 = find_paragraph(doc, lambda t, _: t.startswith("3.2.8 探索便利"))
    anchor = p38
    # 找到 3.2.8 最后一个 bullet（皮肤衣柜）之后、3.4 之前
    for p in doc.paragraphs:
      if (p.text or "").startswith("Heading 2|3.4") or (p.text or "").strip() == "3.4 社交与轻社交":
        break
      if (p.text or "").startswith("测试：OpenWorldP13"):
        anchor = p
    h = insert_paragraph_after(anchor, "3.2.9 体验纵深补齐（P14）", style="Heading 3")
    sections = [
        ("Normal", "操控粘滞感：ClientPredictedActionService 按下即播动画/位移；ServerShadowService 维护 500ms 位置快照；ReactionValidator.dynamicDodgeWindowMs = BaseWindow + RTT×0.5；PrePlaybackService 软回滚仅撤伤害/冷却，保留位移与动画（rubberBand=false）。"),
        ("Normal", "Z 轴衔接：SceneMoveCmd 增加 inheritedVelocity；TraverseMomentumService 滑翔→下落保留 70% 水平速度，钩锁摆荡派生攻击伤害系数；ClimbRestPointService 歇脚后攀爬跳跃 +30%；PoiseService 区分 POISE_ARMOR（可削韧）与 DODGE_IFRAME（闪避无敌帧）。"),
        ("Normal", "破坏即时反馈：DeterministicMutationService 下发 seed+timestamp 本地演算碎片；乐观地貌改写失败 1s 平滑回滚；TerrainInteractionTracker.clientCacheCheck 本地 terrain:cooldown:cache 减少重复请求。"),
        ("Normal", "探索惊喜：ExplorationVitalityService 每日随机调查点（叙事纸条+圣遗物碎片）；EcoMigrationScheduler 高探索度区域生物迁往相邻区；OculiResonanceService 神瞳收集 70% 触发 5 秒共鸣波提示剩余范围。"),
        ("Normal", "社交沉淀：CoopCampService 4 人 SAFE 区野外营地（烹饪锅/合成台/传送点，每日签到维护耐久）；PhantomBorrowService 解谜失败 3 次可借用好友操作影子；SocialTokenService 助战印记兑换限定外观。"),
        ("Normal", "多端适配：CombatAssistService.DeviceType（MOBILE +30ms 闪避窗/+15% 钩锁吸附；PC +5% 掉落）；软锁敌修正 ≤15°；BuildRecommendationService 连招宏 3 槽、触发间隔 ≥2s。"),
        ("Normal", "终局变量：AffixShuffleService 每周从 20 种词缀抽 5 种；RogueFateCardService.exchangeFateEcho 命运残响兑换圣遗物胚子；区域净化至 SAFE 后 Boss 狂暴觉醒（血量+20%/攻速+10%）。"),
        ("List Bullet", "API：POST .../predict/action/start|reconcile、POST .../combat/assist/device|soft-lock、POST .../build/macro/save|trigger、POST .../mutability/deterministic、GET .../explore/vitality/daily、GET .../explore/oculi/resonance、POST .../camp/establish|sign、POST .../phantom/borrow、POST .../social/token/grant、GET .../rogue/weekly-affix、POST .../rogue/fate-echo/exchange；reaction/open-window 支持 playerRttMs。"),
        ("List Bullet", "测试：OpenWorldP14FlowTest、OpenWorldP14EdgeCaseTest（common）；OpenWorldP14BusinessFlowTest（scene Internal API 全链路）；权威 docs/open-world-top-tier.md P14。"),
    ]
    cur = h
    for style, text in sections:
      cur = insert_paragraph_after(cur, text, style=style)

  # ── 第五章：5.2.12 P14 ──
  if not find_paragraph(doc, lambda t, _: t.startswith("5.2.12 P14")):
    p511 = find_paragraph(doc, lambda t, _: t.startswith("5.2.11 P13"))
    anchor = p511
    for p in doc.paragraphs:
      if (p.text or "").startswith("测试：OpenWorldP13FlowTest"):
        anchor = p
    h = insert_paragraph_after(anchor, "5.2.12 P14 体验纵深补齐（手感延迟 / 立体惯性 / 破坏反馈 / 探索惊喜 / 社交沉淀 / 多端适配 / 终局变量）", style="Heading 3")
    cur = insert_paragraph_after(
        h,
        "代码详见 docs/open-world-top-tier.md P14 节；门面 OpenWorldGameplayFacade；GET /gameplay/status 返回 p14 快照（clientPredict、serverShadowWindowMs、weeklyAffixCount 等）。",
        style="Normal",
    )
    bullets = [
        "手感：ClientPredictedActionService、ServerShadowService、ReactionValidator.dynamicDodgeWindowMs、PrePlaybackService 软回滚；SnapshotBuffer 已迁至 mmorpg-common。",
        "移动：SceneMoveCmd.inheritedVelocity*、TraverseMomentumService、ClimbRestPointService 攀爬跳加成、PoiseService 霸体/无敌帧分层。",
        "破坏：DeterministicMutationService、WorldMutabilityService.TerrainInteractionTracker.clientCacheCheck。",
        "探索：ExplorationVitalityService、EcoMigrationScheduler、OculiResonanceService。",
        "社交：CoopCampService、PhantomBorrowService、SocialTokenService。",
        "多端：CombatAssistService.DeviceType + resolveSoftLock；BuildRecommendationService 连招宏。",
        "终局：AffixShuffleService、RogueFateCardService.exchangeFateEcho、Boss 狂暴觉醒闭环。",
        "测试：OpenWorldP14FlowTest / OpenWorldP14EdgeCaseTest（common）；OpenWorldP14BusinessFlowTest（scene）；LagCompensationFlowTest / SnapshotBufferTest 回归。",
    ]
    for b in bullets:
      cur = insert_paragraph_after(cur, b, style="List Bullet")

  # ── 第十章成熟度：P14 bullet ──
  if not find_paragraph(doc, lambda t, _: t.startswith("P14 体验纵深补齐：")):
    p13m = find_paragraph(doc, lambda t, _: t.startswith("P13 体验优化："))
    if p13m:
      insert_paragraph_after(
          p13m,
          "P14 体验纵深补齐：客户端预测与软回滚、动量继承、确定性破坏、探索活力/神瞳共鸣、野外营地/幻影借用/助战印记、多端辅助/连招宏、周词缀与命运残响；已接入 OpenWorldGameplayFacade、InternalOpenWorldController。",
          style="List Bullet",
      )

  # ── 第十一章：11.15 业务故事 ──
  if not find_paragraph(doc, lambda t, _: t.startswith("11.15 客户端预测到命运残响")):
    p1114 = find_paragraph(doc, lambda t, _: t.startswith("11.14 探索罗盘到智能养成"))
    anchor = p1114
    for p in doc.paragraphs:
      if (p.text or "").startswith("肉鸽 POST /rogue/start 带 region_id"):
        anchor = p
    h = insert_paragraph_after(anchor, "11.15 客户端预测到命运残响（P14）", style="Heading 2")
    stories = [
        "玩家按下闪避 → POST /predict/action/start 立即本地播动画 → 服务端 reconcile 校验；失败则 softRollback 保留位移，仅回滚伤害数字。",
        "攻击开窗带 playerRttMs=80 → dodgeWindowMs 从 200 动态放宽到 240；高延迟下仍跟手。",
        "滑翔取消 → TraverseMomentumService 保留 70% 水平速度；攀爬命中歇脚点 → 下次攀爬跳跃距离 +30%。",
        "砍树 → POST /mutability/deterministic 拿 seed 本地碎裂；地形蘑菇冷却命中 clientCache 无需重复发包。",
        "每日 GET /explore/vitality/daily 刷新调查点 → 完成领奖；神瞳 70% GET /explore/oculi/resonance 触发共鸣波。",
        "SAFE 区 POST /camp/establish 建 4 人营地 → 每日 /camp/sign 维护耐久；解谜失败 3 次 POST /phantom/borrow 借用好友影子。",
        "手机端 POST /combat/assist/device MOBILE → 软锁敌 ≤15°；POST /build/macro/save 配置连招，2s 内重复触发被限流。",
        "GET /rogue/weekly-affix 拉取本周 5 词缀；通关获命运残响 POST /rogue/fate-echo/exchange 兑换定向主属性胚子。",
    ]
    cur = h
    for s in stories:
      cur = insert_paragraph_after(cur, s, style="Normal")

  # ── 第十四章 API 速览：追加 P14 ──
  p14api = find_paragraph(doc, lambda t, _: "P13：/compass/*" in t)
  if p14api and "P14：" not in (p14api.text or ""):
    p14api.text = (p14api.text or "").rstrip("。") + "；P14：/predict/action/*、/combat/assist/device|soft-lock、/build/macro/*、/mutability/deterministic、/explore/vitality/*、/explore/oculi/resonance、/camp/*、/phantom/borrow、/social/token/grant、/rogue/weekly-affix、/rogue/fate-echo/exchange；reaction/open-window 支持 playerRttMs。"

  # ── FAQ Q14 ──
  if not find_paragraph(doc, lambda t, _: t.startswith("Q14：P14")):
    q13 = find_paragraph(doc, lambda t, _: t.startswith("Q13：P13"))
    anchor = q13
    for p in doc.paragraphs:
      if (p.text or "").startswith("答：GET /internal/scene/open-world/metrics/performance"):
        anchor = p
    h = insert_paragraph_after(anchor, "Q14：P14 客户端预测/多端辅助/探索活力/肉鸽词缀如何联调？", style="Heading 2")
    insert_paragraph_after(
        h,
        "答：场景侧 InternalOpenWorldController：/predict/action/start|reconcile、/combat/assist/device|soft-lock、/build/macro/save|trigger、/mutability/deterministic、/explore/vitality/daily|complete、/explore/oculi/resonance、/camp/establish|sign、/phantom/borrow、/social/token/grant、/rogue/weekly-affix、/rogue/fate-echo/exchange。回归：mvn -pl mmorpg-common,scene-service test \"-Dtest=OpenWorldP14*\"。权威说明 docs/open-world-top-tier.md（P14）。",
        style="Normal",
    )

  # ── 总结：P5～P14 ──
  replace_if_contains(doc, "它覆盖玩家入口、大世界（含 P5～P13", "它覆盖玩家入口、大世界（含 P5～P14")

  # ── 附录 C ──
  replace_if_contains(
      doc,
      "生成日期：2026年08月22日（P12 性能加固）",
      "生成日期：2026年08月24日（P14 体验纵深补齐；P12 性能加固；P10/P11 原地增补）",
  )

  # ── 附录表格：open-world-top-tier.md 行 ──
  for table in doc.tables:
    for row in table.rows:
      for cell in row.cells:
        if "docs/open-world-top-tier.md" in cell.text and "P0～P11" in cell.text:
          cell.text = cell.text.replace("P0～P11", "P0～P14")

  # ── 移动预测 bullet 补充 SnapshotBuffer 位置 ──
  replace_if_contains(
      doc,
      "移动延迟补偿（SnapshotBuffer）与客户端预测校验",
      "移动延迟补偿（SnapshotBuffer，已迁至 mmorpg-common）与客户端预测校验；P14 扩展 ServerShadowService + ClientPredictedActionService 软回滚",
  )

  doc.save(DOC_PATH)
  print("OK: updated", DOC_PATH)


if __name__ == "__main__":
  main()
